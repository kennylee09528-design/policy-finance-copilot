"""PPT 生成 Agent（大模型以 mock 替代）：需求确认 → 大纲（数字只能占位）→ 生成 .pptx（原生图表/表格）→ 数字溯源；敏感模型约束；提示词版本；模板上传。"""
import io
import json as J

import httpx
import pytest
from pptx import Presentation

from app import admin_auth, db
from app.llm import service as L
from app.ppt import agent as G, render as R

CALLS = []


def fake_llm(url, json=None, headers=None, timeout=None):
    msgs = json["messages"]
    sys_p = next((m["content"] for m in msgs if m["role"] == "system"), "")
    user = msgs[-1]["content"]
    CALLS.append({"model": json["model"], "system": sys_p, "user": user})
    if "先与用户确认需求" in user:
        out = {"questions": ["汇报对象是谁？", "是否突出预警机构？", "页数控制？"],
               "sheet": {"audience": "公司经营层", "purpose": "汇报额度使用", "pages": 9, "tone": "简洁", "sections": ["总体情况", "分机构", "预警"]}}
    else:
        out = {"pages": [
            {"kind": "cover"}, {"kind": "agenda", "title": "目录"},
            {"kind": "section", "title": "总体情况"},
            {"kind": "kpi", "title": "核心指标", "data": ["F1", "F2", "F8"]},
            {"kind": "content", "title": "总体判断", "bullets": ["常规业务使用率为 {{F2}}，略低于时间进度 {{F8}}", "全年预计增长 30%", "需加快推进"], "data": ["F1"]},
            {"kind": "chart", "title": "分机构使用率", "data": "C3"},
            {"kind": "chart", "title": "不存在的图", "data": "C99"},
            {"kind": "table", "title": "明细", "data": "T1"},
            {"kind": "section", "title": "第3部分 预警"},
            {"kind": "text", "title": "预警事项", "data": "X1"},
            {"kind": "closing"}]}
    return httpx.Response(200, json={"choices": [{"message": {"content": "<think>思考</think>```json\n" + J.dumps(out, ensure_ascii=False) + "\n```"}}],
                                     "model": json["model"], "usage": {"total_tokens": 50}})


@pytest.fixture(scope="module")
def env(client):
    admin_auth.reset_password("admin", "Init-pass-ppt")
    client.post("/api/admin/login", json={"username": "admin", "password": "Init-pass-ppt"})
    client.post("/api/admin/password", json={"old": "Init-pass-ppt", "new": "PPT-Admin-2026!x"})
    if not db.q1("SELECT id FROM m5_snapshots"):
        assert client.post("/api/m5/snapshots/sample", json={"stat_date": "2026-09-19"}).status_code == 200
    if not db.q1("SELECT id FROM m4_schemes"):
        assert client.post("/api/m4/import-sample").status_code == 200
    u = {"name": "测试", "role": "admin"}
    pub_p = L.save_provider({"vendor": "deepseek", "base_url": "https://pub-ppt.example/v1", "api_key": "sk-pub-xxxxxxxx"}, u)
    pri_p = L.save_provider({"vendor": "company_private", "base_url": "https://priv-ppt.example/v1", "api_key": "sk-pri-xxxxxxxx"}, u)
    return {"pub": L.save_model({"provider_id": pub_p, "model_id": "ppt-pub"}, u),
            "pri": L.save_model({"provider_id": pri_p, "model_id": "ppt-priv", "allow_sensitive": True}, u)}


def test_sanitize_rules():
    pack = {"facts": {"F1": {"value": "12.34%"}}}
    assert G.sanitize("使用率为 {{F1}}", pack) == "使用率为 12.34%"
    assert G.sanitize("增长 30%", pack) is None and G.sanitize("第３名", pack) is None and G.sanitize("{{F9}}", pack) is None


def test_public_model_allowed_with_notice(client, env, monkeypatch):
    """2026-10-12 用户决定：含内部数据的演示稿也可用公网模型，只给非阻断提示，并在留痕中记录模型。"""
    monkeypatch.setattr("app.llm.client.httpx.post", fake_llm)
    r = client.post("/api/ppt/jobs", json={"deck_type": "usage", "model_id": env["pub"]})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["model_notice"] and "公司自有大模型" in j["model_notice"]
    j = client.post(f"/api/ppt/jobs/{j['id']}/sheet", json={}).json()
    g = client.post(f"/api/ppt/jobs/{j['id']}/generate", json={"confirm": True, "thumbnails": False})
    assert g.status_code == 200, g.text
    a = db.q1("SELECT detail FROM audit_log WHERE action='生成 PPT' AND target=? ORDER BY id DESC", (str(j["id"]),))
    assert "ppt-pub" in a["detail"]
    r = client.post("/api/ppt/jobs", json={"deck_type": "scheme", "model_id": env["pub"]})
    assert r.status_code == 200
    r = client.post("/api/ppt/jobs", json={"deck_type": "usage", "model_id": env["pri"]})
    assert r.status_code == 200 and not r.json()["model_notice"]


def test_usage_flow_with_llm(client, env, monkeypatch, tmp_path):
    monkeypatch.setattr("app.llm.client.httpx.post", fake_llm)
    CALLS.clear()
    j = client.post("/api/ppt/jobs", json={"deck_type": "usage", "model_id": env["pri"], "scope": {"period": "周报"}}).json()
    assert j["status"] == "待确认需求" and j["questions"][0] == "汇报对象是谁？" and j["sheet"]["audience"] == "公司经营层"
    # 大模型看不到任何数值：只看到数据目录
    f1 = j["catalog"]["facts"]["F1"]["value"]
    assert f1.replace(",", "")[:5] not in CALLS[0]["user"].replace(",", "")
    j = client.post(f"/api/ppt/jobs/{j['id']}/sheet", json={"answers": ["经营层，突出预警"], "sheet": {"pages": 10}}).json()
    assert j["status"] == "待确认大纲" and j["sheet_confirmed_at"]
    kinds = [p["kind"] for p in j["outline"]]
    assert kinds[0] == "cover" and kinds[-1] == "closing" and "kpi" in kinds
    content = next(p for p in j["outline"] if p["kind"] == "content")
    assert content["bullets"] == ["常规业务使用率为 {{F2}}，略低于时间进度 {{F8}}", "需加快推进"]       # 含字面数字的要点被丢弃
    assert not any(p.get("data") == "C99" for p in j["outline"])
    assert any("30%" in w for w in j["warnings"]) and any("C99" in w for w in j["warnings"])
    assert any(p["title"] == "预警" or p["kind"] == "section" and "3" not in p["title"] for p in j["outline"])
    prev = next(p for p in j["preview"] if p["kind"] == "content")
    assert j["catalog"]["facts"]["F2"]["value"] in prev["bullets"][0]
    # 用户编辑大纲：删去表格页、修改标题；写入字面数字的要点同样被拒
    pages = [p for p in j["outline"] if p["kind"] != "table"]
    pages[4]["title"] = "总体研判"
    pages[4]["bullets"] = ["人工补充 5 家机构"] + pages[4]["bullets"]
    o = client.put(f"/api/ppt/jobs/{j['id']}/outline", json={"pages": pages}).json()
    assert any("5 家" in w for w in o["outline_warnings"]) and not any(p["kind"] == "table" for p in o["outline"])
    assert client.post(f"/api/ppt/jobs/{j['id']}/generate", json={}).status_code == 400   # 须确认
    g = client.post(f"/api/ppt/jobs/{j['id']}/generate", json={"confirm": True, "thumbnails": False}).json()
    assert g["status"] == "已生成" and len(g["files"]) == 1
    f = g["files"][0]
    assert f["trace"]["ok"], f["trace"]["violations"]
    assert f["trace"]["numbers_checked"] > 5 and f["trace"]["chart_series_checked"] >= 1
    d = client.get(f"/api/ppt/files/{f['id']}/download")
    assert d.status_code == 200 and d.content[:2] == b"PK"
    prs = Presentation(io.BytesIO(d.content))
    assert prs.slide_width == R.Inches(13.333)
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert charts  # 原生可编辑图表
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert j["catalog"]["facts"]["F2"]["value"] in texts and "{{" not in texts and "30%" not in texts
    # 第二版
    g2 = client.post(f"/api/ppt/jobs/{j['id']}/generate", json={"confirm": True, "thumbnails": False}).json()
    assert [x["version"] for x in g2["files"]] == [2, 1]


def test_trace_catches_injected_number(client, env, tmp_path):
    from app.ppt.data import build
    pack = build("usage", {})
    pages, _ = G.normalize(G.rule_outline(pack, {}), pack)
    path = str(tmp_path / "x.pptx")
    R.render(G.rendered(pages, pack), pack, path)
    assert R.trace(path, pack)["ok"]
    prs = Presentation(path)
    R._box(prs.slides[2], 0, 0, R.Inches(3), R.Inches(1), "编造数字 98765.43 万元")
    prs.save(path)
    tr = R.trace(path, pack)
    assert not tr["ok"] and tr["violations"][0]["number"] == "98765.43"


@pytest.mark.parametrize("deck,scope", [("scheme", {}), ("policy", None), ("research", {})])
def test_other_decks_rule_based(client, env, deck, scope, monkeypatch):
    """无模型时按规则生成（需求单默认、规则大纲），仍走两次人工确认。"""
    monkeypatch.setattr(G, "model_of", lambda j=None: None)
    if deck == "policy":
        r = db.q1("SELECT MIN(brief_date) a, MAX(brief_date) b FROM m2_briefs")
        if not r["a"]:
            pytest.skip("无 M2 简报")
        scope = {"date_from": r["a"], "date_to": r["b"]}
    if deck == "research" and not db.q1("SELECT id FROM m3_tasks WHERE status IN ('已定稿','已总结')"):
        pytest.skip("无已定稿 M3 任务")
    j = client.post("/api/ppt/jobs", json={"deck_type": deck, "scope": scope}).json()
    assert j["status"] == "待确认需求", j
    j = client.post(f"/api/ppt/jobs/{j['id']}/sheet", json={}).json()
    g = client.post(f"/api/ppt/jobs/{j['id']}/generate", json={"confirm": True, "thumbnails": False}).json()
    f = g["files"][0]
    assert f["trace"]["ok"], f["trace"]["violations"] and f["pages"] >= 4


def test_admin_prompts_and_templates(client, env, monkeypatch, tmp_path):
    monkeypatch.setattr("app.llm.client.httpx.post", fake_llm)
    c2 = httpx.Client(transport=client._transport, base_url=str(client.base_url))
    assert c2.get("/api/admin/ppt/config").status_code == 403
    assert client.put("/api/admin/ppt/config", json={"model_id": env["pri"], "temperature": 0.2}).json()["model_id"] == env["pri"]
    r = client.post("/api/admin/ppt/prompts", json={"deck_type": "usage", "content": "【V2PPT】只写结构，不写数字。", "note": "v2"}).json()
    vers = [p for p in r["prompts"] if p["deck_type"] == "usage"]
    assert vers[0]["version"] == 2 and vers[0]["active"] == 1
    CALLS.clear()
    j = client.post("/api/ppt/jobs", json={"deck_type": "usage"}).json()
    assert CALLS and "【V2PPT】" in CALLS[0]["system"] and CALLS[0]["model"] == "ppt-priv" and j["prompt_version"] == 2
    client.post(f"/api/admin/ppt/prompts/{vers[-1]['id']}/activate")
    # 上传自定义模板（用默认主题生成一个带版式的 .pptx 作为模板）
    from pptx import Presentation as P
    tp = tmp_path / "tpl.pptx"
    P().save(tp)
    t = client.post("/api/admin/ppt/templates", files={"file": ("公司模板.pptx", tp.read_bytes())}, data={"name": "公司模板"}).json()
    assert t["layouts"] and t["name"] == "公司模板"
    assert client.post("/api/admin/ppt/templates", files={"file": ("x.txt", b"abc")}).status_code == 400
    client.post(f"/api/admin/ppt/templates/{t['id']}/activate")
    client.post(f"/api/ppt/jobs/{j['id']}/sheet", json={})
    g = client.post(f"/api/ppt/jobs/{j['id']}/generate", json={"confirm": True, "thumbnails": False, "template_id": t["id"]}).json()
    assert g["files"][0]["template_id"] == t["id"] and g["files"][0]["trace"]["ok"], g["files"][0]["trace"]
    client.post("/api/admin/ppt/templates/0/activate")
    assert client.delete(f"/api/admin/ppt/templates/{t['id']}").status_code == 200


def test_policy_empty_brief_charts_regression(client, env, monkeypatch):
    """回归 2026-10-12：区间内简报均无条目 → C1/C2 图表无分类，生成时 python-pptx 抛
    ValueError(chart data contains no categories) 导致 500。空图表应在大纲规整时删除并提示。"""
    monkeypatch.setattr(G, "model_of", lambda j=None: None)
    items = {"items": [], "data_rows": [{"name": "常规业务使用率（÷国担额度）", "value": "92.80%", "delta": "—", "source": "M5"}]}
    for d in ("2031-01-05", "2031-01-06"):
        db.execute("INSERT INTO m2_briefs(brief_date,version,status,items,markdown,stats,config,model,generated_at,generated_by) VALUES(?,?,?,?,?,?,?,?,?,?)",
                   (d, 1, "草稿", J.dumps(items), "", "{}", "{}", "", db.now(), "test"))
    j = client.post("/api/ppt/jobs", json={"deck_type": "policy", "scope": {"date_from": "2031-01-05", "date_to": "2031-01-06"}}).json()
    j = client.post(f"/api/ppt/jobs/{j['id']}/sheet", json={}).json()
    pages = [{"kind": "cover", "title": "政策简报汇编"}, {"kind": "kpi", "title": "概况", "data": ["F1", "F5"]},
             {"kind": "chart", "title": "按类别统计", "data": "C1"}, {"kind": "chart", "title": "按来源等级统计", "data": "C2"},
             {"kind": "table", "title": "数据速览", "data": "T4"}, {"kind": "closing", "title": "谢谢"}]
    s = client.put(f"/api/ppt/jobs/{j['id']}/outline", json={"pages": pages}).json()
    assert not any(p["kind"] == "chart" for p in s["outline"]) and any("图表" in w for w in s["outline_warnings"])
    # 旧任务里已存了含空图表的大纲（线上 #4/#7），生成时也必须成功
    db.execute("UPDATE ppt_jobs SET outline=? WHERE id=?", (J.dumps(pages, ensure_ascii=False), j["id"]))
    r = client.post(f"/api/ppt/jobs/{j['id']}/generate", json={"confirm": True, "thumbnails": False})
    assert r.status_code == 200, r.text
    f = r.json()["files"][0]
    assert f["trace"]["ok"]
    prs = Presentation(db.q1("SELECT path FROM ppt_files WHERE id=?", (f["id"],))["path"])
    assert not any(sh.has_chart for sl in prs.slides for sh in sl.shapes)


def _policy_job(client, monkeypatch, dates, briefs):
    monkeypatch.setattr(G, "model_of", lambda j=None: None)
    for d, v, st, items in briefs:
        db.execute("INSERT INTO m2_briefs(brief_date,version,status,items,markdown,stats,config,model,generated_at,generated_by) VALUES(?,?,?,?,?,?,?,?,?,?)",
                   (d, v, st, J.dumps(items, ensure_ascii=False), "", "{}", "{}", "", db.now(), "test"))
    j = client.post("/api/ppt/jobs", json={"deck_type": "policy", "scope": {"date_from": dates[0], "date_to": dates[1]}}).json()
    return client.post(f"/api/ppt/jobs/{j['id']}/sheet", json={"sheet": {"audience": "普惠业务部全体人员"}}).json()


def _rows(stat):
    return [{"name": "常规业务使用率（÷国担额度）", "value": "92.80%", "delta": "—", "source": f"M5 机构看板（统计时点 {stat}，快照 #7）"}]


def test_policy_grounded_bullets_latest_snapshot_cover(client, env, monkeypatch):
    """政策简报：无条目时不得出现大模型自拟要点，只写「本期暂无条目」；数据速览必带 T4 且取最新一期（含草稿）并标注；封面副标题=日期范围+受众。"""
    item = {"url": "https://www.gov.cn/x1", "title": "关于进一步支持小微企业融资担保的通知", "tier": 3, "category": "融资担保",
            "region": "国家", "publisher": "财政部", "level": "A1", "points": ["降低担保费率"]}
    j = _policy_job(client, monkeypatch, ("2032-03-01", "2032-03-02"), [
        ("2032-03-01", 1, "已发布", {"items": [item], "data_rows": _rows("2032-02-20")}),
        ("2032-03-02", 1, "已发布", {"items": [], "data_rows": _rows("2032-02-25")}),
        ("2032-03-02", 2, "草稿", {"items": [], "data_rows": _rows("2032-03-02")})])
    pages = [{"kind": "cover"}, {"kind": "section", "title": "今日必读"},
             {"kind": "content", "title": "今日必读 要点", "bullets": ["聚焦国家重大战略部署与政策动向"]},
             {"kind": "section", "title": "值得关注"},
             {"kind": "content", "title": "值得关注 要点", "bullets": ["关注区域经济发展政策动态"]},
             {"kind": "section", "title": "数据速览"}, {"kind": "closing"}]
    s = client.put(f"/api/ppt/jobs/{j['id']}/outline", json={"pages": pages}).json()
    o = s["outline"]
    allb = [b for p in o for b in p["bullets"]]
    assert "聚焦国家重大战略部署与政策动向" not in allb and "关注区域经济发展政策动态" not in allb
    i = [p["title"] for p in o].index("今日必读")
    assert o[i + 1]["data"] == "X3"  # 有条目 → 挂条目要点
    k = [p["title"] for p in o].index("值得关注")
    assert o[k + 1]["bullets"] == ["本期暂无条目"]
    m = [p["title"] for p in o].index("数据速览")
    assert o[m + 1]["data"] == "T4"
    pack = J.loads(db.q1("SELECT pack FROM ppt_jobs WHERE id=?", (j["id"],))["pack"])
    t4 = pack["tables"]["T4"]
    assert t4["title"] == "数据速览（截至 2032-03-02）" and len(t4["title"]) <= 20
    assert "第2版" in t4["note"] and "草稿" in t4["note"] and "2032-03-02" in t4["rows"][0][3]
    g = client.post(f"/api/ppt/jobs/{j['id']}/generate", json={"confirm": True, "thumbnails": False}).json()
    f = g["files"][-1]
    assert f["trace"]["ok"], f["trace"]
    prs = Presentation(db.q1("SELECT path FROM ppt_files WHERE id=?", (f["id"],))["path"])
    texts = ["\n".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in prs.slides]
    assert "2032-03-01 至 2032-03-02" in texts[0] and "普惠业务部全体人员" in texts[0]
    assert "本期暂无条目" in "\n".join(texts)


def test_template_cover_subtitle_not_title(tmp_path):
    """模板封面：SUBTITLE 占位符不能被当成标题（之前封面「政策简报汇编」出现两遍、封底「谢谢！」两遍）。"""
    from pptx import Presentation as P
    path = tmp_path / "tpl.pptx"
    P().save(path)
    d = R.Deck(str(path))
    d.cover("政策简报汇编", "2026-10-05 至 2026-10-12 · 面向：普惠业务部", "某机构", "")
    d.closing("某机构")
    sl = list(d.prs.slides)
    t0 = [sh.text_frame.text for sh in sl[0].shapes if sh.has_text_frame]
    assert t0.count("政策简报汇编") == 1 and any("2026-10-05" in t for t in t0)
    t1 = [sh.text_frame.text for sh in sl[1].shapes if sh.has_text_frame]
    assert t1.count("谢谢！") == 1


def test_policy_uses_newest_version_items(client, env, monkeypatch):
    """同日已发布旧版无条目、之后草稿有条目 → 取最新草稿的条目。"""
    item = {"url": "https://www.gov.cn/x2", "title": "关于做好政府性融资担保工作的通知", "tier": 2, "category": "融资担保",
            "region": "国家", "publisher": "国务院办公厅", "level": "A1", "points": ["扩大覆盖面"]}
    j = _policy_job(client, monkeypatch, ("2033-04-01", "2033-04-01"), [
        ("2033-04-01", 1, "已发布", {"items": [], "data_rows": _rows("2033-03-30")}),
        ("2033-04-01", 2, "草稿", {"items": [item], "data_rows": _rows("2033-04-01")})])
    pack = J.loads(db.q1("SELECT pack FROM ppt_jobs WHERE id=?", (j["id"],))["pack"])
    assert pack["facts"]["F1"]["raw"] == 1 and "第2版" in pack["source"]["versions"][0]
