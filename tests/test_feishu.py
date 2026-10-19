"""飞书机器人：mock 飞书 API + mock 事件，无需真实凭据。
覆盖：配置加密不回显、绑定码、未绑定只给绑定说明、角色权限（机构只看本机构）、额度/预警/简报/政策/报告/报告文件、
去重、限流、群聊@、群聊不泄露内部数据、发邮箱二次确认、PPT 多轮、LLM 路由只收提问文字、HTTP 回调校验/解密、SDK 长连接适配器、审计。"""
import base64
import hashlib
import json
import os
import smtplib
import uuid
from urllib.parse import quote

import pytest
from fastapi.testclient import TestClient

from app import admin_auth, db
from app import users as U
from app.feishu import api as FA
from app.feishu import config as C
from app.feishu import identity as I
from app.feishu import service as S

INT = {"X-User": quote("fs_zhang"), "X-Role": "internal"}
INST = {"X-User": quote("fs_inst"), "X-Role": "institution"}
BOT = "ou_bot_000001"


class FakeAPI:
    def __init__(self):
        self.sent, self.files = [], []

    def bot_info(self): return {"open_id": BOT, "name": "普惠助手"}
    def token(self): return "t"
    def reply_text(self, mid, text): self.sent.append(("reply", mid, text))
    def send_text(self, chat, text): self.sent.append(("send", chat, text))
    def upload_file(self, name, data): self.files.append((name, len(data), data)); return "fk_" + name
    def send_file(self, chat, key): self.sent.append(("file", chat, key))

    @property
    def texts(self): return [x[2] for x in self.sent if x[0] != "file"]


pytestmark = pytest.mark.usefixtures("isolated_db")


@pytest.fixture(scope="module")
def adm():
    from app.main import app
    c = TestClient(app)
    admin_auth.reset_password("admin", "Feishu-Admin-2026x")
    assert c.post("/api/admin/login", json={"username": "admin", "password": "Feishu-Admin-2026x"}).status_code == 200
    assert c.post("/api/admin/password", json={"old": "Feishu-Admin-2026x", "new": "Feishu-Admin-2026y"}).status_code == 200
    return c


@pytest.fixture(autouse=True)
def fake(monkeypatch):
    f = FakeAPI()
    monkeypatch.setattr(FA, "get_api", lambda: f)
    S._hits.clear()
    return f


def _cfg(adm, **kw):
    body = {"enabled": True, "app_id": "cli_test123456", "app_secret": "secret-value-XYZ-123", "verification_token": "vt-token-123", "encrypt_key": "",
            "mode": "ws", "rate_per_min": 50, **kw}
    r = adm.put("/api/admin/feishu/config", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def ev(open_id, text, chat="p2p", mention=False, mtype="text", event_id=None):
    mid = "om_" + uuid.uuid4().hex[:10]
    msg = {"message_id": mid, "chat_id": "oc_" + ("g1" if chat == "group" else open_id[-4:]), "chat_type": "group" if chat == "group" else "p2p", "message_type": mtype,
           "content": json.dumps({"text": ("@_user_1 " if mention else "") + text}, ensure_ascii=False)}
    if mention:
        msg["mentions"] = [{"key": "@_user_1", "id": {"open_id": BOT}, "name": "普惠助手"}]
    return {"schema": "2.0", "header": {"event_id": event_id or "ev_" + uuid.uuid4().hex[:10], "event_type": "im.message.receive_v1", "token": "vt-token-123", "app_id": "cli_test123456"},
            "event": {"sender": {"sender_id": {"open_id": open_id}, "sender_type": "user"}, "message": msg}}


def say(f, open_id, text, **kw):
    n = len(f.sent)
    S.dispatch(ev(open_id, text, **kw), sync=True)
    return "\n".join(x[2] if x[0] != "file" else f"[文件:{x[2]}]" for x in f.sent[n:])


@pytest.fixture(scope="module")
def world(adm):
    """M5 样例快照 + 一家机构用户 + 内部用户（均已完成改密并绑定飞书）。"""
    from app.main import app
    c = TestClient(app)
    _cfg(adm)
    assert c.post("/api/m5/snapshots/sample", json={"stat_date": "2026-09-19"}, headers=INT).status_code == 200
    s = S._latest_snapshot()
    d = c.get(f"/api/m5/snapshots/{s['id']}/dashboard", headers=INT).json()
    names = [r["name"] for r in d["rows"] if r["regular_limit"]]
    assert len(names) >= 2
    mine, other = names[0], names[1]
    for n in (mine, other):
        if not db.q1("SELECT 1 x FROM institutions WHERE full_name=?", (n,)):
            db.execute("INSERT INTO institutions(full_name,short_name,created_at,updated_at) VALUES(?,?,?,?)", (n, n[:4], db.now(), db.now()))
    admin = {"name": "admin", "role": "admin"}
    U.create({"username": "fs_zhang", "display_name": "张三", "role": "internal"}, admin)
    U.create({"username": "fs_inst", "display_name": "某机构用户", "role": "institution", "org": mine}, admin)
    db.execute("UPDATE users SET must_change=0")
    return {"c": c, "mine": mine, "other": other, "snap": s, "dash": d}


def _bind(c, open_id, headers):
    code = c.post("/api/feishu/bind-code", headers=headers).json()["code"]
    return code


def test_config_secrets_never_returned(adm):
    out = _cfg(adm)
    txt = json.dumps(out) + adm.get("/api/admin/feishu/config").text
    assert "secret-value-XYZ-123" not in txt and "vt-token-123" not in txt
    assert out["has_app_secret"] and out["has_verification_token"] and not out["has_encrypt_key"]
    raw = db.q1("SELECT value FROM settings WHERE key='feishu_config'")["value"]
    assert "secret-value-XYZ-123" not in raw and "vt-token-123" not in raw  # 加密存储
    assert C.secret(C.get(), "app_secret") == "secret-value-XYZ-123"
    _cfg(adm, app_secret="")  # 留空 = 不改
    assert C.secret(C.get(), "app_secret") == "secret-value-XYZ-123"
    assert adm.put("/api/admin/feishu/config", json={"app_id": "bad id!", "mode": "ws"}).status_code == 400
    assert adm.put("/api/admin/feishu/config", json={"app_id": "cli_test123456", "app_secret": "含 空格"}).status_code == 400
    assert TestClient(__import__("app.main", fromlist=["app"]).app).get("/api/admin/feishu/config").status_code == 403


def test_unbound_gets_instructions_only_and_bind_flow(world, fake):
    c = world["c"]
    r = say(fake, "ou_unbound_1", "额度")
    assert "绑定" in r and "万元" not in r
    assert db.q1("SELECT status, username FROM feishu_events WHERE open_id='ou_unbound_1' ORDER BY id DESC")["status"] == "未绑定"
    # 绑定
    assert "无效" in say(fake, "ou_zhang", "绑定 ABCD2345")
    code = _bind(c, "ou_zhang", INT)
    assert "绑定成功" in say(fake, "ou_zhang", f"绑定 {code}")
    assert "无效" in say(fake, "ou_other", f"绑定 {code}")  # 一次性
    assert db.q1("SELECT username FROM feishu_bindings WHERE open_id='ou_zhang'")["username"] == "fs_zhang"
    assert "******" in db.q1("SELECT text FROM feishu_events WHERE open_id='ou_zhang' AND text LIKE '绑定%' ORDER BY id DESC")["text"]  # 日志不留绑定码
    # 群聊里不允许用绑定码
    code2 = _bind(c, "ou_inst", INST)
    assert "私聊" in say(fake, "ou_inst", f"绑定 {code2}", chat="group", mention=True)
    assert "绑定成功" in say(fake, "ou_inst", f"绑定 {code2}")
    # 过期
    db.execute("UPDATE feishu_bind_codes SET expires_at='2000-01-01 00:00:00'")
    code3 = _bind(c, "ou_x", INT)
    db.execute("UPDATE feishu_bind_codes SET expires_at='2000-01-01 00:00:00' WHERE username='fs_zhang' AND used_at IS NULL")
    assert "过期" in say(fake, "ou_y", f"绑定 {code3}")
    # 管理员映射 + 解绑
    assert I.unbind(open_id="ou_inst") == 1 and "绑定" in say(fake, "ou_inst", "额度")
    I._set("ou_inst", "fs_inst", "admin")


def test_quota_and_permissions(world, fake):
    mine, other, s = world["mine"], world["other"], world["snap"]
    r = say(fake, "ou_zhang", "额度")
    assert "数据来源：M5 机构看板" in r and f"快照 #{s['id']}" in r and s["stat_date"] in r and "查询时间" in r
    reg = world["dash"]["summary"]["regular"]
    assert f"{reg['total']['amount']:,.2f}" in r  # 数字来自 M5 看板，非大模型
    r = say(fake, "ou_zhang", f"{mine} 额度使用情况")
    assert mine in r and "使用率" in r
    # 机构用户：只能看到本机构
    r = say(fake, "ou_inst", "额度")
    assert mine in r and other not in r and "全网" not in r
    r2 = say(fake, "ou_inst", f"查一下{other}的额度")
    assert other not in r2 and mine in r2
    assert "仅对普惠部内部" in say(fake, "ou_inst", "政策 担保")
    assert "仅对普惠部内部" in say(fake, "ou_inst", "报告")
    assert "仅对普惠部内部" in say(fake, "ou_inst", "生成PPT")
    assert "仅对普惠部内部" in say(fake, "ou_inst", "发邮箱 报告 #1 a@example.com")


def test_alerts(world, fake):
    r = say(fake, "ou_zhang", "预警")
    assert "预警" in r and ("数据来源：M5" in r)
    r2 = say(fake, "ou_inst", "预警列表")
    assert world["other"] not in r2


def test_brief_policy_reports(world, fake):
    db.execute("INSERT INTO m2_briefs(brief_date,version,status,items,markdown,stats,config,model,published_at) VALUES(?,?,?,?,?,?,?,?,?)",
               ("2026-10-09", 1, "已发布", db.jdump({"items": [{"title": "央行降准释放流动性", "publisher": "中国人民银行", "level": "A1", "url": "http://pbc.example/1", "section": "必读",
                                                              "category": "货币政策", "include": True, "points": ["降准0.5个百分点"], "pub_date": "2026-10-09"}], "data_rows": []}),
                "# x", "{}", "{}", "", db.now()))
    db.execute("INSERT INTO m2_briefs(brief_date,version,status,items,markdown,stats,config,model) VALUES(?,?,?,?,?,?,?,?)",
               ("2026-10-10", 1, "草稿", db.jdump({"items": [{"title": "草稿条目不应对机构可见", "publisher": "x", "level": "B1", "url": "u", "section": "必读", "category": "c", "include": True}], "data_rows": []}), "# y", "{}", "{}", ""))
    r = say(fake, "ou_inst", "今日简报")
    assert "央行降准" in r and "中国人民银行（A1）" in r and "草稿条目" not in r and "已发布" in r  # 机构只看已发布
    r = say(fake, "ou_zhang", "今日简报 2026-10-10")  # 内部账号可指定日期查看草稿（标注状态），默认优先已发布
    assert "草稿条目" in r and "草稿" in r
    assert "央行降准" in say(fake, "ou_zhang", "今日简报")
    db.execute("INSERT INTO m1_items(url,title,publisher,level,pub_date,snippet,collected_at) VALUES(?,?,?,?,?,?,?)",
               ("http://gov.example/p1", "关于进一步降低融资担保费率的通知", "财政部", "A1", "2026-09-30", "担保费率", db.now()))
    r = say(fake, "ou_zhang", "政策 担保费率")
    assert "财政部" in r and "权威分级：A1" in r and "http://gov.example/p1" in r and "数据来源：M1" in r
    # 报告
    done = db.execute("INSERT INTO m3_tasks(title,requirements,status,stage,final,final_check,confirmations,tokens,agents,created_by,created_at,updated_at,finalized_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      ("2026年担保业务研究报告", "{}", "已定稿", "已定稿", "一、总体\n业务稳健。", db.jdump({"red": 0, "items": [], "register": []}), "{}", "{}", "{}", "张三", db.now(), db.now(), db.now()))
    draft = db.execute("INSERT INTO m3_tasks(title,status,created_by,created_at) VALUES(?,?,?,?)", ("草稿报告不得外发", "待人工修改（有标红）", "张三", db.now()))
    r = say(fake, "ou_zhang", "报告")
    assert f"#{done}" in r and "草稿报告" not in r
    r = say(fake, "ou_zhang", f"报告 #{done}")
    assert "业务稳健" in r and "数据来源：M3" in r
    assert "没有找到已定稿" in say(fake, "ou_zhang", f"报告 #{draft}")
    n = len(fake.files)
    r = say(fake, "ou_zhang", f"报告文件 #{done}")
    assert len(fake.files) == n + 1 and fake.files[-1][0].endswith(".docx") and "[文件:" in r and "水印" in r
    from docx import Document
    import io
    assert "内部资料 注意保密" in Document(io.BytesIO(fake.files[-1][2])).sections[0].header.paragraphs[0].text
    assert db.q1("SELECT 1 x FROM audit_log WHERE action='导出 Word' AND detail LIKE '%飞书%'")
    assert "⚠" in say(fake, "ou_zhang", f"报告文件 #{draft}")
    world["done"] = done


def test_audit_has_bound_user(world, fake):
    say(fake, "ou_zhang", "帮助")
    a = db.q1("SELECT * FROM audit_log WHERE module='飞书' AND action LIKE '飞书消息%' AND username='fs_zhang' ORDER BY id DESC")
    assert a and "飞书" in a["detail"]
    e = db.q1("SELECT * FROM feishu_events WHERE username='fs_zhang' ORDER BY id DESC")
    assert e["status"] == "已处理" and e["intent"] == "help"


def test_dedup_rate_limit_group(world, fake, adm):
    payload = ev("ou_zhang", "帮助", event_id="ev_dup_1")
    S.dispatch(payload, sync=True)
    n = len(fake.sent)
    S.dispatch(payload, sync=True)  # 飞书重试：同一 event_id
    p2 = ev("ou_zhang", "帮助"); p2["event"]["message"]["message_id"] = payload["event"]["message"]["message_id"]
    S.dispatch(p2, sync=True)       # 同一 message_id
    assert len(fake.sent) == n
    # 限流
    _cfg(adm, rate_per_min=3)
    S._hits.clear()
    out = [say(fake, "ou_zhang", "帮助") for _ in range(5)]
    assert sum("太频繁" in o for o in out) == 2 and db.q1("SELECT 1 x FROM feishu_events WHERE status='限流'")
    _cfg(adm, rate_per_min=50)
    S._hits.clear()
    # 群聊：未 @ 不回复且不记录；@ 后回答一般问题；内部数据功能引导私聊
    cfg = C.get(); cfg["bot_open_id"] = BOT; db.set_setting(C.CFG_KEY, cfg)
    n_ev = db.q1("SELECT COUNT(*) n FROM feishu_events")["n"]
    assert say(fake, "ou_zhang", "额度", chat="group") == ""
    assert db.q1("SELECT COUNT(*) n FROM feishu_events")["n"] == n_ev
    r = say(fake, "ou_zhang", "今日简报", chat="group", mention=True)
    assert "央行降准" in r and fake.sent[-1][0] == "reply"
    r = say(fake, "ou_zhang", "额度", chat="group", mention=True)
    assert "私聊" in r and "万元" not in r
    # @ 了别人（不是机器人）不响应
    p = ev("ou_zhang", "额度", chat="group", mention=True); p["event"]["message"]["mentions"][0]["id"]["open_id"] = "ou_someone_else"
    n = len(fake.sent); S.dispatch(p, sync=True); assert len(fake.sent) == n
    # 非文字消息
    assert "只支持文字" in say(fake, "ou_zhang", "x", mtype="image")


def test_mail_via_bot(world, fake, adm, monkeypatch):
    sent = []

    class SMTP:
        def __init__(self, *a, **k): pass
        def ehlo(self): pass
        def starttls(self, **k): pass
        def login(self, *a): pass
        def send_message(self, m): sent.append(m)
        def quit(self): pass
    monkeypatch.setattr(smtplib, "SMTP_SSL", SMTP)
    r = adm.put("/api/admin/mail/config", json={"host": "smtp.example.com", "port": 465, "security": "ssl", "username": "u", "password": "p", "sender_addr": "bot@example.com",
                                                 "allowed_domains": ["example.com"], "enabled": True, "hourly_limit": 100})
    assert r.status_code == 200
    done = world["done"]
    r = say(fake, "ou_zhang", f"发邮箱 报告 #{done} lisi@example.com")
    assert "尚未发送" in r and "lisi@example.com" in r and "确认发送" in r and not sent
    assert "⚠" in say(fake, "ou_zhang", f"发邮箱 报告 #{done} out@evil.com") and not sent  # 域名白名单
    say(fake, "ou_zhang", f"发邮箱 报告 #{done} lisi@example.com")
    r = say(fake, "ou_zhang", "确认发送")
    assert "邮件已发送" in r and len(sent) == 1 and sent[0]["To"] == "lisi@example.com"
    m = db.q1("SELECT * FROM mail_log ORDER BY id DESC LIMIT 1")
    assert m["channel"] == "飞书" and m["username"] == "fs_zhang" and m["status"] == "已发送"
    # 取消
    say(fake, "ou_zhang", f"发邮箱 报告 #{done} lisi@example.com")
    assert "已取消" in say(fake, "ou_zhang", "取消") and len(sent) == 1
    assert "没明白" not in say(fake, "ou_zhang", "确认发送")  # 状态已清除，不会误发
    assert len(sent) == 1
    # 敏感报告：必须回复「确认涉密发送」
    sens = db.execute("INSERT INTO m3_tasks(title,requirements,status,stage,final,final_check,confirmations,tokens,agents,created_by,created_at,updated_at,finalized_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      ("含M5数据的报告", db.jdump({"use_m5": True}), "已定稿", "已定稿", "x。", db.jdump({"red": 0, "items": [], "register": []}), "{}", "{}", "{}", "张三", db.now(), db.now(), db.now()))
    r = say(fake, "ou_zhang", f"发邮箱 报告 #{sens} lisi@example.com")
    assert "确认涉密发送" in r
    assert "确认涉密发送" in say(fake, "ou_zhang", "确认") and len(sent) == 1
    assert "邮件已发送" in say(fake, "ou_zhang", "确认涉密发送") and len(sent) == 2
    assert db.q1("SELECT sensitive_confirmed FROM mail_log ORDER BY id DESC LIMIT 1")["sensitive_confirmed"] == 1


def test_ppt_multi_turn(world, fake):
    r = say(fake, "ou_zhang", "帮我生成PPT")
    assert "1." in r and "研究报告配套演示稿" in r
    assert "序号" in say(fake, "ou_zhang", "abc")
    r = say(fake, "ou_zhang", "4")
    assert "数据来源" in r and f"#{world['done']}" in r
    r = say(fake, "ou_zhang", "确认")
    assert "PPT 任务 #" in r and "需求单" in r and "确认" in r
    r = say(fake, "ou_zhang", "确认")
    assert "已生成大纲" in r and "确认大纲" in r
    n = len(fake.files)
    r = say(fake, "ou_zhang", "确认大纲")
    assert len(fake.files) == n + 1 and fake.files[-1][0].endswith(".pptx") and "数字溯源检查" in r
    from pptx import Presentation
    import io
    p = Presentation(io.BytesIO(fake.files[-1][2]))
    assert any("内部资料 注意保密" in sh.text_frame.text for sl in p.slides for sh in sl.shapes if sh.has_text_frame)
    assert db.q1("SELECT 1 x FROM ppt_jobs WHERE created_by='张三' AND status='已生成'")
    # 中途取消
    say(fake, "ou_zhang", "生成PPT"); assert "已取消" in say(fake, "ou_zhang", "取消")
    # 获取已生成 PPT
    r = say(fake, "ou_zhang", "PPT文件")
    assert "PPT 已发送" in r


def test_llm_router_only_sees_question_text(world, fake, monkeypatch):
    seen = []

    def fake_run(slot, messages, user, **kw):
        seen.append((slot, messages))
        return {"content": '{"intent":"alerts"}', "ms": 1, "model": "m", "tokens": 1}
    monkeypatch.setattr(S.L, "run", fake_run)
    r = say(fake, "ou_zhang", "最近有哪些让人担心的情况呀")
    assert "预警" in r and seen and seen[0][0] == "feishu.router"
    blob = json.dumps(seen[0][1], ensure_ascii=False)
    assert "最近有哪些让人担心的情况呀" in blob and world["mine"] not in blob and "万元" not in blob  # 只有提问文字，没有业务数据
    monkeypatch.setattr(S.L, "run", lambda *a, **k: {"content": '{"intent":"drop_database"}', "ms": 1, "model": "m", "tokens": 1})
    assert "没太明白" in say(fake, "ou_zhang", "最近有哪些让人担心的情况呀")  # 非白名单意图被拒
    monkeypatch.setattr(S.L, "run", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no model")))
    assert "没太明白" in say(fake, "ou_zhang", "今天天气不错呀朋友")


def test_http_callback(adm, fake):
    from app.main import app
    c = TestClient(app)
    ek = "my-encrypt-key"
    _cfg(adm, mode="http", encrypt_key=ek)
    assert c.post("/api/feishu/events", json={"type": "url_verification", "token": "wrong", "challenge": "c1"}).status_code in (400, 403)
    # 加密 + 签名
    def enc(obj):
        from cryptography.hazmat.primitives import padding
        from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
        key = hashlib.sha256(ek.encode()).digest(); iv = os.urandom(16)
        pd = padding.PKCS7(128).padder(); data = pd.update(json.dumps(obj).encode()) + pd.finalize()
        e = Cipher(algorithms.AES(key), modes.CBC(iv)).encryptor()
        return base64.b64encode(iv + e.update(data) + e.finalize()).decode()
    def post(obj, sign=True, bad_sig=False):
        raw = json.dumps({"encrypt": enc(obj)})
        h = {"content-type": "application/json"}
        if sign:
            ts, nonce = "1700000000", "n1"
            h.update({"x-lark-request-timestamp": ts, "x-lark-request-nonce": nonce,
                      "x-lark-signature": hashlib.sha256((ts + nonce + ek + raw).encode()).hexdigest() if not bad_sig else "0" * 64})
        return c.post("/api/feishu/events", content=raw, headers=h)
    r = post({"type": "url_verification", "token": "vt-token-123", "challenge": "abc123"})
    assert r.status_code == 200 and r.json() == {"challenge": "abc123"}
    assert post({"type": "url_verification", "token": "vt-token-123", "challenge": "x"}, bad_sig=True).status_code == 403
    assert post({"type": "url_verification", "token": "nope", "challenge": "x"}).status_code == 403
    n = len(fake.sent)
    p = ev("ou_unbound_http", "额度")
    import app.feishu.service as SS
    orig = SS.POOL.submit
    SS.POOL.submit = lambda fn, *a: fn(*a)  # 回调内同步执行，便于断言
    try:
        assert post(p).json() == {"code": 0}
    finally:
        SS.POOL.submit = orig
    assert len(fake.sent) == n + 1 and "绑定" in fake.sent[-1][2]
    assert db.q1("SELECT source FROM feishu_events WHERE open_id='ou_unbound_http'")["source"] == "HTTP回调"
    bad = ev("ou_zhang", "额度"); bad["header"]["token"] = "forged"
    assert post(bad).status_code == 403
    _cfg(adm, mode="ws", encrypt_key="")
    assert c.post("/api/feishu/events", json={}).status_code == 404  # 长连接模式下回调关闭


def test_sdk_adapter_roundtrip():
    from app.feishu import ws_runner
    got = []
    h = ws_runner.build_handler({"encrypt_key_enc": "", "verification_token_enc": ""}, dispatch=lambda p, source=None: got.append((p, source)))
    payload = ev("ou_sdk", "额度", chat="group", mention=True)
    feed = getattr(h, "_do_without_validation", None) or h.do_without_validation
    feed(json.dumps(payload).encode())
    assert got and got[0][1] == "长连接"
    p = got[0][0]
    assert p["header"]["event_type"] == "im.message.receive_v1" and p["event"]["message"]["mentions"][0]["id"]["open_id"] == BOT
    assert p["event"]["sender"]["sender_id"]["open_id"] == "ou_sdk"


def test_feishu_api_client_with_mocked_http(monkeypatch):
    calls = []

    class R:
        def __init__(self, j, code=200): self._j, self.status_code, self.text = j, code, json.dumps(j)
        def json(self): return self._j

    def post(url, **kw):
        calls.append(("POST", url, kw)); 
        if url.endswith("tenant_access_token/internal"):
            assert kw["json"]["app_id"] == "cli_test123456" and kw["json"]["app_secret"] == "secret-value-XYZ-123"
            return R({"code": 0, "tenant_access_token": "t-abc", "expire": 7200})
        if url.endswith("/im/v1/files"):
            return R({"code": 0, "data": {"file_key": "fk1"}})
        return R({"code": 0, "data": {}})

    def get(url, **kw):
        calls.append(("GET", url, kw)); return R({"code": 0, "bot": {"open_id": "ou_b", "app_name": "助手"}})
    monkeypatch.setattr(FA.httpx, "post", post); monkeypatch.setattr(FA.httpx, "get", get)
    api = FA.FeishuAPI(C.get())
    assert api.bot_info() == {"open_id": "ou_b", "name": "助手"}
    assert api.upload_file("a.docx", b"xx") == "fk1"
    api.send_text("oc_1", "你好"); api.send_file("oc_1", "fk1"); api.reply_text("om_1", "r")
    assert all(k[2]["headers"]["Authorization"] == "Bearer t-abc" for k in calls if "tenant_access_token" not in k[1])
    assert sum("tenant_access_token" in k[1] for k in calls) == 1  # token 缓存
    # 飞书返回错误码 → FeishuError
    monkeypatch.setattr(FA.httpx, "get", lambda url, **kw: R({"code": 99991663, "msg": "invalid token"}, 400))
    with pytest.raises(FA.FeishuError):
        FA.FeishuAPI(C.get()).bot_info()


def test_admin_test_connection_events_and_mapping(adm, world, fake):
    r = adm.post("/api/admin/feishu/test").json()
    assert r["ok"] and r["bot_name"] == "普惠助手" and C.get()["bot_open_id"] == BOT
    fake.token = lambda: (_ for _ in ()).throw(FA.FeishuError("app secret invalid (code 10014)", 10014))
    r = adm.post("/api/admin/feishu/test").json()
    assert not r["ok"] and "10014" in r["error"] and "secret-value" not in json.dumps(r)
    evs = adm.get("/api/admin/feishu/events").json()
    assert evs and {"open_id", "text", "intent", "status", "reply"} <= set(evs[0])
    assert adm.post("/api/admin/feishu/bindings", json={"open_id": "bad", "username": "fs_zhang"}).status_code == 400
    assert adm.post("/api/admin/feishu/bindings", json={"open_id": "ou_mapped_01", "username": "fs_zhang"}).status_code == 200
    assert db.q1("SELECT username FROM feishu_bindings WHERE open_id='ou_mapped_01'")["username"] == "fs_zhang"
    assert not db.q1("SELECT 1 x FROM feishu_bindings WHERE open_id='ou_zhang'")  # 1 账号 1 飞书用户
    b = adm.get("/api/admin/feishu/bindings").json()
    assert any(x["open_id"] == "ou_mapped_01" for x in b["bindings"])
    assert adm.delete("/api/admin/feishu/bindings/ou_mapped_01").json()["removed"] == 1
    assert "绑定" in say(fake, "ou_mapped_01", "额度")


def test_disabled_account_and_must_change(world, fake):
    I._set("ou_dis", "fs_inst", "admin")
    db.execute("UPDATE users SET enabled=0 WHERE username='fs_inst'")
    assert "停用" in say(fake, "ou_dis", "额度")
    db.execute("UPDATE users SET enabled=1 WHERE username='fs_inst'")
    assert "额度" in say(fake, "ou_dis", "额度") or "万元" in say(fake, "ou_dis", "额度")
    db.execute("UPDATE users SET must_change=1 WHERE username='fs_inst'")
    assert "修改初始密码" in say(fake, "ou_dis", "额度")
    db.execute("UPDATE users SET must_change=0 WHERE username='fs_inst'")


def test_bot_disabled_ignores_events(adm, fake):
    _cfg(adm, enabled=False)
    S.dispatch(ev("ou_zhang", "帮助"), sync=True)
    assert not fake.sent
    _cfg(adm, enabled=True)
