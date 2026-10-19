"""研究报告/简报/PPT 发送到邮箱：配置、域名白名单、定稿/标红/敏感确认、限流、记录与归档（SMTP 全部 mock）。"""
import io
import smtplib
from email import message_from_bytes
from urllib.parse import quote

import pytest
from fastapi.testclient import TestClient

from app import admin_auth, db

A = {"X-User": quote("张三"), "X-Role": "internal"}
INST = {"X-User": quote("机构"), "X-Role": "institution"}
SENT = []


class FakeSMTP:
    fail = False

    def __init__(self, host, port, timeout=None, context=None):
        self.host, self.port, self.tls, self.login_args = host, port, False, None

    def ehlo(self): pass
    def starttls(self, context=None): self.tls = True
    def login(self, u, p): self.login_args = (u, p)

    def send_message(self, msg):
        if FakeSMTP.fail:
            raise smtplib.SMTPAuthenticationError(535, b"auth failed")
        SENT.append((self.host, self.port, self.login_args, msg))

    def quit(self): pass


pytestmark = pytest.mark.usefixtures("isolated_db")


@pytest.fixture(scope="module")
def adm():
    from app.main import app
    c = TestClient(app)
    admin_auth.reset_password("admin", "Mail-Admin-2026x")
    assert c.post("/api/admin/login", json={"username": "admin", "password": "Mail-Admin-2026x"}).status_code == 200
    assert c.post("/api/admin/password", json={"old": "Mail-Admin-2026x", "new": "Mail-Admin-2026y"}).status_code == 200
    return c


@pytest.fixture(autouse=True)
def fake_smtp(monkeypatch):
    SENT.clear()
    FakeSMTP.fail = False
    monkeypatch.setattr(smtplib, "SMTP_SSL", FakeSMTP)
    monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)


def _task(status="已定稿", red=0, use_m5=False, title="2026年普惠融资担保研究"):
    chk = {"red": red, "risky": 0, "verified": 0, "items": [], "register": [], "llm_warnings": []}
    return db.execute("INSERT INTO m3_tasks(title,requirements,status,stage,final,final_check,confirmations,tokens,agents,created_by,created_at,updated_at,finalized_at) "
                      "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      (title, db.jdump({"use_m5": use_m5}), status, status, "一、总体情况\n业务稳步增长。", db.jdump(chk), "{}", "{}", "{}", "张三", db.now(), db.now(),
                       db.now() if status in ("已定稿", "已总结") else None))


def _cfg(adm, **kw):
    body = {"host": "smtp.example.com", "port": 465, "security": "ssl", "username": "bot@example.com", "password": "smtp-secret-pw-1",
            "sender_name": "普惠业务工作台", "sender_addr": "bot@example.com", "allowed_domains": ["example.com"], "enabled": True, "hourly_limit": 100, **kw}
    r = adm.put("/api/admin/mail/config", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def test_config_never_returns_password_and_blank_keeps(adm):
    c = _cfg(adm)
    assert "password" not in str(c).lower().replace("has_password", "") and c["has_password"] is True
    assert "smtp-secret-pw-1" not in adm.get("/api/admin/mail/config").text
    assert "smtp-secret-pw-1" not in (db.q1("SELECT value FROM settings WHERE key='smtp_config'")["value"])  # 加密存储
    c2 = _cfg(adm, password="")  # 留空不修改
    assert c2["has_password"] is True
    assert adm.put("/api/admin/mail/config", json={"host": "smtp.example.com", "port": 465, "security": "tls?"}).status_code == 400
    assert adm.put("/api/admin/mail/config", json={"host": "smtp.example.com", "port": 465, "allowed_domains": ["bad domain"]}).status_code == 400
    r = client_no_admin().get("/api/admin/mail/config")
    assert r.status_code == 403


def client_no_admin():
    from app.main import app
    return TestClient(app)


def test_test_send_uses_encrypted_password(adm):
    _cfg(adm)
    r = adm.post("/api/admin/mail/test", json={"to": "me@example.com"})
    assert r.status_code == 200 and r.json()["ok"] is True
    host, port, login, msg = SENT[-1]
    assert (host, port) == ("smtp.example.com", 465) and login == ("bot@example.com", "smtp-secret-pw-1")
    FakeSMTP.fail = True
    r = adm.post("/api/admin/mail/test", json={"to": "me@example.com"}).json()
    assert r["ok"] is False and "smtp-secret-pw-1" not in str(r)


def test_send_report_full_flow_and_history(client, adm):
    _cfg(adm)
    t = _task()
    p = client.post("/api/mail/preview", json={"kind": "m3_report", "ref": t, "to": "a@example.com; b@example.com"}, headers=A).json()
    assert p["recipients"] == ["a@example.com", "b@example.com"] and p["attachments"][0]["name"].endswith(".docx") and not p["sensitive"]
    # 未确认不得发送
    body = {"kind": "m3_report", "ref": t, "to": "a@example.com", "subject": "请审阅", "message": "请查收，周五前反馈。"}
    assert client.post("/api/mail/send", json=body, headers=A).status_code == 400
    r = client.post("/api/mail/send", json={**body, "confirm": True}, headers=A)
    assert r.status_code == 200, r.text
    msg = SENT[-1][3]
    assert msg["To"] == "a@example.com" and msg["Subject"] == "请审阅" and "内部资料" in msg["X-Confidential"]
    raw = message_from_bytes(msg.as_bytes())
    parts = list(raw.walk())
    text = next(p for p in parts if p.get_content_type() == "text/plain").get_payload(decode=True).decode()
    assert "内部资料 注意保密" in text and "张三" in text and "请查收" in text and "定稿时间" in text
    att = next(p for p in parts if (p.get_filename() or "").endswith(".docx"))
    docx_bytes = att.get_payload(decode=True)
    from docx import Document
    d = Document(io.BytesIO(docx_bytes))
    assert "内部资料 注意保密" in d.sections[0].header.paragraphs[0].text  # 附件带水印
    log = client.get("/api/mail/log", headers=A).json()
    assert log[0]["status"] == "已发送" and log[0]["to_addrs"] == "a@example.com" and log[0]["doc_version"].startswith("定稿时间")
    assert db.q1("SELECT 1 x FROM audit_log WHERE action='发送邮件' AND target=?", (f"m3_report#{t}",))
    assert db.q1("SELECT 1 x FROM export_log WHERE action='导出邮件附件' AND filename LIKE '%.docx'")  # 附件计入导出记录
    # 归档/删除与现有记录功能一致
    mid = log[0]["id"]
    assert client.post("/api/records/mail/archive", json={"ids": [mid]}, headers=A).status_code == 200
    assert mid not in [x["id"] for x in client.get("/api/mail/log", headers=A).json()]
    assert mid in [x["id"] for x in client.get("/api/mail/log?view=archived", headers=A).json()]
    assert client.post("/api/records/mail/delete", json={"ids": [mid]}, headers=A).status_code == 200


def test_blocking_rules(client, adm):
    _cfg(adm)
    ok = {"kind": "m3_report", "to": "a@example.com", "confirm": True}
    send = lambda t, **kw: client.post("/api/mail/send", json={**ok, "ref": t, **kw}, headers=A)  # noqa: E731
    n0 = len(SENT)
    assert send(_task(status="待人工定稿")).status_code == 409           # 未定稿
    assert send(_task(red=2)).status_code == 409                          # 标红未决项
    assert send(_task(), to="x@evil.com").status_code == 403              # 外部域名
    assert send(_task(), to="a@example.com, x@evil.com").status_code == 403
    assert send(_task(), to="bad-address").status_code == 400
    assert send(_task(), to="").status_code == 400
    assert send(999999).status_code == 404
    assert client.post("/api/mail/send", json={**ok, "ref": _task()}, headers=INST).status_code == 403  # 机构用户不可发
    s = _task(use_m5=True)
    assert send(s).status_code == 409                                      # 敏感须二次确认
    assert send(s, confirm_sensitive=True).status_code == 200
    assert db.q1("SELECT sensitive_confirmed FROM mail_log ORDER BY id DESC LIMIT 1")["sensitive_confirmed"] == 1
    assert len(SENT) == n0 + 1


def test_whitelist_must_be_set(client, adm):
    _cfg(adm, allowed_domains=[])
    r = client.post("/api/mail/send", json={"kind": "m3_report", "ref": _task(), "to": "a@example.com", "confirm": True}, headers=A)
    assert r.status_code == 403 and "白名单" in r.json()["detail"]
    _cfg(adm, allowed_domains=["*.corp.cn"])
    assert client.post("/api/mail/send", json={"kind": "m3_report", "ref": _task(), "to": "a@dept.corp.cn", "confirm": True}, headers=A).status_code == 200
    assert client.post("/api/mail/send", json={"kind": "m3_report", "ref": _task(), "to": "a@corp.cn", "confirm": True}, headers=A).status_code == 403  # *.x 不含根域
    _cfg(adm)


def test_failure_logged_and_rate_limit(client, adm):
    _cfg(adm, hourly_limit=3)
    FakeSMTP.fail = True
    t = _task()
    r = client.post("/api/mail/send", json={"kind": "m3_report", "ref": t, "to": "a@example.com", "confirm": True}, headers={"X-User": quote("赵六"), "X-Role": "internal"})
    assert r.status_code == 502 and "smtp-secret-pw-1" not in r.text
    FakeSMTP.fail = False
    h = {"X-User": quote("赵六"), "X-Role": "internal"}
    for _ in range(2):
        assert client.post("/api/mail/send", json={"kind": "m3_report", "ref": t, "to": "a@example.com", "confirm": True}, headers=h).status_code == 200
    r = client.post("/api/mail/send", json={"kind": "m3_report", "ref": t, "to": "a@example.com", "confirm": True}, headers=h)
    assert r.status_code == 429
    log = client.get("/api/mail/log", headers=h).json()
    assert [x["status"] for x in log[:3]] == ["已发送", "已发送", "失败"]
    assert all(x["user"] == "赵六" for x in log)  # 内部用户只能看到自己的记录
    _cfg(adm)


def test_m2_brief_and_ppt_file(client, adm, tmp_path):
    _cfg(adm)
    bid = db.execute("INSERT INTO m2_briefs(brief_date,version,status,items,markdown,stats,config,model,published_at) VALUES(?,?,?,?,?,?,?,?,?)",
                     ("2026-10-09", 1, "已发布", db.jdump({"items": [], "data_rows": []}), "# 今日行业 Brief · 2026-10-09\n\n## 一、今日必读\n（无）\n", "{}", "{}", "", db.now()))
    draft = db.execute("INSERT INTO m2_briefs(brief_date,version,status,items,markdown,stats,config,model) VALUES(?,?,?,?,?,?,?,?)",
                       ("2026-10-08", 1, "草稿", db.jdump({"items": [], "data_rows": []}), "# x\n", "{}", "{}", ""))
    ok = {"kind": "m2_brief", "to": "a@example.com", "confirm": True}
    assert client.post("/api/mail/send", json={**ok, "ref": bid}, headers=A).status_code == 200
    assert client.post("/api/mail/send", json={**ok, "ref": draft}, headers=A).status_code == 409
    # PPT 文件
    from pptx import Presentation
    pj = db.execute("INSERT INTO ppt_jobs(deck_type,title,scope,status,pack,created_by,created_at) VALUES(?,?,?,?,?,?,?)",
                    ("research", "配套演示稿", db.jdump({"task_id": 0}), "已生成", db.jdump({"sensitive": False}), "张三", db.now()))
    path = tmp_path / "v1.pptx"
    pr = Presentation(); pr.slides.add_slide(pr.slide_layouts[0]); pr.save(path)
    fid = db.execute("INSERT INTO ppt_files(job_id,version,filename,path,pages,thumbs,trace,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                     (pj, 1, "配套演示稿_v1.pptx", str(path), 1, "[]", db.jdump({"ok": True}), "张三", db.now()))
    assert client.post("/api/mail/send", json={"kind": "ppt_file", "ref": fid, "to": "a@example.com", "confirm": True}, headers=A).status_code == 200
    assert SENT[-1][3].get_body(preferencelist=("plain",)) is not None
    # 报告 + 配套 PPT：PPT 必须由该报告生成
    t = _task()
    assert client.post("/api/mail/send", json={"kind": "m3_report", "ref": t, "to": "a@example.com", "confirm": True, "opts": {"ppt_file_id": fid}}, headers=A).status_code == 400
    db.execute("UPDATE ppt_jobs SET scope=? WHERE id=?", (db.jdump({"task_id": t}), pj))
    r = client.post("/api/mail/send", json={"kind": "m3_report", "ref": t, "to": "a@example.com", "confirm": True, "opts": {"ppt_file_id": fid}}, headers=A)
    assert r.status_code == 200 and len(r.json()["attachments"]) == 2 and "PPT v1" in r.json()["version"]
