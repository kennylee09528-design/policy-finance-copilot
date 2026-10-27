"""M4 口径版本：formula_spec 与结果 meta 打戳，不改变测算数值。"""
import hashlib
import io

import openpyxl
import pytest

from app.m4.calc import compute
from app.m4.service import recompute
from app.m4.spec import CALC_PATH, calc_module_hash, load_spec, validate_spec


@pytest.fixture(scope="module")
def imported(client):
    r = client.post("/api/m4/import-sample")
    assert r.status_code == 200, r.text
    return r.json()


def test_spec_file_matches_calc_contract():
    spec = validate_spec(load_spec())
    assert spec["version"]
    by_id = {c["id"]: c for c in spec["columns"]}
    assert by_id["U"]["missing_policy"] == "null_and_alert"
    assert by_id["X"]["missing_policy"] == "null_and_alert"
    assert "告警" in spec["missing_coefficient_rule"] and "X" in spec["missing_coefficient_rule"]
    for col in spec["columns"]:
        assert col["expr_human"] and col["source"] and col["inputs"]


def test_recompute_stamps_meta_without_changing_numbers(client, imported):
    sid = imported["scheme_id"]
    s = client.get(f"/api/m4/schemes/{sid}").json()
    before = [(r["name"], r["W"], r["U"], r["X"], r["AC"], r["final"]) for r in s["result"]["rows"]]
    recompute(sid)
    s2 = client.get(f"/api/m4/schemes/{sid}").json()
    after = [(r["name"], r["W"], r["U"], r["X"], r["AC"], r["final"]) for r in s2["result"]["rows"]]
    assert after == before
    meta = s2["result"]["meta"]
    assert meta["formula_spec_version"] == load_spec()["version"]
    assert meta["param_id"] == s["param_id"] == s2["result"]["param"]["id"]
    assert meta["param_version"] == s2["result"]["param"]["version"]
    assert meta["dataset_id"] == s["dataset_id"]
    assert meta["calc_module_hash"] == calc_module_hash()
    assert meta["calc_module_hash"] == hashlib.sha256(CALC_PATH.read_bytes()).hexdigest()
    cur = client.get("/api/m4/schemes/current").json()
    assert cur["result"]["meta"]["formula_spec_version"] == meta["formula_spec_version"]
    assert cur["result"]["meta"]["dataset_id"] == cur["dataset_id"]


def test_missing_coefficient_still_alerts_and_skips_u_x():
    """缺系数：告警且不算 U/X。与口径说明 null_and_alert 一致，数值行为不改。"""
    row = {"name": "示例机构甲", "scale": 100000, "emp": None, "zxzn_raw": 90, "lj_raw": 1.2,
           "dn_raw": 1.1, "pj": 80, "zc_total": 0, "mintou_prev": 0, "fixed": False}
    params = {"weights": {"L": 0.2, "N": 0.2, "P": 0.2, "R": 0.2, "T": 0.2},
              "constants": {"emp_avg": 1000, "lj_avg": 0.02, "dn_avg": 0.02, "floor": 0.5, "zxzn_threshold": 0.9}}
    out = compute([row], params)
    got = out["rows"][0]
    assert got["U"] is None and got["X"] is None
    assert any(a["type"] == "数据缺失" and "示例机构甲" in a["msg"] for a in out["alerts"])


def test_formula_spec_export(client, imported):
    sid = imported["scheme_id"]
    spec = load_spec()
    r = client.get(f"/api/m4/schemes/{sid}/formula-spec.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "口径说明" in wb.sheetnames
    text = "\n".join(str(c.value) for row in wb["口径说明"].iter_rows() for c in row if c.value)
    assert spec["version"] in text
    assert "缺系数" in text or "不算" in text
    assert "null_and_alert" in text
    full = client.get(f"/api/m4/schemes/{sid}/export.xlsx")
    wb2 = openpyxl.load_workbook(io.BytesIO(full.content))
    assert "口径说明" in wb2.sheetnames
    api = client.get("/api/m4/formula-spec").json()
    assert api["version"] == spec["version"] and api["calc_module_hash"] == calc_module_hash()
    assert client.get("/api/m4/formula-spec", headers={"X-Role": "institution"}).status_code == 403
