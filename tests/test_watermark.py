"""导出水印：Word 页眉/页脚、Excel 打印页眉+导出信息表、PPT 每页小字、Markdown 文末；PDF 打印导出留痕。"""
import io
import zipfile
from urllib.parse import quote

from app import db, watermark as W

U = {"name": "张三", "username": "zhang_san", "role": "internal"}


def test_docx_xlsx_pptx_md():
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    d = Document(); d.add_paragraph("正文")
    b = io.BytesIO(); d.save(b)
    out = Document(io.BytesIO(W.docx_bytes(b.getvalue(), U)))
    hdr = " ".join(p.text for p in out.sections[0].header.paragraphs)
    assert "张三（zhang_san）" in hdr and "内部资料" in hdr and out.core_properties.category == "内部资料"
    wb = Workbook(); wb.active.append([1, 2]); b = io.BytesIO(); wb.save(b)
    wb2 = load_workbook(io.BytesIO(W.xlsx_bytes(b.getvalue(), U)))
    assert "导出信息" in wb2.sheetnames and "张三" in wb2.worksheets[0].oddHeader.right.text
    assert wb2["导出信息"]["B3"].value.startswith("张三")
    p = Presentation(); p.slides.add_slide(p.slide_layouts[6]); p.slides.add_slide(p.slide_layouts[6]); b = io.BytesIO(); p.save(b)
    p2 = Presentation(io.BytesIO(W.pptx_bytes(b.getvalue(), U)))
    assert all(any(s.name == "ph_watermark" and "张三" in s.text_frame.text for s in sl.shapes) for sl in p2.slides)
    assert "导出人：张三" in W.md_text("# 标题", U)


def test_routes_apply_watermark_and_pdf_log(client):
    h = {"X-User": quote("李四"), "X-Role": "internal"}
    if not db.q1("SELECT 1 x FROM m4_schemes WHERE is_current=1 AND deleted_at IS NULL"):
        client.post("/api/m4/import-sample")
    sid = db.q1("SELECT id FROM m4_schemes WHERE is_current=1 AND deleted_at IS NULL")["id"]
    r = client.get(f"/api/m4/schemes/{sid}/export.xlsx", headers=h)
    assert r.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        assert any("李四" in z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    r = client.post("/api/exports/log", json={"module": "M5", "target": 1, "filename": "看板.pdf"}, headers=h)
    assert r.status_code == 200 and r.json()["watermark"]["who"] == "李四"
    e = client.get("/api/exports", headers=h).json()[0]
    assert e["action"] == "导出 PDF（打印）" and e["filename"] == "看板.pdf"
    assert client.post("/api/exports/log", json={"module": "M4", "filename": "x"}, headers={"X-Role": "institution"}).status_code == 403
