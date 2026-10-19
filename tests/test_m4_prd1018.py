"""M4 自动化验收：导入项目内虚构示例《授信额度计算表（示例）》（scripts/gen_sample_data.py 生成），默认权重 0.2/0.2/0.1/0.1/0.4。

示例数据均为虚构。K 列基期常数仅「示例担保公司A」含 +50000（演示该口径），其余机构分母只有 I。
"""
import io
import zipfile

import openpyxl
import pytest

from app.config import M4_SAMPLE

PRD = {  # 机构: J, L, N, P, R, T, U, W, X, AC(公式计算值)
    "示例担保公司A": (24.7890, 0.9848, 0.8113, 0.5, 0.5571, 0.6766, 18.2340, 274301, 183478, 200000),
    "示例担保公司B": (38.6681, 1.1797, 0.9832, 0.5, 1.0264, 0.9248, 36.9337, 309795, 371643, 390000),
    "示例担保公司C": (38.2648, 0.9899, 0.8728, 0.8758, 0.5188, 0.9125, 33.5581, 325914, 337676, 360000),
    "示例担保公司D": (10.8420, 0.8356, 0.8939, 1.0672, 1.7573, 0.9360, 10.8717, 104963, 109395, 110000),
    "示例担保公司E": (22.9534, 0.9504, 0.8, 0.5, 1.6309, 0.8141, 20.4012, 202577, 205285, 210000),
}
CURRENT = {"示例担保公司A": (639000, 30000), "示例担保公司B": (335000, 0), "示例担保公司C": (397000, 0)}
FIXED = ["示例担保公司I", "示例担保公司J"]


@pytest.fixture(scope="module")
def cur(client):
    assert client.post("/api/m4/import-sample").status_code == 200
    return client.get("/api/m4/schemes/current").json()


def _ledger():
    ws = openpyxl.load_workbook(M4_SAMPLE, data_only=True)["计算表"]
    return {ws[f"B{r}"].value.strip(): {c: ws[f"{c}{r}"].value for c in ("C", "D", "X", "AB", "AC", "AI")}
            for r in range(2, ws.max_row + 1) if ws[f"B{r}"].value}


def test_prd_table(cur):
    rows = {r["name"]: r for r in cur["result"]["rows"]}
    assert cur["result"]["param"]["weights"] == {"L": 0.2, "N": 0.2, "P": 0.1, "R": 0.1, "T": 0.4}
    print("\n机构 | J L N P R T U | W | X | AC")
    for name, exp in PRD.items():
        r = rows[name]
        got = (r["J"], r["L"], r["N"], r["P"], r["R"], r["T"], r["U"])
        for g, e in zip(got, exp[:7]):
            assert round(g, 4) == pytest.approx(e, abs=1e-9), (name, got)
        assert round(r["W"]) == exp[7] and r["X"] == exp[8] and r["AC"] == exp[9], (name, r["W"], r["X"], r["AC"])
        print(name, *[f"{g:.4f}" for g in got], round(r["W"]), r["X"], r["AC"])


def test_ac_vs_ledger(cur):
    led = _ledger()
    rows = {r["name"]: r for r in cur["result"]["rows"]}
    same = [n for n in led if rows[n]["AC"] == led[n]["AC"]]
    diff = {n: (rows[n]["AC"], led[n]["AC"]) for n in led if n not in same}
    print(f"\nAC 计算值与台账一致 {len(same)}/{len(led)}；差异：{diff}")
    assert len(led) == 10 and len(same) == 10 and not diff
    for n, r in rows.items():  # AI = AC + D
        assert r["X"] == led[n]["X"]
        if r["AC"] == led[n]["AC"]:
            assert r["AI"] == pytest.approx(led[n]["AI"])


def test_current_values_and_fixed(cur):
    rows = {r["name"]: r for r in cur["result"]["rows"]}
    for n, (c, d) in CURRENT.items():
        assert (rows[n]["cur_regular"], rows[n]["cur_mintou"]) == (c, d)
    led = _ledger()
    for n, r in rows.items():
        assert r["cur_regular"] == (led[n]["C"] or 0) and r["cur_mintou"] == (led[n]["D"] or 0)
    fixed = [r["name"] for r in rows.values() if r["fixed"]]
    assert len(fixed) == 2 and all(any(k in n for k in FIXED) for n in fixed)
    assert rows["示例担保公司A"]["Y"] == pytest.approx(183478 - 35632.376)
    assert rows["示例担保公司D"]["Y"] is None  # 无工行预计


def test_constants(cur):
    c = cur["result"]["param"]["constants"]
    assert round(c["emp_avg"], 4) == 1083.6309 and round(c["lj_avg"], 6) == 0.023372 and round(c["dn_avg"], 6) == 0.028992
    assert c["base_consts"]["示例担保公司A"] == 50000.0


def test_current_adjust_log(client, cur):
    sid = cur["id"]
    n = "示例担保公司D"
    base = {"institution": n, "kind": "cur_regular", "delta": 20000}
    assert client.post(f"/api/m4/schemes/{sid}/cur-adjust", json={**base, "reason": " "}).status_code == 400
    assert client.post(f"/api/m4/schemes/{sid}/cur-adjust", json={**base, "reason": "上会确认"}, headers={"X-User": "lisi"}).status_code == 200
    assert client.post(f"/api/m4/schemes/{sid}/cur-adjust", json={"institution": n, "kind": "cur_mintou", "delta": -10000, "reason": "按请示调减"}).status_code == 200
    r = next(x for x in client.get(f"/api/m4/schemes/{sid}").json()["result"]["rows"] if x["name"] == n)
    assert (r["cur_regular"], r["cur_mintou"], r["AC"]) == (433000, 70000, 110000)  # AC 不受影响
    log = client.get(f"/api/m4/schemes/{sid}/cur-log").json()
    assert log[0]["before"] == 80000 and log[0]["after"] == 70000 and log[0]["delta"] == -10000 and log[0]["reason"] == "按请示调减"
    assert log[1]["operator"] == "lisi" and log[1]["before"] == 413000 and log[1]["after"] == 433000 and log[1]["time"]
    x = client.get(f"/api/m4/schemes/{sid}/cur-log.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    cp = client.post(f"/api/m4/schemes/{sid}/copy", json={}).json()["scheme_id"]  # 复制方案保留全部增减
    r2 = next(x for x in client.get(f"/api/m4/schemes/{cp}").json()["result"]["rows"] if x["name"] == n)
    assert (r2["cur_regular"], r2["cur_mintou"]) == (433000, 70000)


def test_weights_preview_tolerance(client, cur):
    sid = cur["id"]
    ok = client.post(f"/api/m4/schemes/{sid}/preview", json={"weights": {"L": 0.21, "N": 0.2, "P": 0.1, "R": 0.1, "T": 0.4}}).json()
    assert not ok["totals"]["weights_normalized"]
    bad = client.post(f"/api/m4/schemes/{sid}/preview", json={"weights": {"L": 0.4, "N": 0.4, "P": 0.2, "R": 0.2, "T": 0.8}}).json()
    assert bad["totals"]["weights_normalized"] and bad["alerts"][0]["type"] == "权重和≠1"
    dflt = {r["name"]: r for r in client.post(f"/api/m4/schemes/{sid}/preview", json={}).json()["rows"]}
    assert all(r["AC"] == dflt[r["name"]]["AC"] for r in bad["rows"])  # ×2 归一后与默认一致
    assert dflt["示例担保公司A"]["AC"] == 200000


def _rename_sheet(data, old, new):
    zin, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for it in zin.infolist():
            b = zin.read(it.filename)
            if it.filename == "xl/workbook.xml":
                b = b.replace(f'name="{old}"'.encode(), f'name="{new}"'.encode())
            z.writestr(it, b)
    return out.getvalue()


def test_import_checks(client):
    data = M4_SAMPLE.read_bytes()
    r = client.post("/api/m4/import", files={"file": ("缺表.xlsx", _rename_sheet(data, "就业人数", "其他"))}, data={"prefill": "false"})
    assert r.status_code == 400 and "就业人数" in r.json()["detail"]
    # 无「计算表」：按 PRD 列映射兜底导入（关键词定位，年份可变）
    d2 = _rename_sheet(_rename_sheet(data, "计算表", "台账备份"), "2025年规模", "2026年规模")
    r = client.post("/api/m4/import", files={"file": ("无计算表.xlsx", d2)}, data={"prefill": "false"})
    assert r.status_code == 200, r.text
    ds = client.get(f"/api/m4/datasets/{r.json()['dataset_id']}").json()["payload"]
    assert ds["mode"] == "direct" and ds["sheet_map"]["2025年规模"] == "2026年规模"
    rows = {x["name"]: x for x in ds["rows"]}
    assert rows["示例担保公司A"]["apply_regular_tech"] == 639000 and rows["示例担保公司A"]["scale"] == pytest.approx(247889.783751)
