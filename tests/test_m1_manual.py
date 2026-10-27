"""同事补录：链接走现有采集链路，PDF/Word 抽正文；去重与权限。"""
import io

import pytest
from fastapi import HTTPException

from app import db
from app.config import DATA_DIR
from app.m1 import collector

TITLE = "关于示例补录甲号融资担保奖补的正式通知文件"
TITLE_B = "关于示例补录乙号融资担保奖补的正式通知文件"
PUB = "示例发布机构（测试）"
DOC = "示测〔2026〕8号"
H_IN = {"X-Role": "internal", "X-User": "zhangsan"}
H_ADMIN = {"X-Role": "admin", "X-User": "admin"}
H_ORG = {"X-Role": "institution", "X-User": "org-user"}


@pytest.fixture(autouse=True)
def _cleanup():
    yield
    db.execute("DELETE FROM m1_items WHERE provider='manual' OR url LIKE 'https://manual.example/%' OR url LIKE 'https://czt.gd.gov.cn/zcfg/manual-%'")


def _docx(*paragraphs, table=None):
    from docx import Document
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=1, cols=1)
        t.rows[0].cells[0].text = table
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()


def _pdf(text: str) -> bytes:
    safe = "".join(ch if 32 <= ord(ch) < 127 and ch not in "()\\" else " " for ch in text)
    stream = f"BT /F1 24 Tf 72 700 Td ({safe}) Tj ET".encode()
    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    blobs = [f"{i} 0 obj\n".encode() + b + b"\nendobj\n" for i, b in enumerate(bodies, 1)]
    header = b"%PDF-1.4\n"
    offsets, pos = [], len(header)
    for b in blobs:
        offsets.append(pos)
        pos += len(b)
    xref = b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    trailer = f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{pos}\n%%EOF\n".encode()
    return header + b"".join(blobs) + xref + trailer


def _blank_pdf() -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=300, height=300)
    bio = io.BytesIO()
    w.write(bio)
    return bio.getvalue()


def _fields(**kw):
    base = {"title": TITLE, "publisher": PUB, "doc_no": DOC, "pub_date": "2026-03-01",
            "summary": "- 虚构要点：担保费率补贴不超过1%", "level": "A2", "suggested_level": "A2",
            "url": "", "allow_title_dup": "false"}
    base.update(kw)
    return base


def _upload(client, blob, name, fields, headers=None):
    return client.post("/api/m1/manual/file", files={"file": (name, blob, "application/octet-stream")}, data=fields, headers=headers or H_IN)


def test_url_ingest_reuses_pipeline_and_rejects_duplicate(client, monkeypatch):
    calls = []

    def fake_fetch(url):
        calls.append(url)
        return {"title": TITLE, "text": "虚构正文。担保费率补贴不超过1%。" * 4, "date": "2026-03-01", "kind": "html", "final_url": url}

    def fake_sum(iid):
        db.execute("UPDATE m1_items SET summary=? WHERE id=?", ("- 虚构要点", iid))
        return True

    monkeypatch.setattr("app.m1.fetch.fetch", fake_fetch)
    monkeypatch.setattr(collector, "_summarize", fake_sum)
    monkeypatch.setattr("app.search.service.search", lambda *a, **k: (_ for _ in ()).throw(AssertionError("不应联网搜索")))

    url = "https://czt.gd.gov.cn/zcfg/manual-1.html?utm_source=newsletter"
    r = client.post("/api/m1/manual/url", json={"url": url}, headers=H_IN)
    assert r.status_code == 200, r.text
    item = r.json()["item"]
    assert item["provider"] == "manual" and item["uploaded_by"] == "zhangsan"
    assert item["validity_status"] == "未知" and item["reviewed_at"] is None
    assert item["level"] == "A1" and item["publisher"] == "广东省财政厅" and item["level_manual"] == 0
    assert item["url"] == "https://czt.gd.gov.cn/zcfg/manual-1.html" and "utm_" not in item["url"]
    assert item["summary"] == "- 虚构要点" and item["fetch_status"] == "已抓取" and "补贴" in item["content"]
    assert calls == ["https://czt.gd.gov.cn/zcfg/manual-1.html"]
    listed = client.get("/api/m1/items", params={"kw": "示例补录甲号"}).json()["items"]
    assert item["id"] in [x["id"] for x in listed]
    q = client.get("/api/m1/review").json()["queue"]
    hit = next(x for x in q if x["id"] == item["id"])
    assert "同事补录，待人工复核" in hit["reasons"] and hit["uploaded_by"] == "zhangsan"

    dup = client.post("/api/m1/manual/url", json={"url": "https://CZT.gd.gov.cn/zcfg/manual-1.html?utm_source=again"}, headers=H_IN)
    assert dup.status_code == 409
    body = dup.json()["detail"]
    assert body["existing_id"] == item["id"] and body["kind"] == "url" and body["can_force"] is False
    assert "已有条目" in body["message"] and str(item["id"]) in body["message"]
    assert calls == ["https://czt.gd.gov.cn/zcfg/manual-1.html"]  # 重复链接不再抓取

    other = client.post("/api/m1/manual/url", json={"url": "https://czt.gd.gov.cn/zcfg/manual-2.html"}, headers=H_IN)
    assert other.status_code == 409 and other.json()["detail"]["kind"] == "title" and other.json()["detail"]["can_force"] is True
    forced = client.post("/api/m1/manual/url", json={"url": "https://czt.gd.gov.cn/zcfg/manual-2.html", "allow_title_dup": True}, headers=H_ADMIN)
    assert forced.status_code == 200, forced.text
    dup_item = forced.json()["item"]
    assert dup_item["duplicate_of"] == item["id"] and dup_item["uploaded_by"] == "admin"
    assert "疑似重复" in forced.json()["warning"]
    visible = [x["id"] for x in client.get("/api/m1/items", params={"kw": "示例补录甲号"}).json()["items"]]
    assert item["id"] in visible and dup_item["id"] not in visible
    shown = [x["id"] for x in client.get("/api/m1/items", params={"kw": "示例补录甲号", "show_dup": True}).json()["items"]]
    assert dup_item["id"] in shown


def test_manual_url_rejects_intranet_and_non_http(client):
    denied = client.post("/api/m1/manual/url", json={"url": "http://127.0.0.1/secret"}, headers=H_IN)
    assert denied.status_code == 400 and "内网" in denied.json()["detail"]
    assert client.post("/api/m1/manual/url", json={"url": "ftp://manual.example/a"}, headers=H_IN).status_code == 400
    assert client.post("/api/m1/manual/url", json={"url": "  "}, headers=H_IN).status_code == 400


def test_url_pdf_warns_and_summary_failure_keeps_item(client, monkeypatch):
    monkeypatch.setattr("app.m1.fetch.fetch", lambda url: {"title": "示例补录丙号通知文件名称足够长", "text": "", "date": None, "kind": "pdf", "final_url": url})
    r = client.post("/api/m1/manual/url", json={"url": "https://manual.example/policy.pdf"}, headers=H_IN)
    assert r.status_code == 200, r.text
    assert r.json()["item"]["fetch_status"] == "PDF（未解析正文）"
    assert "上传文件" in r.json()["warning"]
    monkeypatch.setattr("app.m1.fetch.fetch", lambda url: {"title": "示例补录丁号通知文件名称足够长", "text": "正文足够长，用于确认摘要失败仍入库。", "date": "2026-04-01", "kind": "html", "final_url": url})

    def boom(_iid):
        raise HTTPException(502, "模型调用失败：超时")

    monkeypatch.setattr(collector, "_summarize", boom)
    r = client.post("/api/m1/manual/url", json={"url": "https://manual.example/b.html"}, headers=H_IN)
    assert r.status_code == 200, r.text
    assert r.json()["item"]["id"] and "摘要未生成" in r.json()["warning"]


def test_institution_forbidden(client):
    blob = _docx("这是一段足够长的虚构政策正文，用于权限测试，不对应任何真实文件。")
    assert client.post("/api/m1/manual/url", json={"url": "https://manual.example/x"}, headers=H_ORG).status_code == 403
    assert client.post("/api/m1/manual/preview", files={"file": ("a.docx", blob, "application/octet-stream")}, headers=H_ORG).status_code == 403
    assert _upload(client, blob, "a.docx", _fields(), H_ORG).status_code == 403
    assert client.get("/api/m1/items", params={"kw": "权限测试"}).status_code == 200


def test_file_extract_clean_dedup_and_level_override(client, monkeypatch):
    monkeypatch.setattr("app.llm.service.run", lambda *a, **k: (_ for _ in ()).throw(AssertionError("字段已齐不应调用模型")))
    blob = _docx(
        "首页", "登录", "索引号：ABC-001",
        TITLE,
        "为落实示例政策，对担保费率不超过1%的机构给予奖补。本段为虚构测试正文，不对应真实文件。",
        table="表内字段：示例文号不计真实文件，仅供测试。",
    )
    r = _upload(client, blob, "../../etc/passwd.docx", _fields(level="A1", suggested_level="A2"))
    assert r.status_code == 200, r.text
    item = r.json()["item"]
    assert item["provider"] == "manual" and item["uploaded_by"] == "zhangsan" and item["validity_status"] == "未知"
    assert item["reviewed_at"] is None and item["level"] == "A1" and item["level_manual"] == 1
    assert item["doc_no"] == DOC and item["fetch_status"] == "已上传解析"
    assert "索引号" not in (item["content"] or "") and "首页" not in (item["content"] or "")
    assert "奖补" in item["content"] and "表内字段" in item["content"]
    assert item["source_filename"] == "passwd.docx"
    assert item["source_path"].startswith("m1_uploads/") and ".." not in item["source_path"]
    stored = DATA_DIR / item["source_path"]
    assert stored.is_file() and stored.parent == (DATA_DIR / "m1_uploads") and stored.read_bytes().startswith(b"PK")
    assert item["id"] in [x["id"] for x in client.get("/api/m1/review").json()["queue"]]

    same = _upload(client, blob, "b.docx", _fields(level="A1", suggested_level="A2", allow_title_dup="true"))
    assert same.status_code == 409
    assert same.json()["detail"]["kind"] == "identity" and same.json()["detail"]["can_force"] is False
    assert same.json()["detail"]["existing_id"] == item["id"]
    assert "完全一致" in same.json()["detail"]["message"]

    soft = _upload(client, blob, "c.docx", _fields(title=TITLE_B, doc_no="", level="A2", suggested_level="A2"))
    assert soft.status_code == 200, soft.text
    assert soft.json()["item"]["level_manual"] == 0
    again = _upload(client, blob, "d.docx", _fields(title=TITLE_B, publisher="另一机构", doc_no="示测〔2026〕9号"))
    assert again.status_code == 409 and again.json()["detail"]["kind"] == "title" and again.json()["detail"]["can_force"] is True
    forced = _upload(client, blob, "e.docx", _fields(title=TITLE_B, publisher="另一机构", doc_no="示测〔2026〕9号", allow_title_dup="true"))
    assert forced.status_code == 200 and forced.json()["item"]["duplicate_of"] == soft.json()["item"]["id"]
    assert forced.json()["item"]["id"] not in [x["id"] for x in client.get("/api/m1/items", params={"kw": "示例补录乙号"}).json()["items"]]

    bad_date = _upload(client, blob, "f.docx", _fields(title="示例补录日期无效的通知文件名称", doc_no="示测〔2026〕10号", pub_date="2026/13/01"))
    assert bad_date.status_code == 400 and "日期" in bad_date.json()["detail"]


def test_preview_llm_and_unreadable_files(client, monkeypatch):
    seen = {}

    def fake_run(slot, messages, user, **kw):
        seen["slot"] = slot
        seen["user"] = user.get("name")
        assert "虚构" in messages[0]["content"]
        return {"content": '{"title":"关于示例补录预填的通知文件","publisher":"示例发布机构（测试）","doc_no":"示测〔2026〕11号","pub_date":"2026年5月2日","summary":["担保费率补贴不超过1%"],"level":"A2"}'}

    monkeypatch.setattr("app.llm.service.run", fake_run)
    blob = _docx("首页", "这是预填测试用的虚构政策正文，提到担保费率补贴不超过1%，不对应真实文件。")
    before = db.q1("SELECT COUNT(*) n FROM m1_items WHERE provider='manual'")["n"]
    r = client.post("/api/m1/manual/preview", files={"file": ("预填.docx", blob, "application/octet-stream")}, headers=H_IN)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["suggestions"]["title"].startswith("关于示例补录预填")
    assert body["suggestions"]["pub_date"] == "2026-05-02" and body["suggestions"]["level"] == "A2"
    assert body["suggestions"]["summary"].startswith("- ")
    assert "首页" not in body["text"] and body["chars"] >= 20
    assert seen["slot"] == "m1.extract" and seen["user"] == "zhangsan"
    assert db.q1("SELECT COUNT(*) n FROM m1_items WHERE provider='manual'")["n"] == before

    monkeypatch.setattr("app.llm.service.run", lambda *a, **k: (_ for _ in ()).throw(HTTPException(400, "该环节尚未选择模型")))
    r = client.post("/api/m1/manual/preview", files={"file": ("预填.docx", blob, "application/octet-stream")}, headers=H_IN)
    assert r.status_code == 200 and "手工填写" in r.json()["llm_warning"]
    missing = _upload(client, blob, "预填.docx", _fields(title="", summary="", level="", suggested_level=""))
    assert missing.status_code == 400 and "标题" in missing.json()["detail"]

    blank = client.post("/api/m1/manual/preview", files={"file": ("scan.pdf", _blank_pdf(), "application/pdf")}, headers=H_IN)
    assert blank.status_code == 400 and "扫描" in blank.json()["detail"]
    old = client.post("/api/m1/manual/preview", files={"file": ("old.doc", b"\xd0\xcf\x11\xe0not-a-docx", "application/msword")}, headers=H_IN)
    assert old.status_code == 400 and "docx" in old.json()["detail"]
    text_pdf = _pdf("Financing guarantee subsidy notice for small firms. Fictional text only.")
    saved = _upload(client, text_pdf, "note.pdf", _fields(title="示例补录 PDF 正文通知文件", doc_no="示测〔2026〕12号", level="B1", suggested_level="B1"))
    assert saved.status_code == 200, saved.text
    assert "Financing guarantee" in saved.json()["item"]["content"]
    monkeypatch.setattr("app.m1.manual.MAX_UPLOAD", 16)
    big = client.post("/api/m1/manual/preview", files={"file": ("big.pdf", b"%PDF" + b"x" * 40, "application/pdf")}, headers=H_IN)
    assert big.status_code == 400 and "过大" in big.json()["detail"]
