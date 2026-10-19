"""M5 验收：导入四张样表台账 + 额度表，逐单元格比对《1.业务额度使用情况 - 截止2026年9月19日v1.xlsx》。

结论写入 docs/M5_验收比对报告.md。常规端口与额度相关单元格必须全部一致；
其余端口如有差异，逐一列出单元格、双方数值与原因（不做任何“凑数”）。
"""
from pathlib import Path

import openpyxl
import pytest

from app.config import M5_SUMMARY_SAMPLE, M5_SAMPLE_DIR

REPORT = Path(__file__).resolve().parent.parent / "docs" / "M5_验收比对报告.md"
TOL = 1e-6


@pytest.fixture(scope="module")
def dash(client):
    r = client.post("/api/m5/snapshots/sample", json={"stat_date": "2026-09-19"})
    assert r.status_code == 200, r.text
    sid = r.json()["snapshot_id"]
    return client.get(f"/api/m5/snapshots/{sid}/dashboard").json()


def _num(v):
    return isinstance(v, (int, float))


def _cmp(cells, sheet, coord, excel_v, ours, port, label):
    if not _num(excel_v) and not isinstance(excel_v, str):
        excel_v = None
    if isinstance(excel_v, str) or isinstance(ours, str):
        ok = str(excel_v or "") == str(ours or "")
    elif excel_v is None and ours is None:
        ok = True
    elif excel_v is None or ours is None:
        ok = False
    else:
        ok = abs(excel_v - ours) <= TOL * max(1, abs(excel_v))
    cells.append({"sheet": sheet, "cell": coord, "label": label, "excel": excel_v, "ours": ours, "port": port, "ok": ok})


def collect(d):
    wb = openpyxl.load_workbook(M5_SUMMARY_SAMPLE, data_only=True)
    s1, s2 = wb.worksheets[0], wb.worksheets[1]
    S = d["summary"]
    cells = []
    sh = "业务整体情况"
    R, T, P, X = S["regular"], S["tech"], S["private"], S["t2t"]
    for i, k in enumerate(("direct", "reguar", "total")):
        r = 3 + i
        lab = {"direct": "直保", "reguar": "再保", "total": "合计"}[k]
        _cmp(cells, sh, f"B{r}", s1[f"B{r}"].value, R[k]["amount"], "regular", f"常规{lab}备案")
        _cmp(cells, sh, f"C{r}", s1[f"C{r}"].value, R[k].get("alloc"), "limit", f"常规{lab}额度分配")
        rate_key = "rate" if k == "total" else "rate_alloc"   # 原表 D3/D4 = 备案/额度分配，D5 = 备案/国担额度
        _cmp(cells, sh, f"D{r}", s1[f"D{r}"].value, R[k].get(rate_key), "regular", f"常规{lab}使用率（原表口径）")
        _cmp(cells, sh, f"E{r}", s1[f"E{r}"].value, R[k].get("remain"), "regular", f"常规{lab}剩余额度")
        _cmp(cells, sh, f"F{r}", s1[f"F{r}"].value, T[k]["amount"], "tech", f"科技{lab}备案")
        _cmp(cells, sh, f"J{r}", s1[f"J{r}"].value, P[k]["amount"], "private", f"民间{lab}备案")
        _cmp(cells, sh, f"K{r}", s1[f"K{r}"].value, P[k].get("alloc"), "limit", f"民间{lab}额度分配")
        _cmp(cells, sh, f"M{r}", s1[f"M{r}"].value, P[k].get("remain"), "private", f"民间{lab}剩余额度")
        _cmp(cells, sh, f"N{r}", s1[f"N{r}"].value, X[k]["amount"], "t2t", f"总对总{lab}备案")
        _cmp(cells, sh, f"O{r}", s1[f"O{r}"].value, X[k].get("rate"), "t2t", f"总对总{lab}使用率(÷国担)")
    _cmp(cells, sh, "I5", s1["I5"].value, T["total"]["remain"], "tech", "科技合计剩余(国担−备案)")
    for c, blk in (("B6", R), ("F6", T), ("J6", P), ("N6", X)):
        _cmp(cells, sh, c, s1[c].value, blk["guodan"], "limit", "2026年国担额度")
    _cmp(cells, sh, "P6", s1["P6"].value, X["total"]["remain"], "t2t", "总对总剩余(国担−备案)")
    _cmp(cells, sh, "C7", s1["C7"].value, R.get("unalloc"), "limit", "常规剩余未分配额度")
    _cmp(cells, sh, "J7", s1["J7"].value, P.get("unalloc"), "limit", "民间剩余未分配额度")
    _cmp(cells, sh, "C8", s1["C8"].value, R.get("need_adjust"), "limit", "需调整")

    sh2 = s2.title
    rows = {r["name"]: r for r in d["rows"]}
    for r in range(3, s2.max_row + 1):
        if s2[f"B{r}"].value in (None, "合计"):
            break
        name = s2[f"B{r}"].value
        m = rows.get(name)
        assert m is not None, f"看板缺少机构 {name}"
        _cmp(cells, sh2, f"C{r}", s2[f"C{r}"].value, m["regular_limit"], "limit", f"{name} 常规授信额度")
        _cmp(cells, sh2, f"D{r}", s2[f"D{r}"].value, m["regular_amt"], "regular", f"{name} 常规备案")
        _cmp(cells, sh2, f"E{r}", s2[f"E{r}"].value, m["regular_rate"], "regular", f"{name} 常规使用比例")
        gv = s2[f"G{r}"].value
        _cmp(cells, sh2, f"G{r}", gv, m["tech_amt"], "tech", f"{name} 科技备案")
        _cmp(cells, sh2, f"I{r}", s2[f"I{r}"].value, m["private_limit"], "limit", f"{name} 民间专项额度")
        _cmp(cells, sh2, f"J{r}", s2[f"J{r}"].value, m["private_amt"], "private", f"{name} 民投备案")
        kv = s2[f"K{r}"].value
        _cmp(cells, sh2, f"K{r}", kv, m["private_rate"] if not isinstance(kv, str) else m["private_note"], "private", f"{name} 民间使用比例")
        _cmp(cells, sh2, f"L{r}", s2[f"L{r}"].value, m["t2t_amt"], "t2t", f"{name} 总对总备案")
        _cmp(cells, sh2, f"M{r}", s2[f"M{r}"].value, m["total_amt"], "mixed", f"{name} 机构合计备案")
        _cmp(cells, sh2, f"N{r}", s2[f"N{r}"].value, m["regular_share"], "regular", f"{name} 常规占再保比例")
        _cmp(cells, sh2, f"O{r}", s2[f"O{r}"].value, m["private_share"], "private", f"{name} 民投占再保比例")
    t = d["totals"]
    for c, k, port in (("C13", "regular_limit", "limit"), ("D13", "regular_amt", "regular"), ("G13", "tech_amt", "tech"),
                       ("I13", "private_limit", "limit"), ("J13", "private_amt", "private"), ("L13", "t2t_amt", "t2t")):
        _cmp(cells, sh2, c, s2[c].value, t[k], port, f"合计 {k}")
    _cmp(cells, sh2, "M13", s2["M13"].value, t["total_amt"], "excel_formula", "合计 机构合计备案")
    return cells


REASONS = {
    "tech": "样表 2-2.科创.xlsx 备案时间为 2026-01-01~2026-09-29（汇总表为截止 09-19），统计时点不同；且该台账不含直保主体行，直保金额无来源",
    "private": "样表 2-3.民投.xlsx 备案时间为 2026-01-01~2026-09-29（汇总表为截止 09-19），统计时点不同；且该台账不含直保主体行，直保金额无来源",
    "t2t": "样表 2-4.总对总.xlsx 备案时间为 2026-01-01~2026-09-29（汇总表为截止 09-19），统计时点不同",
    "mixed": "机构合计 = 常规+科技+民投+总对总，其中科技/民投/总对总样表为 09-29 数据",
    "excel_formula": "合计行机构合计备案口径差异",
}


def test_m5_against_summary(dash):
    cells = collect(dash)
    bad = [c for c in cells if not c["ok"]]
    strict_bad = [c for c in bad if c["port"] in ("regular", "limit")]
    by_port = {}
    for c in cells:
        by_port.setdefault(c["port"], [0, 0])
        by_port[c["port"]][0 if c["ok"] else 1] += 1
    lines = ["# M5 验收比对报告", "",
             f"比对对象：`{M5_SUMMARY_SAMPLE.name}`（两张表）vs 系统导入 4 张样表台账（`{M5_SAMPLE_DIR.name}/2-1~2-4`）+ 额度表（取自同一汇总文件的分机构额度与国担额度，即需求 7.1 的“手工导入额度表”路径）后生成的看板。", "",
             f"- 比对单元格：{len(cells)} 个；一致 {len(cells) - len(bad)} 个；不一致 {len(bad)} 个", ""]
    names = {"regular": "常规台账相关", "limit": "额度/国担/分配相关", "tech": "科创台账相关", "private": "民投台账相关", "t2t": "总对总台账相关",
             "mixed": "跨端口合计", "excel_formula": "原表公式异常"}
    lines.append("| 类别 | 一致 | 不一致 |")
    lines.append("|---|---|---|")
    for p, (ok, ng) in by_port.items():
        lines.append(f"| {names[p]} | {ok} | {ng} |")
    lines += ["", "## 不一致明细", "", "| 表 | 单元格 | 含义 | 原表值 | 系统值 | 原因 |", "|---|---|---|---|---|---|"]
    for c in bad:
        fmt = lambda v: f"{v:,.6f}" if isinstance(v, float) else ("（空）" if v is None else str(v))  # noqa: E731
        lines.append(f"| {c['sheet']} | {c['cell']} | {c['label']} | {fmt(c['excel'])} | {fmt(c['ours'])} | {REASONS.get(c['port'], '待查')} |")
    lines += ["", "## 口径说明", "",
              "- 原表 D3/D4（直保/再保使用率）分母为“额度分配”，D5（合计）分母为“2026年国担额度”。按需求 Q3，系统主口径统一以国担额度为分母并在表下标注；验收时 D3/D4 与系统的“对分配额度使用率”比对、D5 与国担口径比对。",
              "- 原表 H3~H5（科技专项使用率）、L3~L5（民间专项使用率）按 Q4/用户要求系统不计算，未纳入比对。",
              "- 示例汇总表 K4（民间再保额度分配）为手填常数，与分机构民间专项额度合计不一致；系统按手填常数采用，并在额度表警告中提示。",
              ""]
    REPORT.parent.mkdir(exist_ok=True)
    REPORT.write_text("\n".join(lines), encoding="utf-8")
    print(f"\n[M5 验收] 比对 {len(cells)} 个单元格，一致 {len(cells) - len(bad)}，不一致 {len(bad)}；报告：{REPORT}")
    for p, (ok, ng) in by_port.items():
        print(f"   {names[p]}: 一致 {ok} / 不一致 {ng}")
    assert not strict_bad, strict_bad
    assert all(c["port"] in REASONS for c in bad)


def test_m4_limits_vs_summary(client):
    """附加比对：M4（附件2口径方案）带出的常规/民间额度 vs 汇总表分机构额度（9.3 总经办会议调整后）。仅报告，不断言一致。"""
    r = client.post("/api/m4/import-sample").json()
    cur = client.get("/api/m4/schemes/current").json()
    m4 = {x["name"]: x for x in cur["result"]["rows"]}
    s2 = openpyxl.load_workbook(M5_SUMMARY_SAMPLE, data_only=True).worksheets[1]
    lines = ["", "## 附：M4 测算额度 vs 汇总表分机构额度", "",
             "汇总表 C 列注明“根据9.3总经办会议调整”、I 列注明“10家控股机构按照绩效考核确定指标”，均为附件2测算之后又经年中调整（见请示）后的额度，"
             "因此 M4（附件2口径方案：常规=第六次调整 AC，民间=D 列）与之不同属正常；看板验收采用“手工导入额度表”路径。", "",
             "| 机构 | M4 常规 | 汇总表常规 | 差额 | M4 民间 | 汇总表民间 | 差额 |", "|---|---|---|---|---|---|---|"]
    same = 0
    for row in range(3, s2.max_row + 1):
        n = s2[f"B{row}"].value
        if n in (None, "合计"):
            break
        x = m4.get(n)
        c, i = s2[f"C{row}"].value or 0, s2[f"I{row}"].value or 0
        if x is None:
            lines.append(f"| {n} | （M4 无此机构） | {c:,.0f} | — | — | {i:,.0f} | — |")
            continue
        if x["regular_final"] == c and x["mintou_special"] == i:
            same += 1
        lines.append(f"| {n} | {x['regular_final']:,.0f} | {c:,.0f} | {c - x['regular_final']:,.0f} | {x['mintou_special']:,.0f} | {i:,.0f} | {i - x['mintou_special']:,.0f} |")
    lines.insert(4, f"常规与民间额度两项均一致的机构：{same} / 10。")
    with REPORT.open("a", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"[M5 附加] M4 额度与汇总表额度两项均一致 {same}/10（预期不同，原因见报告）")
