"""同事补录：粘贴链接走现有抓取/清洗/分级/摘要；上传 PDF 或 Word 抽正文后由模型预填元数据。

两条路径都记 provider=manual、上传人，效力默认「未知」，reviewed_at 留空进入待复核。
清晰重复（同一链接，或标题+发布机构+文号全等）直接拒绝并返回已有条目；仅标题规范化后相同可拒绝，或仍入库并写入 duplicate_of。
"""
import io
import json
import os
import re
import sqlite3
from datetime import datetime
from urllib.parse import urlparse

from fastapi import HTTPException

from .. import db
from ..config import DATA_DIR
from ..llm import service as L
from . import collector, config as C
from . import fetch as fetch_mod
from .classify import classify, norm_url
from .clean import clean_text
from .validity import annotate

PROVIDER = "manual"
TOPIC = "同事补录"
MAX_UPLOAD = 20 * 1024 * 1024
MAX_TEXT = 20000
MAX_PAGES = 40
MIN_TEXT = 20
LEVELS = C.LEVELS
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATE_FIND = re.compile(r"(20\d{2})\s*[-年/.]\s*(\d{1,2})\s*[-月/.]\s*(\d{1,2})")
_LEVEL = re.compile(r"A1|A2|B1|B2|C1|C2")
EMPTY_META = {"title": "", "publisher": "", "doc_no": "", "pub_date": "", "summary": "", "level": ""}
_PROMPT = (
    "你是政策文件著录助手。请只依据给定正文抽取元数据，正文没有的信息留空字符串，不得推测或编造。\n"
    "只输出一个 JSON 对象，不要 Markdown。字段：\n"
    '{"title":"文件标题","publisher":"发布机构","doc_no":"文号","pub_date":"YYYY-MM-DD或空",'
    '"summary":"3到5条要点，每条一行并以- 开头","level":"A1或A2或B1或B2或C1或C2"}\n'
    "等级：A1 国务院及部委正式文件；A2 部委其他信息或省级及以下政府正式文件；B1 中央媒体；"
    "B2 商业媒体；C1 公众号转载；C2 其他。不确定时选最接近的一项。\n正文：\n"
)


def _flag(v):
    if isinstance(v, bool):
        return v
    return str(v or "").strip().lower() in ("1", "true", "yes", "on")


def _http_url(raw):
    u = (raw or "").strip()
    if not u:
        raise ValueError("请填写政策链接")
    if len(u) > 2000:
        raise ValueError("链接过长")
    p = urlparse(u)
    if p.scheme not in ("http", "https") or not p.hostname:
        raise ValueError("仅支持 http 或 https 链接")
    return norm_url(u)


def _norm_date(s):
    s = str(s or "").strip()
    if not s:
        return ""
    if _DATE.match(s):
        cand = s
    else:
        m = _DATE_FIND.search(s)
        if not m:
            return ""
        cand = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    try:
        datetime.strptime(cand, "%Y-%m-%d")
    except ValueError:
        return ""
    return cand


def _user_date(s):
    raw = (s or "").strip()
    if not raw:
        return ""
    out = _norm_date(raw)
    if not out:
        raise HTTPException(400, "发布日期须为 YYYY-MM-DD，或留空")
    return out


def _clip(s, n):
    return str(s or "").strip()[:n]


def _json_obj(text):
    t = (text or "").strip()
    t = re.sub(r"^```(?:json)?\s*", "", t, flags=re.I)
    t = re.sub(r"\s*```$", "", t)
    i = t.find("{")
    if i < 0:
        return {}
    try:
        obj, _ = json.JSONDecoder().raw_decode(t[i:])
    except json.JSONDecodeError:
        return {}
    return obj if isinstance(obj, dict) else {}


def _summary(v):
    if isinstance(v, list):
        lines = []
        for x in v:
            s = str(x).strip()
            if not s:
                continue
            lines.append(s if s.startswith("-") else "- " + s)
        return "\n".join(lines)[:8000]
    return _clip(v, 8000)


def _level(v):
    m = _LEVEL.search(str(v or "").upper())
    return m.group(0) if m else ""


def _reject(row, kind, can_force, message):
    raise HTTPException(409, {
        "message": message,
        "existing_id": row["id"],
        "existing_title": row.get("title") or "",
        "kind": kind,
        "can_force": can_force,
    })


def _url_hit(url):
    if not url:
        return None
    return db.q1("SELECT id, title, url FROM m1_items WHERE url=?", (url,))


def _identity_hit(title, publisher, doc_no):
    title, publisher, doc_no = (title or "").strip(), (publisher or "").strip(), (doc_no or "").strip()
    if not (title and publisher and doc_no):
        return None
    return db.q1(
        "SELECT id, title, url FROM m1_items WHERE duplicate_of IS NULL AND title=? AND publisher=? AND doc_no=? ORDER BY id LIMIT 1",
        (title, publisher, doc_no),
    )


def _title_hit(title):
    nt = collector._norm_title(title)
    if not nt:
        return None
    for o in db.q("SELECT id, title, url FROM m1_items WHERE duplicate_of IS NULL AND COALESCE(title,'') != ''"):
        if collector._norm_title(o["title"]) == nt:
            return o
    return None


def _resolve_dup(url, title, publisher, doc_no, allow_title_dup):
    """返回应写入的 duplicate_of；清晰重复则 409，标题相同且未允许时也 409。"""
    hit = _url_hit(url)
    if hit:
        _reject(hit, "url", False, f"该链接已采集，未重复入库。已有条目 #{hit['id']}《{hit.get('title') or '无标题'}》。")
    hit = _identity_hit(title, publisher, doc_no)
    if hit:
        _reject(hit, "identity", False,
                f"标题、发布机构与文号与已有条目 #{hit['id']}《{hit.get('title') or ''}》完全一致，已拒绝入库。")
    hit = _title_hit(title)
    if hit:
        if allow_title_dup:
            return hit["id"]
        _reject(hit, "title", True,
                f"规范化标题与已有条目 #{hit['id']}《{hit.get('title') or ''}》相同，未入库。若确认不是同一文件，请勾选「标题相同仍入库」后重试。")
    return None


def _orig_name(filename):
    name = os.path.basename((filename or "").replace("\\", "/"))
    name = re.sub(r"[\x00-\x1f]", "", name).strip()[:120]
    return name or "upload"


def _store_name(iid, filename):
    base = re.sub(r"[^\w.\-\u4e00-\u9fff]+", "_", _orig_name(filename))[:80].strip("._")
    return f"{iid}_{base or 'upload'}"


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:
        raise ValueError("PDF 无法读取（文件可能已损坏或加密）") from e
    if getattr(reader, "is_encrypted", False):
        try:
            if reader.decrypt("") == 0:
                raise ValueError("PDF 已加密，无法抽出正文")
        except ValueError:
            raise
        except Exception as e:
            raise ValueError("PDF 已加密，无法抽出正文") from e
    parts = []
    for page in reader.pages[:MAX_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ValueError("不是有效的 Word（.docx）文件") from e
    lines = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = (c.text or "").strip()
                if t and t not in cells:
                    cells.append(t)
            if cells:
                lines.append(" ".join(cells))
    return "\n".join(lines)


def extract_upload(filename, data: bytes):
    """抽出并清洗正文。失败抛 ValueError（中文，可直接展示）。返回 (text, kind)。"""
    if not data:
        raise ValueError("文件为空")
    if len(data) > MAX_UPLOAD:
        raise ValueError(f"文件过大，请上传 {MAX_UPLOAD // (1024 * 1024)}MB 以内的 PDF 或 Word")
    name = (filename or "").strip().lower()
    if name.endswith(".pdf"):
        if not data.startswith(b"%PDF"):
            raise ValueError("文件扩展名是 PDF，但内容不是 PDF")
        kind = "pdf"
        raw = _pdf_text(data)
    elif name.endswith(".docx"):
        if not data.startswith(b"PK"):
            raise ValueError("文件扩展名是 docx，但内容不是 Word（.docx）")
        kind = "docx"
        raw = _docx_text(data)
    elif name.endswith(".doc"):
        raise ValueError("暂不支持旧版 Word（.doc），请另存为 .docx 后上传")
    else:
        raise ValueError("仅支持 PDF 与 Word（.docx）")
    text = clean_text(raw.replace("\x00", "")) or ""
    if len(text) < MIN_TEXT:
        if kind == "pdf":
            raise ValueError("未能从该 PDF 抽出可复制的正文，可能是扫描件或图片版。请换可复制文字的 PDF，或改为粘贴政策链接。")
        raise ValueError("未能从该 Word 文件抽出正文，请确认文档中含有文字。")
    return text[:MAX_TEXT], kind


def suggest(text, user):
    """调用 M1 摘要模型抽取著录字段。模型不可用时由调用方转为可读提示。"""
    r = L.run("m1.extract", [{"role": "user", "content": _PROMPT + (text or "")[:6000]}], user, ref="m1_manual")
    data = _json_obj(r.get("content") or "")
    out = {
        "title": _clip(data.get("title"), 300),
        "publisher": _clip(data.get("publisher"), 200),
        "doc_no": _clip(data.get("doc_no"), 80),
        "pub_date": _norm_date(data.get("pub_date")),
        "summary": _summary(data.get("summary")),
        "level": _level(data.get("level")),
    }
    warn = "" if any(out.values()) else "模型未返回可解析的元数据，请手工填写。"
    return out, warn


def try_suggest(text, user):
    try:
        return suggest(text, user)
    except HTTPException as e:
        msg = e.detail if isinstance(e.detail, str) else "模型调用失败"
        return dict(EMPTY_META), f"模型未能预填元数据：{msg}。请手工填写后再入库。"


def preview_file(filename, data, user):
    try:
        text, _kind = extract_upload(filename, data)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    sugg, warn = try_suggest(text, user)
    return {
        "filename": _orig_name(filename),
        "chars": len(text),
        "text": text[:4000],
        "suggestions": sugg,
        "llm_warning": warn,
    }


def _insert(user, **kw):
    now = db.now()
    try:
        return db.execute(
            "INSERT INTO m1_items(url,title,publisher,level,layer,region,topic,pub_date,snippet,content,summary,query,provider,"
            "fetch_status,duplicate_of,level_manual,collected_at,fetched_at,validity_status,doc_no,uploaded_by,source_filename)"
            " VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (kw.get("url") or None, kw.get("title"), kw.get("publisher"), kw.get("level"), kw.get("layer") or "",
             kw.get("region") or "", TOPIC, kw.get("pub_date") or "", (kw.get("snippet") or "")[:180], kw.get("content") or "",
             kw.get("summary") or "", "同事补录", PROVIDER, kw.get("fetch_status") or "", kw.get("duplicate_of"),
             1 if kw.get("level_manual") else 0, now, now, "未知", kw.get("doc_no") or "",
             (user or {}).get("name") or (user or {}).get("username") or "", kw.get("source_filename")),
        )
    except sqlite3.IntegrityError:
        hit = _url_hit(kw.get("url"))
        if hit:
            _reject(hit, "url", False, f"该链接已采集，未重复入库。已有条目 #{hit['id']}《{hit.get('title') or '无标题'}》。")
        raise HTTPException(409, "入库失败：与已有条目冲突") from None


def _store_upload(iid, filename, data):
    folder = DATA_DIR / "m1_uploads"
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / _store_name(iid, filename)
    if dest.resolve().parent != folder.resolve():
        raise HTTPException(400, "文件名无效")
    try:
        dest.write_bytes(data)
    except OSError as e:
        raise HTTPException(500, "文件保存失败") from e
    rel = f"m1_uploads/{dest.name}"
    db.execute("UPDATE m1_items SET source_path=?, source_filename=? WHERE id=?", (rel, _orig_name(filename), iid))
    return rel


def _finish(iid, user, via, warning, duplicate_of=None, filename=""):
    it = db.q1("SELECT * FROM m1_items WHERE id=?", (iid,))
    db.audit(user, "M1", "同事补录", iid, {
        "via": via, "provider": PROVIDER, "title": (it or {}).get("title") or "",
        "filename": filename or "", "duplicate_of": duplicate_of,
    })
    return {"ok": True, "item": annotate(it), "warning": warning or ""}


def _drop(iid):
    db.execute("DELETE FROM m1_items WHERE id=?", (iid,))


def ingest_url(raw, user, allow_title_dup=False):
    try:
        url = _http_url(raw)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    hit = _url_hit(url)
    if hit:
        _reject(hit, "url", False, f"该链接已采集，未重复入库。已有条目 #{hit['id']}《{hit.get('title') or '无标题'}》。")
    try:
        fetched = fetch_mod.fetch(url)
    except ValueError as e:
        raise HTTPException(400, str(e)[:200]) from e
    except Exception as e:
        raise HTTPException(502, f"抓取失败：{str(e)[:200]}") from e
    c = classify(url, C.get()) or {}
    title = _clip(fetched.get("title"), 300)
    if not title:
        title = _clip("未命名政策 " + (urlparse(url).path or url), 300)
    text = (fetched.get("text") or "").replace("\x00", "")
    pub_date = _norm_date(fetched.get("date") or "")
    level = c.get("level") if c.get("level") in LEVELS else "C2"
    publisher = _clip(c.get("publisher"), 200)
    kind = fetched.get("kind") or "html"
    warns = []
    if kind == "pdf":
        status = "PDF（未解析正文）"
        warns.append("该链接指向 PDF，采集链路不解析 PDF 正文。如需正文，请改用「上传文件」。")
    elif not text.strip():
        status = "抓取失败：页面未抽出正文"
        warns.append("页面未抽出正文，条目已入库并进入待复核。")
    else:
        status = "已抓取"
    dup_of = _resolve_dup(url, title, publisher, "", _flag(allow_title_dup))
    if dup_of:
        warns.append(f"已入库，并标记为与 #{dup_of} 疑似重复（默认列表不显示）。")
    iid = _insert(user, url=url, title=title, publisher=publisher, level=level, layer=c.get("layer") or "",
                  pub_date=pub_date, snippet=text[:180], content=text[:MAX_TEXT], summary="", fetch_status=status,
                  duplicate_of=dup_of, level_manual=0, doc_no="")
    if text.strip():
        try:
            if not collector._summarize(iid):
                warns.append("未配置 M1 要点摘要模型，条目已入库但没有 AI 要点，可在详情中稍后提炼。")
        except Exception as e:  # noqa: BLE001  摘要失败不回滚已入库条目
            detail = getattr(e, "detail", None)
            msg = detail if isinstance(detail, str) else str(e)
            warns.append(f"AI 摘要未生成：{str(msg)[:120]}。条目已入库，可稍后在详情中重试。")
    return _finish(iid, user, "url", " ".join(warns), dup_of)


def ingest_file(filename, data, user, fields):
    try:
        text, _kind = extract_upload(filename, data)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    fields = fields or {}
    given = {
        "title": _clip(fields.get("title"), 300),
        "publisher": _clip(fields.get("publisher"), 200),
        "doc_no": _clip(fields.get("doc_no"), 80),
        "summary": _clip(fields.get("summary"), 8000),
        "level": _clip(fields.get("level"), 20),
        "url": str(fields.get("url") or "").strip(),
    }
    ai_level = _level(fields.get("suggested_level"))
    sugg, swarn = dict(EMPTY_META), ""
    if not given["title"] or not given["summary"] or not given["level"]:
        sugg, swarn = try_suggest(text, user)
        if not ai_level:
            ai_level = sugg.get("level") or ""
    title = given["title"] or sugg.get("title") or ""
    if not title:
        raise HTTPException(400, "请填写标题后再入库" + (f"（{swarn}）" if swarn else ""))
    publisher = given["publisher"] or sugg.get("publisher") or ""
    doc_no = given["doc_no"] or sugg.get("doc_no") or ""
    summary = given["summary"] or sugg.get("summary") or ""
    if (fields.get("pub_date") or "").strip():
        pub_date = _user_date(fields.get("pub_date"))
    else:
        pub_date = _norm_date(sugg.get("pub_date"))
    level_in = _level(given["level"]) if given["level"] else ""
    if given["level"] and not level_in:
        raise HTTPException(400, "来源等级无效")
    if level_in:
        level = level_in
        level_manual = 0 if ai_level and level_in == ai_level else 1
    elif sugg.get("level") in LEVELS:
        level = sugg["level"]
        level_manual = 0
    else:
        level, level_manual = "", 0
    url, layer = "", ""
    if given["url"]:
        try:
            url = _http_url(given["url"])
        except ValueError as e:
            raise HTTPException(400, str(e)) from e
        c = classify(url, C.get()) or {}
        layer = c.get("layer") or ""
        if not publisher:
            publisher = _clip(c.get("publisher"), 200)
        if not level and c.get("level") in LEVELS:
            level = c["level"]
            level_manual = 0
    if level not in LEVELS:
        raise HTTPException(400, "请选择来源等级（A1 / A2 / B1 / B2 / C1 / C2）")
    dup_of = _resolve_dup(url, title, publisher, doc_no, _flag(fields.get("allow_title_dup")))
    warns = []
    if swarn:
        warns.append(swarn)
    if dup_of:
        warns.append(f"已入库，并标记为与 #{dup_of} 疑似重复（默认列表不显示）。")
    iid = _insert(user, url=url, title=title, publisher=publisher, level=level, layer=layer, pub_date=pub_date,
                  snippet=text[:180], content=text, summary=summary, fetch_status="已上传解析", duplicate_of=dup_of,
                  level_manual=level_manual, doc_no=doc_no, source_filename=_orig_name(filename))
    try:
        _store_upload(iid, filename, data)
    except HTTPException:
        _drop(iid)
        raise
    return _finish(iid, user, "file", " ".join(warns), dup_of, _orig_name(filename))
