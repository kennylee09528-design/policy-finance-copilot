"""导出文件水印 / 导出人信息（需求第 10 章：内部资料、注意保密、可追溯）。

统一在路由层对导出字节做后处理，不改各模块生成逻辑：
- Word：页眉（导出人 · 时间 · 内部资料 注意保密）+ 页脚 + 文档属性
- Excel：每个工作表打印页眉/页脚 + 文档属性 + 末尾「导出信息」表
- PPT：每页右下角小字 + 文档属性
- Markdown：文末引用块
- PDF：浏览器打印视图（前端 printView，带平铺水印），并通过 /api/exports/log 记入导出历史
"""
import io
from datetime import datetime

SYSTEM = "普惠业务工作台"


def info(user):
    who = user.get("name") or "未知"
    if user.get("username") and user.get("username") != who:
        who += f"（{user['username']}）"
    if user.get("org"):
        who += f" · {user['org']}"
    return {"who": who, "at": datetime.now().strftime("%Y-%m-%d %H:%M"), "text": f"内部资料 注意保密 · 导出人：{who} · 导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')} · {SYSTEM}"}


def _gray(run, size=8):
    from docx.shared import Pt, RGBColor
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def docx_bytes(data: bytes, user) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    i = info(user)
    d = Document(io.BytesIO(data))
    for s in d.sections:
        hp = s.header.add_paragraph() if s.header.paragraphs and s.header.paragraphs[0].text else (s.header.paragraphs[0] if s.header.paragraphs else s.header.add_paragraph())
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _gray(hp.add_run(i["text"]))
        fp = s.footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _gray(fp.add_run(f"本文件由{SYSTEM}导出，仅限内部使用，请勿外传 · {i['who']} · {i['at']}"))
    cp = d.core_properties
    cp.last_modified_by = i["who"][:250]
    cp.comments = i["text"][:250]
    cp.category = "内部资料"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def xlsx_bytes(data: bytes, user) -> bytes:
    from openpyxl import load_workbook
    i = info(user)
    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        ws.oddHeader.right.text = f"内部资料 注意保密 · {i['who']} · {i['at']}"
        ws.oddHeader.right.size = 8
        ws.oddFooter.center.text = f"{SYSTEM} 导出 · 第 &P 页 / 共 &N 页"
        ws.oddFooter.center.size = 8
    if "导出信息" not in wb.sheetnames:
        ws = wb.create_sheet("导出信息")
        for row in (["项目", "内容"], ["密级", "内部资料 注意保密"], ["导出人", i["who"]], ["导出时间", i["at"]], ["来源系统", SYSTEM],
                    ["说明", "金额单位：万元；本文件由系统根据数据库实时数据生成，人工修改后与系统数据可能不一致"]):
            ws.append(row)
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 70
    wb.properties.lastModifiedBy = i["who"][:250]
    wb.properties.description = i["text"][:250]
    wb.properties.category = "内部资料"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pptx_bytes(data: bytes, user) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Pt
    i = info(user)
    p = Presentation(io.BytesIO(data))
    W, H = p.slide_width, p.slide_height
    for s in p.slides:
        tb = s.shapes.add_textbox(Emu(int(W * 0.45)), Emu(int(H - Pt(18))), Emu(int(W * 0.54)), Emu(int(Pt(14))))
        tb.name = "ph_watermark"
        tf = tb.text_frame
        tf.margin_top = tf.margin_bottom = 0
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.RIGHT
        r = para.add_run()
        r.text = f"内部资料 注意保密 · {i['who']} · {i['at']}"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xA0, 0xA0, 0xA0)
    cp = p.core_properties
    cp.last_modified_by = i["who"][:250]
    cp.comments = i["text"][:250]
    cp.category = "内部资料"
    out = io.BytesIO()
    p.save(out)
    return out.getvalue()


def md_text(text: str, user) -> str:
    return (text or "").rstrip() + f"\n\n---\n> {info(user)['text']}\n"
