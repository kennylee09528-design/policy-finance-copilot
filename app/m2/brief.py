"""今日行业 Brief：M1 增量条目 → 重要度打分（4.2）→ 分档 → 要点/关键数字（含来源校验）→ 按 4.3 模板生成 → 编辑 → 发布 → 归档/导出。"""
import io
import json
import os
import re
import threading
import time
from datetime import datetime, timedelta

from fastapi import HTTPException

from .. import db
from ..m1.config import GD_CITIES
from ..textcheck import key_numbers, number_in_source
from . import config as C

WEEK = "一二三四五六日"
CATS = [("国担基金", r"国家融资担保基金|国担基金"), ("风险补偿与代偿", r"风险补偿|代偿"), ("银担/政银担合作", r"政银担|银担|分险"),
        ("科技金融/专精特新", r"科技金融|科创|专精特新|科技型"), ("民间投资", r"民间投资"), ("三农", r"三农|农业|乡村|农户"),
        ("就业创业", r"就业|创业"), ("政府性融资担保", r"政府性融资担保|再担保|融资担保"), ("普惠金融", r"普惠|小微"),
        ("财政金融协同", r"财政|贴息|补贴|专项资金")]
SYS = {"name": "M2定时生成", "role": "internal"}


def layer_key(it):
    lay, reg = it.get("layer") or "", it.get("region") or ""
    if lay == "国家" or reg == "全国":
        return "国家"
    if lay == "省级":
        return "广东本省" if reg == "广东" else "其他省级"
    if lay == "重点特殊财政管理城市":
        return "重点城市"
    if lay in ("普通地市", "地市"):
        return "地市"
    return "其他"


def domestic(it):
    return (it.get("region") or "") in ["全国", "广东", "深圳"] + GD_CITIES or layer_key(it) == "国家"


def _parse_dt(s):
    for f in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime((s or "")[:19], f)
        except ValueError:
            continue
    return None


def score(it, cfg, ref):
    text = " ".join([it.get("title") or "", it.get("summary") or "", it.get("snippet") or "", (it.get("content") or "")[:3000]])
    lv = cfg["level_scores"].get(it.get("level"), 0)
    ly = cfg["layer_scores"].get(layer_key(it), cfg["layer_scores"].get("其他", 50))
    hits = [k for k in cfg["keywords"] if k in text]
    rel = min(100, sum(cfg["keywords"][k] for k in hits))
    d = _parse_dt(it.get("pub_date")) or _parse_dt(it.get("collected_at")) or ref
    age = max(0.0, (ref - d).total_seconds() / 86400)
    tm = cfg["timeliness_older"]
    for t in sorted(cfg["timeliness"], key=lambda x: x["days"]):
        if age <= t["days"]:
            tm = t["score"]
            break
    w = cfg["weights"]
    total = round(lv * w["level"] + ly * w["layer"] + rel * w["relevance"] + tm * w["timeliness"], 1)
    tier = 3 if total >= cfg["tiers"]["three"] else 2 if total >= cfg["tiers"]["two"] else 1
    return {"total": total, "tier": tier, "parts": {"来源等级": lv, "政策层级": ly, "业务相关度": rel, "时效性": tm}, "hits": hits}


def category(it):
    t = (it.get("title") or "") + (it.get("summary") or "") + (it.get("snippet") or "")
    for name, pat in CATS:
        if re.search(pat, t):
            return name
    return "其他"


def _first_sentences(text, n=2):
    from ..m1.clean import clean_text
    # 只取含句读标点的正文行（跳过标题、日期/来源行等），再按句切分
    body = "\n".join(l for l in clean_text(text or "").splitlines() if re.search(r"[。，；！？]", l) and not re.match(r"^(日期|时间|来源|作者)[：:]", l))
    ss = [s.strip() for s in re.split(r"(?<=[。！？；])", body) if len(s.strip()) > 8]
    return ss[:n]


def points_and_numbers(it, cfg, user):
    """要点与关键数字。优先：M1 已有 AI 要点 → M2 模型提炼 → 抽取式（原文句子 + 原文数字）。"""
    src = "\n".join([it.get("title") or "", it.get("snippet") or "", it.get("content") or ""])
    method, points, nums = "抽取", [], []
    if it.get("summary"):
        lines = [l.strip()[2:].strip() for l in it["summary"].splitlines() if l.strip().startswith(("- ", "• ", "* "))]
        if lines:
            points, method = lines[:5], "M1 要点"
    if not points and cfg.get("use_llm"):
        from ..llm.service import assignments, run
        if assignments().get("m2.brief") and (it.get("content") or it.get("snippet")):
            prompt = ("严格依据下列原文，输出 JSON：{\"points\": [不超过3条要点], \"key_numbers\": [原文出现的关键金额/比例/期限，原样摘录]}。"
                      "只能使用原文中的信息，原文没有就输出空数组，不得推测。\n标题：" + (it.get("title") or "") + "\n原文：\n" + src[:5000])
            try:
                r = run("m2.brief", [{"role": "user", "content": prompt}], user, overrides={"temperature": 0.1}, ref=f"m1_item:{it.get('id') or it.get('m1_id') or ''}")
                j = json.loads(re.search(r"\{.*\}", r["content"], re.S).group(0))
                points = [str(x) for x in j.get("points", [])][:3]
                nums = [str(x) for x in j.get("key_numbers", [])][:6]
                method = "M2 模型"
            except Exception:  # noqa: BLE001  模型失败回退抽取式
                points, nums = [], []
    if not points:
        points = _first_sentences(it.get("content") or it.get("snippet") or "")[:2] or ([it["snippet"][:120]] if it.get("snippet") else [])
        method = "抽取"
    from ..m1.clean import clean_point
    points = [q for q in (clean_point(x) for x in points) if q]
    if not nums:
        seen = []
        for frag, _n, _u in key_numbers(" ".join(points) if method != "抽取" else (it.get("content") or it.get("snippet") or "")):
            if frag not in seen:
                seen.append(frag)
        nums = seen[:5]
    # 来源校验：要点与关键数字中的每个数字都必须能在原文中找到
    checks = []
    for frag, n, u in key_numbers(" ".join(points + nums)):
        checks.append({"num": frag, "ok": number_in_source(n, u, src)})
    return points, nums, checks, method


def _window(brief_date, cfg):
    d = datetime.strptime(brief_date, "%Y-%m-%d")
    h, m = map(int, cfg["schedule"]["time"].split(":"))
    end = d.replace(hour=h, minute=m)
    now = datetime.now()
    if d.date() == now.date() and now > end:
        end = now
    return end - timedelta(hours=cfg["window_hours"]), end


def _m5_rows(brief_date):
    s = db.q1("SELECT id, stat_date FROM m5_snapshots WHERE stat_date<=? ORDER BY stat_date DESC, id DESC LIMIT 1", (brief_date,))
    if not s:
        return []
    from ..m5 import service as S
    d = S.dashboard(s["id"], {"name": "M2", "role": "internal", "org": None})
    src = f"M5 机构看板（统计时点 {s['stat_date']}，快照 #{s['id']}）"
    sm = d["summary"]
    rows = [("常规业务累计备案（直保+再保）", sm["regular"]["total"]["amount"], sm["regular"]["total"].get("delta"), "万元"),
            ("常规业务使用率（÷国担额度）", sm["regular"]["total"].get("rate"), None, "%"),
            ("科技创新专项累计备案", sm["tech"]["total"]["amount"], sm["tech"]["total"].get("delta"), "万元"),
            ("民间投资专项累计备案", sm["private"]["total"]["amount"], sm["private"]["total"].get("delta"), "万元"),
            ("总对总累计备案", sm["t2t"]["total"]["amount"], sm["t2t"]["total"].get("delta"), "万元")]
    out = []
    for name, v, dl, unit in rows:
        if v is None:
            continue
        val = f"{v * 100:.2f}%" if unit == "%" else f"{v:,.2f} 万元"
        out.append({"name": name, "value": val, "delta": "—" if dl is None else f"{dl:+,.2f} 万元", "source": src})
    return out


RELEVANCE = {"国担基金": "直接涉及国家融资担保基金合作与再担保分险安排，影响再担保业务规模与分险比例",
             "风险补偿与代偿": "涉及代偿补偿/风险补偿资金，关系再担保代偿分担与风险池安排",
             "银担/政银担合作": "涉及银担分险合作，影响再担保体系内合作机构与银行分险模式",
             "科技金融/专精特新": "涉及科技型企业融资支持，关联科技创新专项担保/再担保业务",
             "民间投资": "涉及民间投资融资支持，关联民间投资专项担保业务",
             "三农": "涉及涉农融资支持，关联农业担保与再担保合作",
             "就业创业": "涉及创业担保贷款等，关联政策性担保业务",
             "政府性融资担保": "直接涉及政府性融资担保/再担保体系的政策与监管",
             "普惠金融": "涉及普惠小微融资，关系担保体系服务对象与业务导向",
             "财政金融协同": "涉及财政贴息/补贴/专项资金，关系担保降费奖补与财政支持",
             "其他": "与融资担保相关度较低，供参考"}


def _section(it):
    if it["tier"] == 3:
        return "必读"
    if not domestic(it):
        return "同业动态"
    return "关注" if it["tier"] == 2 else "浏览"


def render(brief_date, items, data_rows, notes=None, meta=None):
    d = datetime.strptime(brief_date, "%Y-%m-%d")
    L = [f"# 今日行业 Brief · {brief_date}（周{WEEK[d.weekday()]}）", ""]
    inc = sorted([i for i in items if i.get("include", True) and not i.get("blocked")], key=lambda x: (x.get("order", 0), -x["score"]))
    meta = meta or {}
    secs = {k: sum(1 for i in inc if i["section"] == k) for k in ("必读", "关注", "浏览", "同业动态")}
    L += ["## 今日要点"]
    if meta.get("window_label"):
        L.append(f"- 取材范围：{meta['window_label']}")
    L.append(f"- 本期收录 {len(inc)} 条：必读 {secs['必读']}、关注 {secs['关注']}、浏览 {secs['浏览']}、同业动态 {secs['同业动态']}")
    for i in inc[:3]:
        L.append(f"- {i['title']}（{i['publisher']}，{i['level']}，{i.get('pub_date') or '未标注'}）")
    L.append("")

    def block(sec_items, start=1):
        out = []
        for n, i in enumerate(sec_items, start):
            out.append(f"{n}. 【{i['category']}｜{i.get('region') or '—'}】{i['title']}")
            pts = [re.sub(r"\s+", " ", p).strip() for p in (i.get("points") or [])]
            good = [p for p in pts if len(p) >= 15 and not re.search(r"首页|发文字号|索引号|主题分类|打开|登录|导航", p)]
            one = good[0] if good else ""
            out.append("   - 一句话摘要：" + (one[:120] + ("…" if len(one) > 120 else "") if one else "（原文正文未能自动提炼，请点原文查看）"))
            rest = [p for p in good[1:]]
            if rest:
                out.append("   - 要点：" + "；".join(x[:200] for x in rest))
            out.append("   - 与再担保关联：" + RELEVANCE.get(i["category"], RELEVANCE["其他"]))
            if i.get("key_numbers"):
                out.append("   - 关键数字：" + "；".join(i["key_numbers"]))
            out.append(f"   - 来源：{i['publisher']}（{i['level']}）｜发布：{i.get('pub_date') or '未标注'}｜[原文]({i['url']})")
        return out or ["（无）"]
    L += ["## 一、今日必读（★★★）"] + block([i for i in inc if i["section"] == "必读"]) + [""]
    L += ["## 二、值得关注（★★）"] + block([i for i in inc if i["section"] == "关注"])
    browse = [i for i in inc if i["section"] == "浏览"]
    if browse:
        L += ["", "**★ 浏览（默认折叠）**"] + [f"- 【{i['category']}｜{i.get('region') or '—'}】{i['title']}（{i['publisher']}，{i['level']}）{i['url']}" for i in browse]
    L += ["", "## 三、同业动态 / 他省做法"] + block([i for i in inc if i["section"] == "同业动态"]) + [""]
    L += ["## 四、数据速览"]
    if data_rows:
        L += ["| 指标 | 数值 | 较上期 | 来源 |", "|---|---|---|---|"] + [f"| {r['name']} | {r['value']} | {r['delta']} | {r['source']} |" for r in data_rows]
    else:
        L += ["（暂无 M5 看板数据）"]
    blocked = [i for i in items if i.get("blocked")]
    single = [i for i in inc if i["level"] != "A1" and i.get("key_numbers")]
    L += ["", "## 五、待核实提示"]
    L += [f"- 关键数字未能在原文中找到来源、已拦截未入简报的条目 {len(blocked)} 条" + ("：" + "；".join(f"{i['title']}（{'、'.join(i.get('block_reason') or [])}）" for i in blocked) if blocked else "")]
    L += [f"- 关键数字仅有非正式文件来源（非 A1）、需以原始文件印证的条目 {len(single)} 条"]
    for n in notes or []:
        L.append(f"- {n}")
    return "\n".join(L) + "\n"


def generate(brief_date, user, trigger="手动"):
    cfg = C.get()
    start, end = _window(brief_date, cfg)
    fmt = "%Y-%m-%d %H:%M:%S"
    oldest = (end - timedelta(days=cfg.get("max_pub_age_days", 30))).strftime("%Y-%m-%d")
    rows = db.q("SELECT * FROM m1_items WHERE duplicate_of IS NULL AND collected_at>? AND collected_at<=? AND (pub_date>=? OR (COALESCE(pub_date,'')='' AND ?)) ORDER BY id",
                (start.strftime(fmt), end.strftime(fmt), oldest, 0 if cfg.get("require_pub_date", True) else 1))
    label = f"近 {cfg['window_hours']} 小时 M1 新采集条目（{start:%m-%d %H:%M} 至 {end:%m-%d %H:%M}）"
    if len(rows) < cfg.get("min_items", 3):
        for nd in cfg.get("fallback_days") or []:
            lo = (end - timedelta(days=nd)).strftime("%Y-%m-%d")
            more = db.q("SELECT * FROM m1_items WHERE duplicate_of IS NULL AND pub_date>=? AND pub_date<=? AND collected_at<=? ORDER BY id", (lo, end.strftime("%Y-%m-%d"), end.strftime(fmt)))
            if len(more) > len(rows):
                rows = more
                label = f"今日窗口内新条目不足 {cfg.get('min_items', 3)} 条，已扩展为近 {nd} 天发布的 M1 条目（{lo} 至 {end:%Y-%m-%d}）"
            if len(rows) >= cfg.get("min_items", 3):
                break
    items = []
    for it in rows:
        sc = score(it, cfg, end)
        items.append({"m1_id": it["id"], "title": it["title"], "url": it["url"], "publisher": it["publisher"], "level": it["level"],
                      "region": it["region"], "layer": it["layer"], "layer_key": layer_key(it), "pub_date": it["pub_date"],
                      "category": category(it), "score": sc["total"], "tier": sc["tier"], "parts": sc["parts"], "hits": sc["hits"], "include": True})
    items.sort(key=lambda x: -x["score"])
    items = items[: cfg["max_items"]]
    full = {r["id"]: r for r in rows}
    for n, i in enumerate(items):
        pts, nums, checks, method = points_and_numbers(full[i["m1_id"]], cfg, user)
        bad = [c["num"] for c in checks if not c["ok"]]
        i.update(points=pts, key_numbers=nums, number_check=checks, points_method=method, order=n,
                 blocked=bool(bad) or not i["url"], block_reason=([f"数字无来源：{'、'.join(bad)}"] if bad else []) + ([] if i["url"] else ["无原文链接"]))
        i["section"] = _section(i)
    data_rows = _m5_rows(brief_date) if cfg["include_m5"] else []
    md = render(brief_date, items, data_rows, meta={"window_label": label})
    ver = (db.q1("SELECT MAX(version) v FROM m2_briefs WHERE brief_date=?", (brief_date,)) or {}).get("v") or 0
    stats = {"window": [start.strftime("%Y-%m-%d %H:%M"), end.strftime("%Y-%m-%d %H:%M")], "candidates": len(rows), "items": len(items),
             "three": sum(1 for i in items if i["tier"] == 3), "two": sum(1 for i in items if i["tier"] == 2), "one": sum(1 for i in items if i["tier"] == 1),
             "blocked": sum(1 for i in items if i["blocked"]), "trigger": trigger, "window_label": label}
    from ..llm.service import assignments
    bid = db.execute("INSERT INTO m2_briefs(brief_date,version,status,items,markdown,stats,config,model,generated_at,generated_by) VALUES(?,?,?,?,?,?,?,?,?,?)",
                     (brief_date, ver + 1, "草稿", db.jdump({"items": items, "data_rows": data_rows}), md, db.jdump(stats), db.jdump(cfg),
                      str(assignments().get("m2.brief") or ""), db.now(), user["name"]))
    db.audit(user, "M2", "生成简报", brief_date, {"brief_id": bid, **stats})
    return bid


def get(bid):
    b = db.q1("SELECT * FROM m2_briefs WHERE id=?", (bid,))
    if not b:
        raise HTTPException(404, "简报不存在")
    p = db.jload(b["items"], {})
    b["items"], b["data_rows"] = p.get("items", []), p.get("data_rows", [])
    b["stats"], b["config"] = db.jload(b["stats"], {}), db.jload(b["config"], {})
    return b


def edit(bid, body, user):
    b = get(bid)
    if b["status"] == "已发布":
        raise HTTPException(409, "已发布的简报不可修改，请重新生成新版本")
    items = b["items"]
    if "items" in body:
        by = {i["m1_id"]: i for i in items}
        for e in body["items"]:
            i = by.get(e.get("m1_id"))
            if not i:
                continue
            changed = False
            for k in ("include", "order"):
                if k in e:
                    i[k] = e[k]
            if "tier" in e and int(e["tier"]) in (1, 2, 3) and int(e["tier"]) != i["tier"]:
                i["tier"], changed = int(e["tier"]), True
            if "points" in e and e["points"] != i.get("points"):
                i["points"], changed = [p for p in e["points"] if p.strip()], True
                src = db.q1("SELECT title,snippet,content FROM m1_items WHERE id=?", (i["m1_id"],)) or {}
                txt = "\n".join(str(src.get(k) or "") for k in ("title", "snippet", "content"))
                i["number_check"] = [{"num": f, "ok": number_in_source(n, u, txt)} for f, n, u in key_numbers(" ".join(i["points"] + i.get("key_numbers", [])))]
                bad = [c["num"] for c in i["number_check"] if not c["ok"]]
                i["blocked"] = bool(bad)
                i["block_reason"] = [f"数字无来源：{'、'.join(bad)}"] if bad else []
            i["section"] = _section(i)
            if changed:  # 反哺 M1：人工校正标记
                db.execute("UPDATE m1_items SET summary=?, manual_corrected=1 WHERE id=?", ("\n".join("- " + p for p in i["points"]), i["m1_id"]))
        md = render(b["brief_date"], items, b["data_rows"], meta={"window_label": (b.get("stats") or {}).get("window_label")})
    else:
        md = b["markdown"]
    if body.get("markdown") is not None and "items" not in body:
        md = body["markdown"]
    db.execute("UPDATE m2_briefs SET items=?, markdown=?, edited_at=?, edited_by=? WHERE id=?",
               (db.jdump({"items": items, "data_rows": b["data_rows"]}), md, db.now(), user["name"], bid))
    db.audit(user, "M2", "编辑简报", bid, {"items": len(body.get("items", [])), "markdown": "markdown" in body})
    return get(bid)


def publish(bid, user):
    b = get(bid)
    body = b["markdown"].split("## 五、")[0]  # 「五、待核实提示」中列出被拦截条目是允许的
    bad = [i["title"] for i in b["items"] if i.get("blocked") and i["title"] in body]
    if bad:
        raise HTTPException(409, f"以下条目的关键数字无来源，已被拦截，不能发布：{'；'.join(bad[:5])}")
    db.execute("UPDATE m2_briefs SET status='已发布', published_at=?, published_by=? WHERE id=?", (db.now(), user["name"], bid))
    db.audit(user, "M2", "发布简报", bid)
    return get(bid)


def to_docx(bid):
    from docx import Document
    from docx.shared import Pt
    b = get(bid)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "仿宋"
    st.font.size = Pt(12)
    lines = b["markdown"].splitlines()
    k = 0
    while k < len(lines):
        ln = lines[k]
        if ln.startswith("# "):
            doc.add_heading(ln[2:], 0)
        elif ln.startswith("## "):
            doc.add_heading(ln[3:], 1)
        elif ln.startswith("|"):
            rows = []
            while k < len(lines) and lines[k].startswith("|"):
                if not re.match(r"^\|[-| ]+\|$", lines[k]):
                    rows.append([c.strip() for c in lines[k].strip("|").split("|")])
                k += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    t.cell(r, c).text = v
            continue
        elif ln.strip():
            doc.add_paragraph(ln.replace("**", ""))
        k += 1
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue(), f"今日行业Brief_{b['brief_date']}_v{b['version']}.docx"


# ---------------- 定时：每日 schedule.time（Asia/Shanghai，容器 TZ） ----------------
def _loop():
    from .. import ops
    while True:
        try:
            ops.beat("m2_brief")
            cfg = C.get()
            sch = cfg.get("schedule") or {}
            now = datetime.now()
            today = now.strftime("%Y-%m-%d")
            if sch.get("enabled") and sch["time"] <= now.strftime("%H:%M") < _plus30(sch["time"]) and db.get_setting("m2_sched_last") != today:
                db.set_setting("m2_sched_last", today)
                ops.run_job("m2_brief", generate, today, SYS, trigger=f"定时 {sch['time']}")
        except Exception:  # noqa: BLE001
            pass
        time.sleep(30)


def _plus30(t):
    h, m = map(int, t.split(":"))
    x = min(h * 60 + m + 30, 23 * 60 + 59)
    return f"{x // 60:02d}:{x % 60:02d}"


def start_scheduler():
    from .. import ops
    if os.environ.get("PH_DISABLE_SCHEDULER") or not ops.acquire_scheduler_lock():
        return
    threading.Thread(target=_loop, daemon=True, name="m2-scheduler").start()
