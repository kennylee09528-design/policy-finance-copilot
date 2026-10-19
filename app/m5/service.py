"""M5 机构看板：快照、看板计算、环比说明、预警、导出。单位：万元。"""
from datetime import date
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .. import db
from ..m4.xlsxformula import norm_name
from . import limits as L
from .ledger import PORTS, ManualUploadSource

DEFAULT_RULES = {"red": 0.95, "over": 1.0, "lag_time": 0.60, "lag_rate": 0.40, "swing": 0.50, "low_hint": 0.50,
                 "direct_entities": ["示例省级融资再担保有限责任公司"]}
TYPE_CN = {"regular": "常规业务", "tech": "科技专项", "private": "民间投资", "t2t": "总对总"}


def rules():
    return {**DEFAULT_RULES, **(db.get_setting("m5_rules", {}) or {})}


# ---------------- 快照 ----------------
def create_snapshot(stat_date, label, limit_set_id, user):
    if not limit_set_id:
        ls = L.latest_limit_set()
        limit_set_id = ls["id"] if ls else None
    sid = db.execute("INSERT INTO m5_snapshots(stat_date,label,limit_set_id,created_at,created_by) VALUES(?,?,?,?,?)",
                     (stat_date, label or f"截止{stat_date}", limit_set_id, db.now(), user["name"]))
    db.audit(user, "M5", "新建快照", sid, {"stat_date": stat_date, "limit_set_id": limit_set_id})
    return sid


def upload_port(sid, port, data, filename, user):
    if port not in PORTS:
        raise HTTPException(400, "端口必须是 regular/tech/private/t2t")
    snap = get_snapshot(sid)
    try:
        led = ManualUploadSource().fetch(port, data=data, filename=filename, direct_entities=rules()["direct_entities"])
    except ValueError as e:
        raise HTTPException(400, f"{PORTS[port]}台账解析失败：{e}")
    if led["period_end"] and snap["stat_date"] and led["period_end"] != snap["stat_date"]:
        led["warnings"].append(f"台账备案截止日 {led['period_end']} 与快照统计时点 {snap['stat_date']} 不一致")
    with db.tx() as c:
        c.execute("DELETE FROM m5_ledgers WHERE snapshot_id=? AND port=?", (sid, port))
        c.execute("INSERT INTO m5_ledgers(snapshot_id,port,file_name,period_start,period_end,updated_on,total,direct,direct_source,rows,warnings,uploaded_at,uploaded_by) "
                  "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  (sid, port, filename, led["period_start"], led["period_end"], led["updated_on"], led["total"], led["direct"],
                   led["direct_source"], db.jdump(led["rows"] + led["direct_rows"]), db.jdump(led["warnings"]), db.now(), user["name"]))
    db.audit(user, "M5", f"上传{PORTS[port]}台账", sid, {"file": filename, "total": led["total"], "period_end": led["period_end"]})
    return led


def get_snapshot(sid):
    s = db.q1("SELECT * FROM m5_snapshots WHERE id=?", (sid,))
    if not s:
        raise HTTPException(404, "快照不存在")
    s["notes"] = db.jload(s["notes"], {})
    s["manual_direct"] = db.jload(s["manual_direct"], {})
    s["ledgers"] = {}
    for l in db.q("SELECT * FROM m5_ledgers WHERE snapshot_id=?", (sid,)):
        l["rows"] = db.jload(l["rows"], [])
        l["warnings"] = db.jload(l["warnings"], [])
        s["ledgers"][l["port"]] = l
    return s


def list_snapshots():
    out = db.q("SELECT id,stat_date,label,limit_set_id,created_at,created_by FROM m5_snapshots ORDER BY stat_date DESC, id DESC")
    for s in out:
        s["ports"] = [r["port"] for r in db.q("SELECT port FROM m5_ledgers WHERE snapshot_id=?", (s["id"],))]
    return out


def set_manual_direct(sid, port, value, user, reason):
    if not reason:
        raise HTTPException(400, "手工填写直保金额必须注明来源/原因")
    s = get_snapshot(sid)
    md = s["manual_direct"]
    if value is None:
        md.pop(port, None)
    else:
        md[port] = {"value": float(value), "reason": reason, "by": user["name"], "at": db.now()}
    db.execute("UPDATE m5_snapshots SET manual_direct=? WHERE id=?", (db.jdump(md), sid))
    db.audit(user, "M5", "手工直保金额", sid, {"port": port, "value": value, "reason": reason})


def save_notes(sid, notes, user):
    db.execute("UPDATE m5_snapshots SET notes=? WHERE id=?", (db.jdump(notes), sid))
    db.audit(user, "M5", "编辑环比说明", sid, notes)


def set_limit_set(sid, lid, user):
    L.get_limit_set(lid)
    db.execute("UPDATE m5_snapshots SET limit_set_id=? WHERE id=?", (lid, sid))
    db.audit(user, "M5", "切换额度表", sid, {"limit_set_id": lid})


# ---------------- 计算 ----------------
def _port_figures(s):
    out = {}
    for p in PORTS:
        l = s["ledgers"].get(p)
        if not l:
            out[p] = None
            continue
        direct, src = l["direct"], l["direct_source"]
        md = s["manual_direct"].get(p)
        if direct is None and md:
            direct, src = md["value"], f"手工填写（{md['reason']}）"
        inst = {}
        for r in l["rows"]:
            if not r.get("is_direct_entity"):
                inst[r["name"]] = inst.get(r["name"], 0.0) + r["amount"]
        out[p] = {"total": l["total"], "direct": direct, "direct_source": src,
                  "reguar": l["total"] - (direct or 0.0), "inst": inst,
                  "period_end": l["period_end"], "file": l["file_name"], "warnings": l["warnings"]}
    return out


def _prev_snapshot(s):
    r = db.q1("SELECT id FROM m5_snapshots WHERE deleted_at IS NULL AND (stat_date < ? OR (stat_date = ? AND id < ?)) ORDER BY stat_date DESC, id DESC LIMIT 1",
              (s["stat_date"], s["stat_date"], s["id"]))
    return get_snapshot(r["id"]) if r else None


def _yi(v):
    return f"{abs(v) / 10000:.2f}亿"


def _chg(v):
    return ("增加" if v >= 0 else "减少") + _yi(v)


def auto_notes(cur, prev, prev_date):
    notes = {}
    for p, cn in TYPE_CN.items():
        c = cur.get(p)
        pv = prev.get(p) if prev else None
        if not c:
            notes[p] = f"{cn}：本期未导入台账。"
            continue
        if not pv:
            notes[p] = f"{cn}：暂无上一统计时点快照，无法计算环比（本期累计备案 {c['total']:,.2f} 万元）。"
            continue
        d = c["total"] - pv["total"]
        head = f"比上一个统计时点({prev_date.month}月{prev_date.day}日)，" if p == "regular" else ""
        if c["direct"] is not None and pv["direct"] is not None:
            notes[p] = f"{head}{cn}{_chg(d)}，其中，直保{_chg(c['direct'] - pv['direct'])}，再保{_chg(c['reguar'] - pv['reguar'])}。"
        else:
            notes[p] = f"{head}{cn}{_chg(d)}（直保数据缺失，未拆分直保/再保）。"
    return notes


def _rate(a, b):
    return a / b if (a is not None and b) else None


def dashboard(sid, user=None):
    s = get_snapshot(sid)
    R = rules()
    figs = _port_figures(s)
    prev = _prev_snapshot(s)
    prev_figs = _port_figures(prev) if prev else None
    prev2 = _prev_snapshot(prev) if prev else None
    prev2_figs = _port_figures(prev2) if prev2 else None
    ls = L.get_limit_set(s["limit_set_id"]) if s["limit_set_id"] else None
    alerts = []
    if not ls:
        alerts.append({"level": "error", "msg": "未关联授信额度表（请从 M4 方案带入或导入额度表）"})
    insts = ls["data"]["institutions"] if ls else []
    ov = ls["data"]["overall"] if ls else {"guodan": {}, "direct_alloc": {}, "reguar_alloc": {}, "unalloc_const": {}}
    sd = date.fromisoformat(s["stat_date"])
    days = (date(sd.year, 12, 31) - date(sd.year, 1, 1)).days + 1
    time_prog = sd.timetuple().tm_yday / days

    def g(p, k):
        f = figs.get(p)
        return f[k] if f else None

    # ---------- 汇总页 ----------
    summ = {}
    for p in PORTS:
        f = figs.get(p)
        pf = prev_figs.get(p) if prev_figs else None
        guodan = (ov.get("guodan") or {}).get(p)
        blk = {"guodan": guodan, "loaded": f is not None, "direct_source": f["direct_source"] if f else None,
               "period_end": f["period_end"] if f else None}
        for k in ("direct", "reguar", "total"):
            amt = f[k] if f else None
            pamt = pf[k] if pf else None
            blk[k] = {"amount": amt, "delta": (amt - pamt) if (amt is not None and pamt is not None) else None}
        if p in ("regular", "private"):
            da = (ov.get("direct_alloc") or {}).get(p)
            ra = (ov.get("reguar_alloc") or {}).get(p)
            key = "regular" if p == "regular" else "private"
            inst_sum = sum((x.get(key) or 0) for x in insts)
            ra = ra if ra is not None else inst_sum
            blk["direct"]["alloc"], blk["reguar"]["alloc"] = da, ra
            blk["total"]["alloc"] = (da or 0) + (ra or 0) if (da is not None or ra is not None) else None
            for k in ("direct", "reguar"):
                a, al = blk[k]["amount"], blk[k]["alloc"]
                blk[k]["remain"] = (al - a) if (al is not None and a is not None) else None
            uc = (ov.get("unalloc_const") or {}).get(p)
            blk["unalloc"] = uc if uc is not None else ((guodan - blk["total"]["alloc"]) if (guodan is not None and blk["total"]["alloc"] is not None) else None)
        tot = blk["total"]["amount"]
        blk["total"]["remain"] = (guodan - tot) if (guodan is not None and tot is not None) else None
        if p in ("regular", "t2t"):  # Q3：使用率统一以 2026 年国担额度为分母
            for k in ("direct", "reguar", "total"):
                blk[k]["rate"] = _rate(blk[k]["amount"], guodan)
            if p == "regular":
                for k in ("direct", "reguar", "total"):
                    blk[k]["rate_alloc"] = _rate(blk[k]["amount"], blk[k].get("alloc"))
        summ[p] = blk
    summ["regular"]["need_adjust"] = ov.get("need_adjust", "")
    tr = summ["regular"]["total"].get("rate")
    if tr is not None:
        if tr >= R["over"]:
            alerts.append({"level": "error", "msg": f"常规业务合计使用率 {tr:.2%} ≥ {R['over']:.0%}：超额告警"})
        elif tr >= R["red"]:
            alerts.append({"level": "error", "msg": f"常规业务合计使用率 {tr:.2%} ≥ {R['red']:.0%}：额度即将用尽"})
    if summ["regular"].get("unalloc") is not None and summ["regular"]["unalloc"] < 0:
        alerts.append({"level": "warn", "msg": f"常规额度分配合计超出国担额度 {-summ['regular']['unalloc']:,.2f} 万元（剩余未分配为负）"})
    for p in PORTS:
        f = figs.get(p)
        if f is None:
            alerts.append({"level": "warn", "msg": f"{PORTS[p]}台账未导入"})
        else:
            for w in f["warnings"]:
                alerts.append({"level": "warn", "msg": f"{PORTS[p]}：{w}"})

    # ---------- 分机构 ----------
    order = [x["name"] for x in insts]
    lim = {x["name"]: x for x in insts}
    names = list(order)
    master = {norm_name(r["full_name"]) for r in db.q("SELECT full_name FROM institutions")}
    for p in PORTS:
        for n in (figs[p]["inst"] if figs.get(p) else {}):
            if n not in lim and n not in names:
                names.append(n)
                alerts.append({"level": "warn", "msg": f"「{n}」出现在{PORTS[p]}台账，但不在授信额度表中（名单不匹配）"})
    for n in names:
        if master and norm_name(n) not in master:
            alerts.append({"level": "info", "msg": f"「{n}」不在机构主数据中"})
    reg_tot = summ["regular"]["reguar"]["amount"] if figs.get("regular") else None
    pri_tot = sum(figs["private"]["inst"].values()) if figs.get("private") else None
    rows = []
    for i, n in enumerate(names, 1):
        lm = lim.get(n, {})
        a = {p: (figs[p]["inst"].get(n, 0.0) if figs.get(p) else None) for p in PORTS}
        pa = {p: (prev_figs[p]["inst"].get(n, 0.0) if prev_figs and prev_figs.get(p) else None) for p in PORTS}
        pa2 = {p: (prev2_figs[p]["inst"].get(n, 0.0) if prev2_figs and prev2_figs.get(p) else None) for p in PORTS}
        row = {"seq": i, "name": n, "in_limits": n in lim,
               "regular_limit": lm.get("regular"), "regular_amt": a["regular"],
               "regular_rate": _rate(a["regular"], lm.get("regular")),
               "tech_amt": a["tech"],
               "private_limit": lm.get("private"), "private_amt": a["private"],
               "private_rate": _rate(a["private"], lm.get("private")),
               "private_note": lm.get("private_note") if not lm.get("private") else None,
               "t2t_amt": a["t2t"],
               "total_amt": sum(v for v in a.values() if v is not None),
               "regular_share": _rate(a["regular"], reg_tot), "private_share": _rate(a["private"], pri_tot),
               "regular_remain": (lm.get("regular") - a["regular"]) if (lm.get("regular") is not None and a["regular"] is not None) else None,
               "delta": {p: (a[p] - pa[p]) if (a[p] is not None and pa[p] is not None) else None for p in PORTS},
               "warn": []}
        rr = row["regular_rate"]
        if rr is not None:
            if rr >= R["over"]:
                row["warn"].append({"level": "over", "msg": f"常规使用率 {rr:.2%} 超额"})
            elif rr >= R["red"]:
                row["warn"].append({"level": "red", "msg": f"常规使用率 {rr:.2%} ≥{R['red']:.0%}，额度即将用尽"})
            if time_prog >= R["lag_time"] and rr < R["lag_rate"]:
                row["warn"].append({"level": "yellow", "msg": f"时间进度 {time_prog:.0%} 而使用率仅 {rr:.2%}：进度滞后"})
        for p in PORTS:
            cur_new = row["delta"][p]
            prev_new = (pa[p] - pa2[p]) if (pa[p] is not None and pa2[p] is not None) else None
            if cur_new is not None and prev_new:
                sw = (cur_new - prev_new) / abs(prev_new)
                if abs(sw) > R["swing"]:
                    row["warn"].append({"level": "check", "msg": f"{PORTS[p]}新增备案环比波动 {sw:+.0%}，请核查"})
        rows.append(row)
    if user:
        from ..auth import row_filter
        rows = row_filter(user, rows)

    notes_auto = auto_notes(figs, prev_figs, date.fromisoformat(prev["stat_date"]) if prev else None)
    notes = {p: (s["notes"].get(p) or notes_auto[p]) for p in PORTS}
    return {"snapshot": {k: s[k] for k in ("id", "stat_date", "label", "limit_set_id", "created_at", "created_by")},
            "limit_set": {k: ls[k] for k in ("id", "name", "source")} if ls else None,
            "limit_warnings": ls["data"].get("warnings", []) if ls else [],
            "prev": {"id": prev["id"], "stat_date": prev["stat_date"]} if prev else None,
            "time_progress": time_prog, "summary": summ, "rows": rows, "notes": notes, "notes_auto": notes_auto,
            "notes_edited": {p: bool(s["notes"].get(p)) for p in PORTS},
            "ports": {p: ({k: figs[p][k] for k in ("total", "direct", "direct_source", "reguar", "period_end", "file")} if figs.get(p) else None) for p in PORTS},
            "totals": {"regular_limit": sum((r["regular_limit"] or 0) for r in rows),
                       "regular_amt": sum((r["regular_amt"] or 0) for r in rows),
                       "tech_amt": sum((r["tech_amt"] or 0) for r in rows),
                       "private_limit": sum((r["private_limit"] or 0) for r in rows),
                       "private_amt": sum((r["private_amt"] or 0) for r in rows),
                       "t2t_amt": sum((r["t2t_amt"] or 0) for r in rows),
                       "total_amt": sum((r["total_amt"] or 0) for r in rows)},
            "rules": R, "alerts": alerts}


# ---------------- 导出 ----------------
THIN = Side(style="thin", color="999999")
B = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HF = PatternFill("solid", fgColor="F3E3E4")
M2 = "#,##0.00"
PCT = "0.00%"


def export_xlsx(sid):
    d = dashboard(sid)
    s = d["summary"]
    sd = d["snapshot"]["stat_date"]
    md = f"{int(sd[5:7])}月{int(sd[8:10])}日"
    wb = Workbook()
    ws = wb.active
    ws.title = "业务整体情况"
    ws.merge_cells("A1:A2")
    ws["A1"] = "项目"
    groups = [("B", "E", "常规额度"), ("F", "I", "科技专项额度"), ("J", "M", "民间专项额度"), ("N", "P", "总对总额度")]
    for a, b, t in groups:
        ws.merge_cells(f"{a}1:{b}1")
        ws[f"{a}1"] = t
    sub = [f"截止{md}备案金额", "额度分配", "使用率(÷国担额度)", "剩余额度",
           f"截止{md}备案金额", "额度分配", "较上期变化", "剩余额度",
           f"截止{md}备案金额", "额度分配", "较上期变化", "剩余额度",
           f"截止{md}备案金额", "使用率(÷国担额度)", "剩余额度"]
    for i, t in enumerate(sub):
        ws.cell(2, 2 + i, t)
    labels = ["直保", "再保", "合计", "2026年国担额度", "剩余未分配额度", "需调整"]
    for i, t in enumerate(labels):
        ws.cell(3 + i, 1, t)
    R, T, P, X = s["regular"], s["tech"], s["private"], s["t2t"]
    for i, k in enumerate(("direct", "reguar", "total")):
        r = 3 + i
        ws.cell(r, 2, R[k]["amount"]); ws.cell(r, 3, R[k].get("alloc")); ws.cell(r, 4, R[k].get("rate")); ws.cell(r, 5, R[k].get("remain"))
        ws.cell(r, 6, T[k]["amount"]); ws.cell(r, 7, "无上限额度，先到先得" if k != "total" else None); ws.cell(r, 8, T[k]["delta"])
        ws.cell(r, 10, P[k]["amount"]); ws.cell(r, 11, P[k].get("alloc")); ws.cell(r, 12, P[k]["delta"]); ws.cell(r, 13, P[k].get("remain"))
        ws.cell(r, 14, X[k]["amount"]); ws.cell(r, 15, X[k].get("rate"))
    ws.cell(5, 9, T["total"]["remain"]); ws.cell(5, 16, None)
    ws.cell(6, 2, R["guodan"]); ws.cell(6, 6, T["guodan"]); ws.cell(6, 10, P["guodan"]); ws.cell(6, 14, X["guodan"]); ws.cell(6, 16, X["total"]["remain"])
    ws.cell(7, 3, R.get("unalloc")); ws.cell(7, 10, P.get("unalloc"))
    ws.cell(8, 3, R.get("need_adjust") or None)
    ws.cell(9, 1, "注：使用率统一以「2026年国担额度」为分母（需求 Q3）；科技专项、民间专项不计算使用率（Q4），仅展示累计备案与较上期变化；单位：万元。")
    n = d["notes"]
    ws.merge_cells("A11:E11"); ws["A11"] = "变化情况:" + n["regular"]
    ws.merge_cells("F11:I11"); ws["F11"] = n["tech"]
    ws.merge_cells("J11:M11"); ws["J11"] = n["private"]
    ws.merge_cells("N11:P11"); ws["N11"] = n["t2t"]
    for row in ws.iter_rows(min_row=1, max_row=8, max_col=16):
        for c in row:
            c.border = B
            if c.row <= 2:
                c.font = Font(bold=True); c.fill = HF; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            elif isinstance(c.value, float):
                c.number_format = PCT if (c.column in (4, 15) and c.row <= 5) else M2
    for row in ws["A11:P11"]:
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[11].height = 60
    ws.column_dimensions["A"].width = 16
    for col in "BCDEFGHIJKLMNOP":
        ws.column_dimensions[col].width = 15

    ws2 = wb.create_sheet("各机构明细")
    heads = ["序号", "机构", "常规业务授信额度（万元）", f"截止{md}常规业务备案金额", "额度使用比例", "科技创新计划额度（万元）",
             f"截止{md}科技创新计划备案金额", "额度使用比例", "民间专项计划额度（万元）", f"截止{md}民间专项备案金额", "额度使用比例",
             f"截止{md}总对总备案金额", "机构合计备案金额", "常规备案占再保总备案比例", "民投备案占再保总备案比例"]
    ws2.append(heads)
    for r in d["rows"]:
        ws2.append([r["seq"], r["name"], r["regular_limit"], r["regular_amt"], r["regular_rate"], "无上限额度，先到先得" if r["seq"] == 1 else None,
                    r["tech_amt"], None, r["private_limit"], r["private_amt"], r["private_rate"] if r["private_rate"] is not None else (r["private_note"] or None),
                    r["t2t_amt"], r["total_amt"], r["regular_share"], r["private_share"]])
    t = d["totals"]
    ws2.append([None, "合计", t["regular_limit"], t["regular_amt"], _rate(t["regular_amt"], t["regular_limit"]), None, t["tech_amt"], None,
                t["private_limit"], t["private_amt"], _rate(t["private_amt"], t["private_limit"]), t["t2t_amt"], t["total_amt"], 1, 1])
    for row in ws2.iter_rows():
        for c in row:
            c.border = B
            if c.row == 1:
                c.font = Font(bold=True); c.fill = HF; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            elif isinstance(c.value, (int, float)) and c.column > 2:
                c.number_format = PCT if c.column in (5, 8, 11, 14, 15) else M2
    ws2.column_dimensions["B"].width = 40
    for col in "CDEFGHIJKLMNO":
        ws2.column_dimensions[col].width = 15
    ws2.freeze_panes = "C2"
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue(), f"业务额度使用情况_截止{sd}.xlsx"
