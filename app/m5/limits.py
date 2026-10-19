"""M5 授信额度来源：① M4 方案结果；② 手工导入额度表（《业务额度使用情况》同格式）。"""
import re
import warnings
from io import BytesIO

import openpyxl
from fastapi import HTTPException

from .. import db
from ..m4.xlsxformula import ftext, norm_name

warnings.filterwarnings("ignore", message="Workbook contains no default style")

TYPE_HEAD = {"regular": "常规", "tech": "科技", "private": "民间", "t2t": "总对总"}
ROW_LABELS = {"直保": "direct", "再保": "reguar", "合计": "total", "2026年国担额度": "guodan", "剩余未分配额度": "unalloc",
              "需调整": "need_adjust"}


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_limit_workbook(data: bytes):
    wf = openpyxl.load_workbook(BytesIO(data))
    wv = openpyxl.load_workbook(BytesIO(data), data_only=True)
    warn = []
    # ---- 分机构表：含「机构」与「常规业务授信额度」表头 ----
    inst_ws = None
    for ws in wv.worksheets:
        for row in ws.iter_rows(max_row=5):
            if any(isinstance(c.value, str) and "常规业务授信额度" in c.value for c in row):
                inst_ws = ws
                break
        if inst_ws:
            break
    if inst_ws is None:
        raise HTTPException(400, "未找到含「常规业务授信额度」表头的分机构额度表")
    hdr = None
    for row in inst_ws.iter_rows(max_row=5):
        vals = [str(c.value or "") for c in row]
        if any("常规业务授信额度" in v for v in vals):
            hdr = vals
            break
    name_c = next(i for i, v in enumerate(hdr) if v.strip() in ("机构", "机构名称"))
    reg_c = next(i for i, v in enumerate(hdr) if "常规业务授信额度" in v)
    pri_c = next((i for i, v in enumerate(hdr) if "民间专项计划额度" in v or "民间专项额度" in v), None)
    insts = []
    for row in inst_ws.iter_rows(min_row=3):
        name = row[name_c].value
        if not name or str(name).strip() in ("合计", "总计"):
            continue
        reg = _num(row[reg_c].value)
        pri = _num(row[pri_c].value) if pri_c is not None else None
        note = None
        if pri_c is not None and pri_c + 2 < len(row) and isinstance(row[pri_c + 2].value, str):
            note = row[pri_c + 2].value  # 如“未申报额度”
        insts.append({"name": str(name).strip(), "regular": reg, "private": pri, "private_note": note})

    # ---- 汇总表：A 列行标签 + 第 1 行业务类型分组 ----
    overall = {"guodan": {}, "direct_alloc": {}, "reguar_alloc": {}, "need_adjust": "", "unalloc_const": {}}
    sum_ws = None
    for ws in wv.worksheets:
        if any(isinstance(c.value, str) and c.value.strip() == "2026年国担额度" for c in ws["A"]):
            sum_ws = ws
            break
    if sum_ws is not None:
        fws = wf[sum_ws.title]
        groups = {}
        for c in sum_ws[1]:
            if isinstance(c.value, str):
                for k, kw in TYPE_HEAD.items():
                    if kw in c.value and k not in groups:
                        groups[k] = c.column
        rowmap = {}
        for c in sum_ws["A"]:
            if isinstance(c.value, str) and c.value.strip() in ROW_LABELS:
                rowmap[ROW_LABELS[c.value.strip()]] = c.row
        sumreg = sum(x["regular"] or 0 for x in insts)
        sumpri = sum(x["private"] or 0 for x in insts)
        for k, col in groups.items():
            gv = _num(sum_ws.cell(rowmap["guodan"], col).value)
            overall["guodan"][k] = gv
            if k in ("regular", "private"):
                overall["direct_alloc"][k] = _num(sum_ws.cell(rowmap["direct"], col + 1).value)
                cell_f = fws.cell(rowmap["reguar"], col + 1)
                f = ftext(cell_f.value)
                if f is None and _num(cell_f.value) is not None:
                    const = _num(cell_f.value)
                    overall["reguar_alloc"][k] = const
                    s = sumreg if k == "regular" else sumpri
                    if abs(const - s) > 0.01:
                        warn.append(f"{TYPE_HEAD[k]}再保额度分配为手填常数 {const:,.2f}，与分机构额度合计 {s:,.2f} 不一致（相差 {s - const:,.2f}），按手填常数采用")
                else:
                    overall["reguar_alloc"][k] = None  # 公式引用分机构合计 → 按合计计算
            if "unalloc" in rowmap:
                ucell = fws.cell(rowmap["unalloc"], col + 1)
                if ftext(ucell.value) is None and _num(ucell.value) is not None:
                    overall["unalloc_const"][k] = _num(ucell.value)
        if "need_adjust" in rowmap:
            for col in range(2, sum_ws.max_column + 1):
                v = sum_ws.cell(rowmap["need_adjust"], col).value
                if v not in (None, ""):
                    overall["need_adjust"] = str(v)
                    break
    else:
        warn.append("未找到汇总表（含“2026年国担额度”行），国担额度等需手工填写")
    return {"institutions": insts, "overall": overall, "warnings": warn}


def save_limit_set(name, source, data, user):
    lid = db.execute("INSERT INTO m5_limit_sets(name,source,created_at,created_by,data) VALUES(?,?,?,?,?)",
                     (name, source, db.now(), user["name"], db.jdump(data)))
    db.audit(user, "M5", "新建额度表", name, {"source": source, "institutions": len(data["institutions"])})
    return lid


def from_m4_scheme(sid, user, base_limit_id=None):
    from ..m4 import service as m4
    s = m4.get_scheme(sid)
    insts = [{"name": r["name"], "regular": r["regular_final"], "private": r["mintou_special"] or 0.0,
              "private_note": None if r["mintou_special"] else "未申报额度"} for r in s["result"]["rows"]]
    base = get_limit_set(base_limit_id) if base_limit_id else latest_limit_set()
    overall = base["data"]["overall"] if base else {"guodan": {}, "direct_alloc": {}, "reguar_alloc": {}, "need_adjust": "", "unalloc_const": {}}
    overall = {**overall, "reguar_alloc": {}}  # 再保额度分配按 M4 分机构额度合计
    data = {"institutions": insts, "overall": overall,
            "warnings": [f"分机构额度取自 M4 方案 v{s['version']}「{s['name']}」（常规=常规授信额度，民间=民间投资专项）"] +
                        ([f"国担额度/直保额度分配沿用额度表「{base['name']}」"] if base else ["国担额度、直保额度分配未设置，请在额度表中填写"])}
    return save_limit_set(f"M4 方案 v{s['version']} 额度", f"M4方案#{sid}", data, user)


def get_limit_set(lid):
    r = db.q1("SELECT * FROM m5_limit_sets WHERE id=?", (lid,))
    if not r:
        raise HTTPException(404, "额度表不存在")
    r["data"] = db.jload(r["data"])
    return r


def latest_limit_set():
    r = db.q1("SELECT id FROM m5_limit_sets ORDER BY id DESC LIMIT 1")
    return get_limit_set(r["id"]) if r else None


def list_limit_sets():
    return db.q("SELECT id,name,source,created_at,created_by FROM m5_limit_sets ORDER BY id DESC")


def update_limit_set(lid, data, user, reason):
    if not reason:
        raise HTTPException(400, "修改额度必须填写原因")
    old = get_limit_set(lid)
    nid = save_limit_set(f"{old['name']}（修订）", f"修订自#{lid}", data, user)
    db.audit(user, "M5", "修订额度表", lid, {"new": nid, "reason": reason})
    return nid
