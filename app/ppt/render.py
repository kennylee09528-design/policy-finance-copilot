"""python-pptx 渲染：默认工作台深红主题（16:9），或管理员上传的 .pptx 模板（使用其母版与版式）。
图表为原生可编辑图表（柱/条/堆积/折线），表格为原生表格；所有数字来自数据包（代码从数据库取值）。"""
import os
import re
import shutil
import subprocess
import tempfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .data import NUMTOK, allowed_tokens

RED, RED_D, GOLD, INK, SUB, BG, WHITE = (RGBColor(0xB0, 0x24, 0x2E), RGBColor(0x7A, 0x14, 0x1C), RGBColor(0xC9, 0xA2, 0x27),
                                          RGBColor(0x22, 0x22, 0x22), RGBColor(0x66, 0x66, 0x66), RGBColor(0xF7, 0xF3, 0xF0), RGBColor(0xFF, 0xFF, 0xFF))
SERIES = [RED, GOLD, RGBColor(0x8C, 0x8C, 0x8C), RGBColor(0x2E, 0x7D, 0x7A), RGBColor(0x5B, 0x6E, 0xA8)]
CHART_TYPES = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "column_stacked": XL_CHART_TYPE.COLUMN_STACKED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
               "bar_stacked": XL_CHART_TYPE.BAR_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS}
FONT = "微软雅黑"
TABLE_ROWS, CHART_CATS, TEXT_ITEMS = 12, 18, 8


def _font(run, size, color=INK, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT


def _box(slide, x, y, w, h, text, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        _font(r, size, color, bold)
    return tb


def _rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


class Deck:
    def __init__(self, template_path=None):
        self.tpl = template_path
        if template_path:
            self.prs = Presentation(template_path)
            sl = self.prs.slides._sldIdLst  # noqa: SLF001  清空模板中的示例页，仅保留母版与版式
            for sid in list(sl):
                self.prs.part.drop_rel(sid.rId)
                sl.remove(sid)
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.W, self.H = self.prs.slide_width, self.prs.slide_height
        self.footers = []

    # ---------- 版式 ----------
    def _layout(self, kind):
        names = {"cover": ["title slide", "标题幻灯片", "封面"], "section": ["section", "节标题", "章节"],
                 "content": ["title and content", "标题和内容", "内容"], "title_only": ["title only", "仅标题"], "closing": ["title slide", "标题幻灯片", "封底", "结束"]}
        for l in self.prs.slide_layouts:
            if any(n in l.name.lower() for n in names.get(kind, [])):
                return l
        idx = {"cover": 0, "section": 2, "content": 1, "title_only": 5, "closing": 0}.get(kind, 6)
        return self.prs.slide_layouts[min(idx, len(self.prs.slide_layouts) - 1)]

    def _tpl_slide(self, kind, title, subtitle=None):
        s = self.prs.slides.add_slide(self._layout(kind))
        for ph in list(s.placeholders):
            t = ph.placeholder_format.type
            is_sub = "SUBTITLE" in str(t)  # 注意 SUBTITLE 也含 "TITLE"，须先判断，否则副标题位被写成标题（封面标题重复）
            if not is_sub and (ph.placeholder_format.idx == 0 or "TITLE" in str(t)):
                ph.text = title
            elif subtitle is not None and (is_sub or ph.placeholder_format.idx == 1) and kind in ("cover", "closing", "section"):
                ph.text = subtitle
                subtitle = None
            else:
                ph._element.getparent().remove(ph._element)  # noqa: SLF001  去掉未用占位符
        return s

    def _area(self):
        return Inches(0.6), Inches(1.35), self.W - Inches(1.2), self.H - Inches(2.0)

    def content_slide(self, title):
        if self.tpl:
            s = self._tpl_slide("title_only", title)
        else:
            s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            _rect(s, 0, 0, self.W, self.H, BG)
            _rect(s, 0, 0, self.W, Inches(1.05), RED)
            _rect(s, 0, Inches(1.05), self.W, Inches(0.06), GOLD)
            _box(s, Inches(0.6), Inches(0.22), self.W - Inches(1.2), Inches(0.7), title, 26, WHITE, True)
        self.footers.append(s)
        return s

    # ---------- 页面 ----------
    def cover(self, title, subtitle, org, date_text):
        if self.tpl:
            self._tpl_slide("cover", title, f"{subtitle}\n{org}")
            return
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        _rect(s, 0, 0, self.W, self.H, RED_D)
        _rect(s, 0, Inches(4.6), self.W, Inches(0.08), GOLD)
        _box(s, Inches(0.9), Inches(2.0), self.W - Inches(1.8), Inches(1.6), title, 40, WHITE, True)
        _box(s, Inches(0.9), Inches(3.6), self.W - Inches(1.8), Inches(0.8), subtitle, 20, GOLD)
        _box(s, Inches(0.9), Inches(5.0), self.W - Inches(1.8), Inches(1.0), [org, date_text], 16, WHITE)
        _box(s, Inches(0.9), Inches(6.6), Inches(6), Inches(0.4), "内部资料 · 注意保密", 12, GOLD)

    def agenda(self, title, items):
        s = self.content_slide(title)
        x, y, w, h = self._area()
        for i, it in enumerate(items[:8]):
            yy = y + Inches(0.15) + i * Inches(0.62)
            _rect(s, x, yy, Inches(0.8), Inches(0.5), RED)
            _box(s, x, yy + Inches(0.05), Inches(0.8), Inches(0.45), f"{i + 1:02d}", 18, WHITE, True, PP_ALIGN.CENTER)
            _box(s, x + Inches(1.0), yy + Inches(0.05), w - Inches(1.0), Inches(0.5), it, 20, INK, True)

    def section(self, no, title):
        if self.tpl:
            self._tpl_slide("section", title, f"第 {no:02d} 部分")
            return
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        _rect(s, 0, 0, self.W, self.H, BG)
        _rect(s, 0, 0, Inches(4.2), self.H, RED)
        _box(s, Inches(0.6), Inches(2.6), Inches(3.2), Inches(1.5), f"{no:02d}", 72, WHITE, True)
        _box(s, Inches(4.8), Inches(3.0), self.W - Inches(5.4), Inches(1.2), title, 34, RED_D, True)
        _rect(s, Inches(4.8), Inches(4.2), Inches(1.5), Inches(0.08), GOLD)

    def bullets(self, title, bullets, facts=None):
        s = self.content_slide(title)
        x, y, w, h = self._area()
        bw = w if not facts else int(w * 0.6)
        tb = s.shapes.add_textbox(x, y, bw, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets or ["（无）"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            r = p.add_run()
            r.text = "■  "
            _font(r, 14, RED, True)
            r = p.add_run()
            r.text = b
            _font(r, 18)
        if facts:
            self._cards(s, x + bw + Inches(0.3), y, w - bw - Inches(0.3), facts, cols=1)

    def _cards(self, s, x, y, w, facts, cols=3):
        cw = int((w - Inches(0.25) * (cols - 1)) / cols)
        ch = Inches(1.25)
        for i, f in enumerate(facts):
            cx = x + (i % cols) * (cw + Inches(0.25))
            cy = y + (i // cols) * (ch + Inches(0.2))
            card = _rect(s, cx, cy, cw, ch, WHITE)
            card.line.color.rgb = RGBColor(0xE6, 0xD8, 0xD0)
            _rect(s, cx, cy, Inches(0.08), ch, RED)
            _box(s, cx + Inches(0.2), cy + Inches(0.1), cw - Inches(0.3), Inches(0.45), f["name"], 12, SUB)
            _box(s, cx + Inches(0.2), cy + Inches(0.52), cw - Inches(0.3), Inches(0.65), f["value"], 24, RED_D, True)

    def kpi(self, title, facts, bullets=None):
        s = self.content_slide(title)
        x, y, w, h = self._area()
        self._cards(s, x, y, w, facts[:9], cols=3)
        if bullets:
            rows = (min(len(facts), 9) + 2) // 3
            _box(s, x, y + rows * Inches(1.45) + Inches(0.1), w, Inches(1.2), ["■ " + b for b in bullets], 14, INK)

    def chart(self, title, c, part=None):
        s = self.content_slide(title + (f"（{part}）" if part else ""))
        x, y, w, h = self._area()
        cd = CategoryChartData()
        cd.categories = c["categories"]
        for se in c["series"]:
            cd.add_series(se["name"], se["values"])
        gf = s.shapes.add_chart(CHART_TYPES.get(c["type"], XL_CHART_TYPE.COLUMN_CLUSTERED), x, y, w, h, cd)
        ch = gf.chart
        ch.font.size = Pt(11)
        ch.font.name = FONT
        ch.has_title = False
        ch.has_legend = len(c["series"]) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
        plot = ch.plots[0]
        plot.has_data_labels = len(c["categories"]) <= 20 and len(c["series"]) <= 2 and "stacked" not in c["type"]
        if plot.has_data_labels:
            plot.data_labels.number_format = c.get("fmt") or "#,##0.00"
            plot.data_labels.number_format_is_linked = False
            plot.data_labels.font.size = Pt(9)
        try:
            ch.value_axis.tick_labels.number_format = "0%" if "%" in (c.get("fmt") or "") else "#,##0"
            ch.value_axis.tick_labels.number_format_is_linked = False
            ch.value_axis.has_major_gridlines = True
        except (ValueError, AttributeError):
            pass
        if c["type"].startswith("bar"):
            ch.category_axis.reverse_order = True
        for i, se in enumerate(plot.series):
            fill = se.format.fill
            fill.solid()
            fill.fore_color.rgb = SERIES[i % len(SERIES)]
            if c["type"] == "line":
                se.format.line.color.rgb = SERIES[i % len(SERIES)]

    def table(self, title, t, part=None):
        s = self.content_slide(title + (f"（{part}）" if part else ""))
        x, y, w, h = self._area()
        rows = t["rows"] or [["（无数据）"] + [""] * (len(t["columns"]) - 1)]
        n = len(rows) + 1
        gs = s.shapes.add_table(n, len(t["columns"]), x, y, w, min(h, Inches(0.42) * n))
        tb = gs.table
        lens = [max([len(str(c)) for c in [t["columns"][j]] + [r[j] for r in rows]] + [2]) for j in range(len(t["columns"]))]
        lens = [min(max(L, 4), 40) for L in lens]
        tot = sum(lens)
        for j, L in enumerate(lens):
            tb.columns[j].width = Emu(int(w * L / tot))
        small = len(t["columns"]) > 6 or any(len(str(c)) > 40 for r in rows for c in r)
        for j, col in enumerate(t["columns"]):
            cell = tb.cell(0, j)
            cell.text = col
            cell.fill.solid()
            cell.fill.fore_color.rgb = RED
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    _font(r, 11, WHITE, True)
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                cell = tb.cell(i, j)
                cell.text = str(v)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xF6, 0xEE, 0xEA)
                for p in cell.text_frame.paragraphs:
                    if NUMTOK.fullmatch(str(v).replace("%", "").replace("+", "").replace("-", "").strip() or "x"):
                        p.alignment = PP_ALIGN.RIGHT
                    for r in p.runs:
                        _font(r, 9 if small else 11)
        if t.get("note"):
            _box(s, x, y + h - Inches(0.1), w, Inches(0.4), t["note"], 10, SUB)

    def text(self, title, items, part=None):
        self.bullets(title + (f"（{part}）" if part else ""), items)

    def closing(self, org):
        if self.tpl:
            self._tpl_slide("closing", "谢谢！", org)
            return
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        _rect(s, 0, 0, self.W, self.H, RED_D)
        _box(s, 0, Inches(2.7), self.W, Inches(1.2), "谢  谢！", 48, WHITE, True, PP_ALIGN.CENTER)
        _box(s, 0, Inches(4.1), self.W, Inches(0.6), org, 18, GOLD, False, PP_ALIGN.CENTER)

    def save(self, path):
        n = len(self.prs.slides)
        for s in self.footers:
            no = list(self.prs.slides).index(s) + 1
            _box(s, Inches(0.6), self.H - Inches(0.5), self.W - Inches(1.2), Inches(0.35), f"内部资料 注意保密 · 示例再担保公司普惠业务部 · {no} / {n}", 10, SUB, False, PP_ALIGN.RIGHT)
        self.prs.save(path)
        return n


def _chunks(seq, n):
    return [seq[i:i + n] for i in range(0, len(seq), n)] or [seq]


def _budget(items, n, chars):
    out, cur, size = [], [], 0
    for it in items:
        if cur and (len(cur) >= n or size + len(it) > chars):
            out.append(cur)
            cur, size = [], 0
        cur.append(it)
        size += len(it)
    return out + [cur] if cur else (out or [items])


def render(pages, pack, path, template_path=None):
    """pages：已校验的大纲；返回页数。"""
    d = Deck(template_path)
    meta = pack["meta"]
    sec_no = 0
    agenda_items = [p["title"] for p in pages if p["kind"] == "section"]
    for p in pages:
        k = p["kind"]
        if k == "cover":
            d.cover(p.get("title") or meta["title"], p.get("subtitle") or meta["subtitle"], meta["org"], p.get("date_text") or "")
        elif k == "agenda":
            d.agenda(p["title"] or "目录", agenda_items or [b for b in p.get("bullets", [])])
        elif k == "section":
            sec_no += 1
            d.section(sec_no, p["title"])
        elif k == "content":
            facts = [pack["facts"][f] for f in (p.get("data") or []) if f in pack["facts"]] if isinstance(p.get("data"), list) else None
            d.bullets(p["title"], p.get("bullets"), facts)
        elif k == "kpi":
            d.kpi(p["title"], [pack["facts"][f] for f in p["data"] if f in pack["facts"]], p.get("bullets"))
        elif k == "chart":
            c = pack["charts"][p["data"]]
            if not c.get("categories"):  # 空图表 python-pptx 会抛 ValueError，兜底跳过
                continue
            parts = _chunks(list(range(len(c["categories"]))), CHART_CATS) if c["type"].startswith("bar") else [list(range(len(c["categories"])))]
            for n, idx in enumerate(parts, 1):
                sub = {**c, "categories": [c["categories"][i] for i in idx], "series": [{"name": s["name"], "values": [s["values"][i] for i in idx]} for s in c["series"]]}
                d.chart(p["title"], sub, f"{n}/{len(parts)}" if len(parts) > 1 else None)
        elif k == "table":
            t = pack["tables"][p["data"]]
            longest = max([len(str(c)) for r in t["rows"] for c in r] + [0])
            parts = _chunks(t["rows"], TABLE_ROWS if longest <= 25 else 8 if longest <= 50 else 5)
            for n, rows in enumerate(parts, 1):
                d.table(p["title"], {**t, "rows": rows}, f"{n}/{len(parts)}" if len(parts) > 1 else None)
        elif k == "text":
            items = pack["texts"][p["data"]]["items"]
            parts = _budget(items, TEXT_ITEMS, 520)
            for n, it in enumerate(parts, 1):
                d.text(p["title"], it or ["（无）"], f"{n}/{len(parts)}" if len(parts) > 1 else None)
        elif k == "closing":
            d.closing(meta["org"])
    return d.save(path)


# ---------------- 数字溯源检查 ----------------
def _texts_of(shape):
    if shape.has_text_frame:
        yield shape.text_frame.text
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for c in row.cells:
                yield c.text
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            yield from _texts_of(s)


def trace(path, pack, extra_texts=()):
    """逐页检查：文本/表格中出现的每个数字都必须出现在数据包（数据库取值）中；图表数据必须与数据包图表序列一致。"""
    prs = Presentation(path)
    n = len(prs.slides)
    ok_tok = allowed_tokens(pack, extra_texts)
    struct = {str(i) for i in range(0, n + 1)} | {f"{i:02d}" for i in range(0, n + 1)}
    pack_series = [tuple(round(v or 0, 6) for v in s["values"]) for c in pack["charts"].values() for s in c["series"]]
    viol, checked, charts = [], 0, 0
    for no, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            for txt in _texts_of(sh):
                for m in NUMTOK.finditer(txt or ""):
                    tok = m.group(0).replace(",", "").replace("，", "")
                    checked += 1
                    if tok not in ok_tok and tok not in struct:
                        viol.append({"slide": no, "number": m.group(0), "text": txt[:80]})
            if getattr(sh, "has_chart", False) and sh.has_chart:
                for se in sh.chart.plots[0].series:
                    charts += 1
                    vals = tuple(round(v or 0, 6) for v in se.values)
                    if not any(_is_sub(vals, ps) for ps in pack_series):
                        viol.append({"slide": no, "number": "图表序列", "text": se.name})
    return {"ok": not viol, "numbers_checked": checked, "chart_series_checked": charts, "violations": viol[:50]}


def _is_sub(a, b):
    """图表分页后序列是数据包序列的连续片段。"""
    if not a:
        return True
    for i in range(0, len(b) - len(a) + 1):
        if b[i:i + len(a)] == a:
            return True
    return False


# ---------------- 缩略图（LibreOffice，可选） ----------------
def thumbnails(pptx_path, out_dir, dpi=50):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or not shutil.which("pdftoppm"):
        return []
    os.makedirs(out_dir, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        try:
            subprocess.run([soffice, "--headless", "-env:UserInstallation=file://" + td + "/u", "--convert-to", "pdf", "--outdir", td, pptx_path],
                           check=True, timeout=180, capture_output=True)
            pdf = os.path.join(td, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
            subprocess.run(["pdftoppm", "-png", "-r", str(dpi), pdf, os.path.join(out_dir, "slide")], check=True, timeout=180, capture_output=True)
        except (subprocess.SubprocessError, OSError):
            return []
    return sorted([f for f in os.listdir(out_dir) if f.endswith(".png")], key=lambda f: int(re.sub(r"\D", "", f) or 0))
