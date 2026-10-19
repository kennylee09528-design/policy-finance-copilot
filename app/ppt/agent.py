"""PPT 生成 Agent：新建任务 → Agent 提出澄清问题与需求单（人工确认）→ 生成大纲（人工可编辑、确认）→ 渲染 .pptx → 数字溯源检查。

大模型只产出结构、标题与要点文字；任何数字只能以 {{Fn}} 占位符出现，由代码替换为数据库取值；
图表/表格/文字块由页面 data 字段引用编号，内容全部由代码插入。含字面数字的标题/要点一律丢弃。"""
import json
import os
import re
from datetime import datetime

from fastapi import HTTPException

from .. import db
from ..config import DATA_DIR
from ..llm import service as L
from . import config as C
from . import render as R
from .data import DECKS, build, catalog

SLOT = "ppt.agent"
PPT_DIR = DATA_DIR / "ppt"
TPL_DIR = PPT_DIR / "templates"
KINDS = {"cover": "封面", "agenda": "目录", "section": "章节页", "content": "要点页", "kpi": "指标卡片", "chart": "图表", "table": "表格", "text": "文字块", "closing": "封底"}
PH = re.compile(r"\{\{\s*(F\d+)\s*\}\}")
DIGIT = re.compile(r"[0-9０-９]")
SAFE = re.compile(r'[\\/:*?"<>|\s]+')


def _job(jid):
    j = db.q1("SELECT * FROM ppt_jobs WHERE id=?", (jid,))
    if not j:
        raise HTTPException(404, "PPT 任务不存在")
    for k, d in (("scope", {}), ("pack", {}), ("questions", []), ("sheet", {}), ("outline", []), ("messages", []), ("warnings", [])):
        j[k] = db.jload(j[k], d)
    return j


def _upd(jid, **kw):
    for k in ("scope", "pack", "questions", "sheet", "outline", "messages", "warnings"):
        if k in kw and not isinstance(kw[k], str):
            kw[k] = db.jdump(kw[k])
    kw["updated_at"] = db.now()
    db.execute(f"UPDATE ppt_jobs SET {', '.join(k + '=?' for k in kw)} WHERE id=?", (*kw.values(), jid))


def _msg(j, role, text, **extra):
    j["messages"].append({"role": role, "text": text, "at": db.now(), **extra})


def _json(text, default):
    m = re.search(r"(\{.*\}|\[.*\])", text or "", re.S)
    try:
        return json.loads(m.group(1)) if m else default
    except ValueError:
        return default


def model_of(j=None):
    if (j or {}).get("scope", {}).get("_rule_only") if isinstance((j or {}).get("scope"), dict) else False:
        return None
    return (j or {}).get("model_id") or C.get().get("model_id") or L.assignments().get(SLOT)


def needs_sensitive(pack):
    return bool(pack.get("sensitive"))


def _check_model(mid, pack):
    if not mid:
        return
    L.check_allowed(SLOT, mid)
    # 2026-10-12 用户决定：四类演示稿均可使用公网/未授权敏感数据的模型（不再硬性拦截），
    # 仅软提示；大模型仍只看到数据目录（编号与名称），数字全部由系统从数据库填入。


def model_notice(mid, pack):
    """含内部数据且所选模型未标记「允许处理敏感数据」时返回非阻断提示，否则 None。"""
    if not mid or not needs_sensitive(pack):
        return None
    m = db.q1("SELECT model_id, allow_sensitive FROM llm_models WHERE id=?", (mid,))
    if not m or m["allow_sensitive"]:
        return None
    return (f"当前使用的模型「{m['model_id']}」不是公司自有/授权模型。本类演示稿含内部数据，建议使用公司自有大模型；"
            "继续使用不影响生成：大模型只看到数据目录（编号与名称），不接触具体数值，数字均由系统从数据库填入。")


def _model_label(mid):
    if not mid:
        return "规则（未调用大模型）"
    m = db.q1("SELECT model_id, allow_sensitive FROM llm_models WHERE id=?", (mid,))
    return f"{m['model_id']}{'（授权）' if m['allow_sensitive'] else '（公网/未授权）'}" if m else str(mid)


# ---------------- 文字清洗：数字只能来自占位符 ----------------
def sanitize(text, pack):
    """返回替换占位符后的文字；含字面数字或未知占位符时返回 None（整条丢弃）。"""
    if text is None:
        return None
    text = str(text).strip()
    out = []
    for i, part in enumerate(PH.split(text)):
        if i % 2:
            f = pack["facts"].get(part)
            if not f:
                return None
            out.append(f["value"])
        else:
            if DIGIT.search(part):
                return None
            out.append(part)
    return "".join(out).strip() or None


def normalize(pages, pack, max_pages=30):
    """校验大纲：类型、数据引用、文字清洗；返回 (pages, warnings)。页面保留原始文字（含占位符）以便再编辑。"""
    warns, out = [], []
    for n, p in enumerate(pages or [], 1):
        if not isinstance(p, dict):
            continue
        k = p.get("kind") if p.get("kind") in KINDS else "content"
        data = p.get("data")
        if k == "kpi":
            data = [x.strip() for x in (data.split(",") if isinstance(data, str) else (data or [])) if str(x).strip() in pack["facts"]]
            if not data:
                warns.append(f"第 {n} 页指标卡片未引用有效指标编号，已删除")
                continue
        elif k == "content":
            data = [x for x in (data if isinstance(data, list) else []) if x in pack["facts"]][:4] or None
        elif k in ("chart", "table", "text"):
            key = {"chart": "charts", "table": "tables", "text": "texts"}[k]
            if data not in pack[key]:
                warns.append(f"第 {n} 页引用的数据块 {data!r} 不存在或类型不符，已删除")
                continue
            if k == "chart" and not (pack["charts"][data].get("categories") and any(s.get("values") for s in pack["charts"][data].get("series") or [])):
                warns.append(f"第 {n} 页图表「{pack['charts'][data]['title']}」无数据，已删除")
                continue
            if k == "table" and not pack["tables"][data]["rows"]:
                warns.append(f"第 {n} 页表格「{pack['tables'][data]['title']}」无数据，已删除")
                continue
            if k == "text" and not pack["texts"][data]["items"]:
                warns.append(f"第 {n} 页文字块「{pack['texts'][data]['title']}」无内容，已删除")
                continue
        else:
            data = None
        default_title = {"chart": lambda: pack["charts"][data]["title"], "table": lambda: pack["tables"][data]["title"],
                         "text": lambda: pack["texts"][data]["title"]}.get(k, lambda: {"agenda": "目录", "closing": "谢谢", "cover": pack["meta"]["title"]}.get(k, "要点"))
        title = str(p.get("title") or "").strip()
        if not title or sanitize(title, pack) is None:
            if title:
                warns.append(f"第 {n} 页标题「{title[:20]}」含非系统数字，已改用默认标题")
            title = default_title()
        bullets = []
        for b in (p.get("bullets") or [])[:6]:
            if sanitize(b, pack) is None:
                warns.append(f"第 {n} 页要点「{str(b)[:24]}」含非系统数字（数字须用 {{{{Fn}}}} 占位符），已删除")
            else:
                bullets.append(str(b).strip())
        if k == "content" and not bullets and not data:
            warns.append(f"第 {n} 页「{title}」没有有效要点，已删除")
            continue
        out.append({"kind": k, "title": title[:40], "bullets": bullets, "data": data, "note": str(p.get("note") or "")[:200]})
    if pack.get("deck_type") == "policy":
        out = _policy_ground(out, pack, warns)
    out = [p for i, p in enumerate(out) if p["kind"] != "cover" or i == 0]
    if not out or out[0]["kind"] != "cover":
        out.insert(0, {"kind": "cover", "title": pack["meta"]["title"], "bullets": [], "data": None, "note": ""})
    out = [p for i, p in enumerate(out) if p["kind"] != "closing" or i == len(out) - 1]
    if out[-1]["kind"] != "closing":
        out.append({"kind": "closing", "title": "谢谢", "bullets": [], "data": None, "note": ""})
    if len(out) > max_pages:
        warns.append(f"大纲超过上限 {max_pages} 页，已截断")
        out = out[:max_pages - 1] + [out[-1]]
    return out, warns


def cover_subtitle(pack, sheet):
    """封面副标题：数据范围（日期/快照/方案）+ 受众，不与标题重复。"""
    sub = pack["meta"].get("subtitle") or ""
    aud = str((sheet or {}).get("audience") or "").strip()
    return f"{sub} · 面向：{aud}" if aud and DIGIT.search(aud) is None else sub


NO_ITEMS = "本期暂无条目"
POLICY_TIERS = (("今日必读", 3), ("值得关注", 2), ("同业动态", 1), ("浏览", 1))


def _policy_ground(out, pack, warns):
    """政策简报：要点必须来自简报条目。自由要点（无 {{Fn}} 占位符）删除；各星级章节后挂对应要点块，
    无条目时只放一行「本期暂无条目」；数据速览章节必须带 T4 表。"""
    def tier_of(title):
        return next((t for k, t in POLICY_TIERS if k in (title or "")), None)
    res = []
    for p in out:
        if p["kind"] == "content":
            free = [b for b in p["bullets"] if b != NO_ITEMS and not PH.search(b)]
            if free:
                warns.append(f"「{p['title']}」中 {len(free)} 条要点不是来自简报条目，已删除")
                p = {**p, "bullets": [b for b in p["bullets"] if b not in free]}
            if not p["bullets"] and not p["data"]:
                continue
        res.append(p)
    out, res = res, []
    for i, p in enumerate(out):
        res.append(p)
        if p["kind"] != "section":
            continue
        nxt = out[i + 1] if i + 1 < len(out) else {}
        tier = tier_of(p["title"])
        if tier:
            x, t = f"X{tier}", f"T{tier}"
            if nxt.get("data") in (x, t) or (nxt.get("kind") == "content" and nxt.get("bullets") == [NO_ITEMS]):
                continue
            if pack["texts"].get(x, {}).get("items"):
                res.append({"kind": "text", "title": pack["texts"][x]["title"], "bullets": [], "data": x, "note": ""})
            else:
                res.append({"kind": "content", "title": p["title"], "bullets": [NO_ITEMS], "data": None, "note": ""})
        elif "数据速览" in p["title"] and nxt.get("data") != "T4" and pack["tables"].get("T4", {}).get("rows"):
            res.append({"kind": "table", "title": pack["tables"]["T4"]["title"], "bullets": [], "data": "T4", "note": ""})
    if pack["tables"].get("T4", {}).get("rows") and not any(p.get("data") == "T4" for p in res):
        at = len(res) - 1 if res and res[-1]["kind"] == "closing" else len(res)
        res[at:at] = [{"kind": "section", "title": "数据速览", "bullets": [], "data": None, "note": ""},
                      {"kind": "table", "title": pack["tables"]["T4"]["title"], "bullets": [], "data": "T4", "note": ""}]
    return res


def rendered(pages, pack):
    """渲染用：占位符替换为数据库数值。"""
    return [{**p, "title": sanitize(p["title"], pack) or p["title"], "bullets": [sanitize(b, pack) for b in p["bullets"] if sanitize(b, pack)]} for p in pages]


# ---------------- 规则大纲（无模型时的回退，也是默认结构） ----------------
def rule_outline(pack, sheet):
    t = pack["deck_type"]
    P = [{"kind": "cover"}, {"kind": "agenda", "title": "目录"}]
    f = pack["facts"]
    if t == "usage":
        P += [{"kind": "section", "title": "总体情况"},
              {"kind": "kpi", "title": "额度使用总体情况", "data": [k for k in ("F1", "F2", "F3", "F4", "F5", "F6", "F7", "F8", "F9") if k in f],
               "bullets": ["常规业务使用率（÷国担额度）为 {{F2}}，年度时间进度为 {{F8}}"]},
              {"kind": "chart", "data": "C1"},
              {"kind": "section", "title": "分机构使用进度"}, {"kind": "chart", "data": "C3"}, {"kind": "chart", "data": "C2"},
              {"kind": "table", "data": "T1"}, {"kind": "chart", "data": "C4"},
              {"kind": "section", "title": "预警事项"}, {"kind": "text", "data": "X1"}]
        if pack["texts"]["X2"]["items"]:
            P += [{"kind": "section", "title": "环比说明"}, {"kind": "text", "data": "X2"}]
        P += [{"kind": "content", "title": "下一步工作建议", "bullets": ["对使用率偏低机构开展专项督导，逐户梳理储备项目", "加强民间投资专项额度推介，提升专项使用效率",
                                                                  "按周跟踪预警机构，及时报送额度调剂建议"]}]
    elif t == "scheme":
        P += [{"kind": "section", "title": "测算依据与系数规则"},
              {"kind": "content", "title": "测算依据", "bullets": ["以各合作机构上年度业务数据与评价结果为基础", "按就业贡献、支小支农、代偿率、评价分等系数综合测算",
                                                               "经三步调整形成常规授信额度，另行核定民间投资专项"]},
              {"kind": "table", "data": "T1"},
              {"kind": "section", "title": "三步调整结果"},
              {"kind": "kpi", "title": "授信方案总体情况", "data": [k for k in ("F1", "F2", "F3", "F4", "F5", "F6", "F7", "F8", "F9") if k in f]},
              {"kind": "chart", "data": "C2"}, {"kind": "chart", "data": "C1"}, {"kind": "table", "data": "T2"},
              {"kind": "section", "title": "人工调整及原因"}, {"kind": "table", "data": "T3"}, {"kind": "text", "data": "X1"},
              {"kind": "section", "title": "提请审议事项"},
              {"kind": "content", "title": "提请审议事项", "bullets": ["审议本授信方案及各合作担保机构授信额度", "审议人工调整事项及调整原因",
                                                                 "授权普惠业务部按方案办理额度下达及后续调剂"], "data": ["F9"]}]
    elif t == "policy":
        P += [{"kind": "kpi", "title": "本期概览", "data": ["F1", "F2", "F3", "F4", "F5"]}, {"kind": "chart", "data": "C1"}]
        for tier, name in ((3, "今日必读"), (2, "值得关注"), (1, "同业动态与浏览")):
            if pack["tables"][f"T{tier}"]["rows"]:
                P += [{"kind": "section", "title": name}, {"kind": "text", "data": f"X{tier}"}, {"kind": "table", "data": f"T{tier}"}]
        P += [{"kind": "chart", "data": "C2"}]
        if "T4" in pack["tables"]:
            P += [{"kind": "section", "title": "数据速览"}, {"kind": "table", "data": "T4"}]
    else:
        for k, s in enumerate(pack["sections"], 1):
            P.append({"kind": "section", "title": s})
            plain = pack["texts"].get(f"P{k}", {}).get("items", [])
            if plain:
                P.append({"kind": "content", "title": s.split("、", 1)[-1][:20], "bullets": [x[:60] for x in plain[:4]]})
            if pack["texts"].get(f"X{k}", {}).get("items"):
                P.append({"kind": "text", "data": f"X{k}"})
        if pack["tables"]["T1"]["rows"] and sheet.get("appendix", True):
            P += [{"kind": "section", "title": "附录"}, {"kind": "table", "data": "T1"}]
    P.append({"kind": "closing"})
    return P


# ---------------- 主流程 ----------------
def create(body, user):
    t = body.get("deck_type")
    if t not in DECKS:
        raise HTTPException(400, "请选择演示稿类型")
    scope = body.get("scope") or {}
    pack = build(t, scope)
    mid = int(body["model_id"]) if body.get("model_id") not in (None, "") else (None if scope.get("_rule_only") else model_of())
    _check_model(mid, pack)
    tpl = body.get("template_id")
    if tpl in (None, ""):
        a = db.q1("SELECT id FROM ppt_templates WHERE active=1 ORDER BY id DESC LIMIT 1")
        tpl = a["id"] if a else None
    pv = C.active_prompt(t)["version"]
    jid = db.execute("INSERT INTO ppt_jobs(deck_type,title,scope,model_id,template_id,status,step,pack,messages,warnings,tokens,prompt_version,created_by,created_at,updated_at) "
                     "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                     (t, body.get("title") or pack["meta"]["title"], db.jdump(scope), mid, tpl, "需求确认中", 1, db.jdump(pack), "[]", "[]", 0, pv,
                      user["name"], db.now(), db.now()))
    db.audit(user, "PPT", "新建 PPT 任务", jid, {"deck_type": t, "scope": scope, "model": mid})
    j = _job(jid)
    _propose_sheet(j, user)
    return detail(jid)


def _call(j, user, system, prompt):
    mid = model_of(j)
    _check_model(mid, j["pack"])
    cfg = C.get()
    r = L.run(SLOT, [{"role": "system", "content": system}, {"role": "user", "content": prompt}], user, mid=mid,
              overrides={"temperature": cfg.get("temperature"), "top_p": cfg.get("top_p")}, timeout=cfg.get("timeout") or 120,
              ref=f"ppt_job:{j['id']}", prompt_version=f"{j.get('deck_type')} v{j.get('prompt_version') or 0}")
    j["tokens"] = (j.get("tokens") or 0) + r["tokens"]
    return r


def _propose_sheet(j, user):
    t, pack = j["deck_type"], j["pack"]
    sheet = {**C.DEFAULT_SHEET[t], "title": pack["meta"]["title"], "sections": list(pack.get("sections") or [])}
    questions = list(C.DEFAULT_QUESTIONS[t])
    via = "规则"
    if model_of(j):
        prompt = (f"{catalog(pack)}\n\n【默认需求单】{json.dumps(sheet, ensure_ascii=False)}\n\n"
                  "请作为 PPT 策划助手，先与用户确认需求：提出 3~5 个最关键的澄清问题，并给出建议需求单（可调整受众、目的、风格、语气、重点、章节）。"
                  '仅输出 JSON：{"questions":["..."],"sheet":{"audience":"","purpose":"","pages":12,"style":"","tone":"","emphasis":"","sections":["..."]}}')
        try:
            r = _call(j, user, C.active_prompt(t)["content"], prompt)
            o = _json(r["content"], {})
            if isinstance(o.get("questions"), list) and o["questions"]:
                questions = [str(x)[:120] for x in o["questions"][:6]]
            for k, v in (o.get("sheet") or {}).items():
                if k in ("audience", "purpose", "style", "tone", "emphasis") and isinstance(v, str) and v.strip():
                    sheet[k] = v.strip()[:200]
                elif k == "pages" and str(v).isdigit():
                    sheet[k] = max(4, min(int(v), C.get().get("max_pages") or 30))
                elif k == "sections" and isinstance(v, list) and v:
                    sheet[k] = [str(x)[:30] for x in v[:10]]
            via = r["model"]
        except HTTPException as e:
            j["warnings"].append(f"需求确认阶段模型调用失败，已使用默认问题：{e.detail}")
    _msg(j, "agent", "我先确认几项需求，请回答下列问题，并检查右侧的需求单（可直接修改）。确认后我会生成大纲。", questions=questions, via=via)
    _upd(j["id"], questions=questions, sheet=sheet, messages=j["messages"], warnings=j["warnings"], tokens=j["tokens"], status="待确认需求", step=2)


def confirm_sheet(jid, body, user):
    j = _job(jid)
    if j["step"] < 2:
        raise HTTPException(400, "需求单尚未生成")
    sheet = {**j["sheet"], **{k: v for k, v in (body.get("sheet") or {}).items() if k in ("title", "audience", "purpose", "pages", "style", "tone", "emphasis", "sections", "period", "appendix")}}
    answers = [str(a).strip() for a in (body.get("answers") or []) if str(a).strip()]
    if body.get("message"):
        answers.append(str(body["message"]).strip())
    _msg(j, "user", "需求单已确认。" + ("\n" + "\n".join(f"· {a}" for a in answers) if answers else ""), sheet=sheet)
    if sheet.get("title") and sanitize(sheet["title"], j["pack"]) is None:
        j["warnings"].append("演示稿标题含非系统数字，封面仍使用系统标题")
    _upd(jid, sheet=sheet, sheet_confirmed_at=db.now(), messages=j["messages"], warnings=j["warnings"])
    db.audit(user, "PPT", "确认需求单", jid, {"pages": sheet.get("pages")})
    j = _job(jid)
    _make_outline(j, user, answers)
    return detail(jid)


def _make_outline(j, user, answers=()):
    t, pack, sheet = j["deck_type"], j["pack"], j["sheet"]
    pages, via, warns = None, "规则", []
    if model_of(j):
        ref = ""
        if t == "research":
            ref = "\n【报告各章无数字句（供提炼要点）】\n" + "\n".join(f"{k}：" + " / ".join(v["items"][:8]) for k, v in pack["texts"].items() if k.startswith("P"))
        elif t == "policy":
            ref = "\n【入选条目标题】\n" + "\n".join(DIGIT.sub("#", x.split("：")[0]) for k in ("X3", "X2", "X1") for x in pack["texts"][k]["items"][:15])
        prompt = (f"{catalog(pack)}{ref}\n\n【已确认需求单】{json.dumps(sheet, ensure_ascii=False)}\n【用户补充】{'；'.join(answers) or '无'}\n\n"
                  f"请生成演示稿大纲（约 {sheet.get('pages') or 12} 页）。页面类型 kind 只能是：{', '.join(KINDS)}。\n"
                  "chart/table/text 页用 data 字段写一个对应编号；kpi 页 data 写指标编号数组；content 页写 bullets（3~5 条），需要数据时用 {{Fn}} 占位符；"
                  "section 页只写 title。严禁任何阿拉伯数字。\n"
                  '仅输出 JSON：{"pages":[{"kind":"cover"},{"kind":"agenda","title":"目录"},{"kind":"section","title":""},'
                  '{"kind":"content","title":"","bullets":[""],"data":["F1"]},{"kind":"chart","title":"","data":"C1"},{"kind":"closing"}]}')
        try:
            r = _call(j, user, C.active_prompt(t)["content"], prompt)
            o = _json(r["content"], {})
            cand = o.get("pages") if isinstance(o, dict) else o
            if isinstance(cand, list) and len(cand) >= 3:
                pages, via = cand, r["model"]
            else:
                warns.append("模型输出的大纲无法解析，已使用规则大纲")
        except HTTPException as e:
            warns.append(f"大纲生成阶段模型调用失败，已使用规则大纲：{e.detail}")
    if pages is None:
        pages = rule_outline(pack, sheet)
    pages, w2 = normalize(pages, pack, C.get().get("max_pages") or 30)
    warns += w2
    if via != "规则" and not any(p["kind"] in ("chart", "table", "kpi") for p in pages):
        warns.append("模型大纲未引用任何图表/表格，已追加系统数据页")
        extra, _ = normalize(rule_outline(pack, sheet), pack)
        pages = pages[:-1] + [p for p in extra if p["kind"] in ("chart", "table", "kpi")] + pages[-1:]
    _msg(j, "agent", f"已生成大纲，共 {len(pages)} 页（不含表格/图表自动分页）。请逐页检查，可修改标题、要点、调整顺序或删除页面；确认后生成 PPT。"
         + (f"\n注意：{len(warns)} 条自动修正，见提示。" if warns else ""), via=via)
    _upd(j["id"], outline=pages, messages=j["messages"], warnings=j["warnings"] + warns, tokens=j["tokens"], status="待确认大纲", step=3)


def regenerate_outline(jid, body, user):
    j = _job(jid)
    if not j["sheet_confirmed_at"]:
        raise HTTPException(400, "请先确认需求单")
    msg = str(body.get("message") or "").strip()
    if msg:
        _msg(j, "user", msg)
    _make_outline(j, user, [msg] if msg else [])
    return detail(jid)


def save_outline(jid, body, user):
    j = _job(jid)
    if j["step"] < 3:
        raise HTTPException(400, "请先确认需求单")
    pages, warns = normalize(body.get("pages"), j["pack"], C.get().get("max_pages") or 30)
    _msg(j, "user", f"修改了大纲（{len(pages)} 页）。")
    _upd(jid, outline=pages, messages=j["messages"], warnings=warns, outline_confirmed_at=None)
    db.audit(user, "PPT", "编辑大纲", jid, {"pages": len(pages)})
    return {**detail(jid), "outline_warnings": warns}


def generate(jid, body, user):
    j = _job(jid)
    if j["step"] < 3 or not j["outline"]:
        raise HTTPException(400, "请先确认需求单并生成大纲")
    if not body.get("confirm"):
        raise HTTPException(400, "请先确认大纲（confirm=true）后再生成 PPT")
    if body.get("refresh_data"):
        j["pack"] = build(j["deck_type"], j["scope"])
        _upd(jid, pack=j["pack"])
    pack = j["pack"]
    _check_model(model_of(j), pack)
    pages, warns = normalize(j["outline"], pack, C.get().get("max_pages") or 30)
    tpl_id = body.get("template_id", j["template_id"])
    tpl_path = None
    if tpl_id not in (None, ""):
        tp = db.q1("SELECT * FROM ppt_templates WHERE id=?", (int(tpl_id),))
        if tp and os.path.exists(tp["path"]):
            tpl_path = tp["path"]
        else:
            warns.append("所选模板不存在，已使用默认主题")
            tpl_id = None
    v = (db.q1("SELECT MAX(version) v FROM ppt_files WHERE job_id=?", (jid,))["v"] or 0) + 1
    out_dir = PPT_DIR / f"job{jid}"
    out_dir.mkdir(parents=True, exist_ok=True)
    fname = SAFE.sub("_", j["title"])[:40] + f"_v{v}.pptx"
    path = str(out_dir / f"v{v}.pptx")
    date_text = datetime.now().strftime("%Y年%m月%d日")
    rp = rendered(pages, pack)
    rp[0]["date_text"] = date_text
    rp[0]["subtitle"] = cover_subtitle(pack, j["sheet"])
    n = R.render(rp, pack, path, tpl_path)
    tr = R.trace(path, pack, extra_texts=[date_text])
    thumbs = R.thumbnails(path, str(out_dir / f"v{v}_thumbs")) if body.get("thumbnails", True) else []
    fid = db.execute("INSERT INTO ppt_files(job_id,version,filename,path,pages,thumbs,trace,outline,sheet,template_id,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                     (jid, v, fname, path, n, db.jdump(thumbs), db.jdump(tr), db.jdump(pages), db.jdump(j["sheet"]), tpl_id, user["name"], db.now()))
    _msg(j, "agent", f"已生成 PPT 第 {v} 版，共 {n} 页。数字溯源检查：{'通过' if tr['ok'] else '发现 ' + str(len(tr['violations'])) + ' 处疑点'}"
         f"（检查数字 {tr['numbers_checked']} 个、图表序列 {tr['chart_series_checked']} 个，均须来自数据库）。", file_id=fid)
    _upd(jid, outline=pages, outline_confirmed_at=db.now(), messages=j["messages"], warnings=warns, status="已生成", step=4)
    db.audit(user, "PPT", "生成 PPT", jid, {"version": v, "pages": n, "trace_ok": tr["ok"], "template": tpl_id,
                                           "model": _model_label(model_of(j)), "model_notice": bool(model_notice(model_of(j), pack))})
    return detail(jid)


def detail(jid):
    j = _job(jid)
    pack = j.pop("pack")
    j["deck_name"] = DECKS[j["deck_type"]]["name"]
    j["sensitive"] = needs_sensitive(pack)
    j["model_notice"] = model_notice(model_of(j), pack)
    j["catalog"] = {"facts": {k: {"name": f["name"], "value": f["value"]} for k, f in pack.get("facts", {}).items()},
                    "charts": {k: {"title": c["title"], "type": c["type"], "n": len(c["categories"])} for k, c in pack.get("charts", {}).items()},
                    "tables": {k: {"title": t["title"], "n": len(t["rows"])} for k, t in pack.get("tables", {}).items()},
                    "texts": {k: {"title": t["title"], "n": len(t["items"])} for k, t in pack.get("texts", {}).items()}}
    j["meta"] = pack.get("meta")
    j["source"] = pack.get("source")
    j["preview"] = rendered(j["outline"], pack) if j["outline"] else []
    j["files"] = files(jid)
    return j


def files(jid):
    rows = db.q("SELECT id, job_id, version, filename, pages, thumbs, trace, template_id, created_by, created_at FROM ppt_files WHERE job_id=? ORDER BY version DESC", (jid,))
    for r in rows:
        r["thumbs"] = db.jload(r["thumbs"], [])
        r["trace"] = db.jload(r["trace"], {})
    return rows


def list_jobs():
    rows = db.q("SELECT j.id, j.deck_type, j.title, j.status, j.step, j.created_by, j.created_at, j.updated_at, "
                "(SELECT MAX(version) FROM ppt_files f WHERE f.job_id=j.id) versions FROM ppt_jobs j ORDER BY j.id DESC LIMIT 2000")
    for r in rows:
        r["deck_name"] = DECKS[r["deck_type"]]["name"]
    return rows


def file_row(fid):
    f = db.q1("SELECT * FROM ppt_files WHERE id=?", (fid,))
    if not f or not os.path.exists(f["path"]):
        raise HTTPException(404, "文件不存在")
    return f


# ---------------- 模板（管理员上传） ----------------
def upload_template(name, filename, data, user):
    from io import BytesIO

    from pptx import Presentation
    if not filename.lower().endswith(".pptx"):
        raise HTTPException(400, "仅支持 .pptx 模板")
    if len(data) > 30 * 1024 * 1024:
        raise HTTPException(400, "模板文件不能超过 30MB")
    try:
        prs = Presentation(BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"无法解析该 .pptx：{e}") from e
    layouts = [{"name": l.name, "placeholders": [p.placeholder_format.type.__str__() for p in l.placeholders]} for l in prs.slide_layouts]
    TPL_DIR.mkdir(parents=True, exist_ok=True)
    tid = db.execute("INSERT INTO ppt_templates(name,filename,path,layouts,active,uploaded_by,uploaded_at) VALUES(?,?,?,?,0,?,?)",
                     (name or filename, filename, "", db.jdump(layouts), user["name"], db.now()))
    path = str(TPL_DIR / f"tpl{tid}.pptx")
    with open(path, "wb") as fh:
        fh.write(data)
    db.execute("UPDATE ppt_templates SET path=? WHERE id=?", (path, tid))
    db.audit(user, "PPT", "上传 PPT 模板", tid, {"name": name, "layouts": len(layouts)})
    return template(tid)


def template(tid):
    t = db.q1("SELECT id, name, filename, layouts, active, uploaded_by, uploaded_at FROM ppt_templates WHERE id=?", (tid,))
    if not t:
        raise HTTPException(404, "模板不存在")
    t["layouts"] = db.jload(t["layouts"], [])
    return t


def templates():
    return [template(r["id"]) for r in db.q("SELECT id FROM ppt_templates ORDER BY id DESC")]


def activate_template(tid, user):
    if tid:
        template(tid)
    db.execute("UPDATE ppt_templates SET active=CASE WHEN id=? THEN 1 ELSE 0 END", (tid or 0,))
    db.audit(user, "PPT", "设置默认 PPT 模板", tid or "默认深红主题")


def delete_template(tid, user):
    t = db.q1("SELECT * FROM ppt_templates WHERE id=?", (tid,))
    if not t:
        raise HTTPException(404, "模板不存在")
    if t["path"] and os.path.exists(t["path"]):
        os.remove(t["path"])
    db.execute("DELETE FROM ppt_templates WHERE id=?", (tid,))
    db.audit(user, "PPT", "删除 PPT 模板", tid, {"name": t["name"]})
