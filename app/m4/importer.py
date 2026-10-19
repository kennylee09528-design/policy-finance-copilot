"""附件2《各合作担保机构授信额度计算表》导入（10 张分表）。

做法：同时以「公式」和「缓存值」两种模式打开工作簿。
  * 从计算表每一行的公式中解析出每个指标取自哪张分表、哪一列（FILTER 查找），以及
    基期常数（K 列）、权重（U 列）、全省均值单元格（L/P/R 列引用）、参与分配的行范围（X 列 SUM 区间）、
    固定额度机构（X 列为常数而非公式）；
  * 再按解析出的口径到分表中按机构名取原始数值；
  * 原计算表各列缓存值一并保存，用于逐行验收比对。
"""
import re
from io import BytesIO

import openpyxl

from .xlsxformula import Lookup, eval_simple, ftext, parse_filters, parse_ref, norm_name

CALC_SHEET = "计算表"
EXPECTED_SHEETS = ["计算表", "合作业务授信额度申请表", "2025年民投", "2025年规模", "就业人数", "支小支农业务统计表",
                   "近3年累计代偿率", "2025年代偿率", "2025年机构评价", "2025年互联网落地"]
CACHED_COLS = ["C", "D", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V", "W", "X",
               "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI"]


def _num(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# PRD（feedback_1018）：按工作表名关键词定位 9 张分表；前 8 张为核心表（缺任何一张即拒绝导入）
SHEET_KEYS = [  # (标准名, 关键词正则, 是否核心)
    ("合作业务授信额度申请表", r"申请表", True),
    ("2025年民投", r"(\d{4}年)?民投", True),
    ("2025年规模", r"(\d{4}年)?规模", True),
    ("就业人数", r"就业人数", True),
    ("支小支农业务统计表", r"支小支农", True),
    ("近3年累计代偿率", r"累计代偿率", True),
    ("2025年代偿率", r"^(?!.*累计).*\d{4}年代偿率", True),
    ("2025年机构评价", r"机构评价", True),
    ("2025年互联网落地", r"互联网落地", False),
]
FIXED_KEYWORDS = ["示例担保公司I", "示例担保公司J"]
DEFAULT_BASE_CONST = 50000.0


def _key_of(name):
    for std, pat, _ in SHEET_KEYS:
        if name == std or re.search(pat, name or ""):
            return std
    return None


def _find_sheet(wb, name):
    if name in wb.sheetnames:
        return name
    for s in wb.sheetnames:  # 容错：年份变化（如 2026年规模）
        if re.sub(r"\d{4}年", "", s) == re.sub(r"\d{4}年", "", name):
            return s
    std = _key_of(name)  # 关键词定位
    if std:
        pat = next(p for n, p, _ in SHEET_KEYS if n == std)
        for s in wb.sheetnames:
            if s != CALC_SHEET and re.search(pat, s) and not (std == "2025年民投" and "申请" in s):
                return s
    return None


def locate_sheets(wb):
    """返回 {标准名: 实际表名}，缺失核心表时抛 ValueError。"""
    found = {std: _find_sheet(wb, std) for std, _, _ in SHEET_KEYS}
    lack = [std for std, _, core in SHEET_KEYS if core and not found[std]]
    if lack:
        raise ValueError(f"缺少核心分表（需至少识别前 8 张）：{'、'.join(lack)}；已有工作表：{'、'.join(wb.sheetnames)}")
    return found


def _names_col(ws, col, start):
    out = []
    for r in range(start, ws.max_row + 1):
        v = ws[f"{col}{r}"].value
        if isinstance(v, str) and v.strip() and v.strip() not in ("合计", "总计") and not re.match(r"^\d+、", v.strip()):
            out.append((r, v.strip()))
    return out


def base_list(wv, found):
    """以「2025年规模」C 列为基准名单（第 7 行起，跳过合计与注释行）。"""
    return [n for _, n in _names_col(wv[found["2025年规模"]], "C", 7)]


def _direct_map(wv, found):
    """PRD 第 3 节列映射：{字段: {机构名: 值}}，用于无「计算表」时的兜底导入。"""
    def col(std, kcol, vcol, start):
        if not found.get(std):
            return {}
        ws = wv[found[std]]
        return {norm_name(n): _num(ws[f"{vcol}{r}"].value) or 0.0 for r, n in _names_col(ws, kcol, start)}
    return {
        "app_c": col("合作业务授信额度申请表", "B", "C", 5), "app_d": col("合作业务授信额度申请表", "B", "D", 5),
        "app_f": col("合作业务授信额度申请表", "B", "F", 5), "app_i": col("合作业务授信额度申请表", "B", "I", 5),
        "mintou": col("2025年民投", "G", "H", 2), "scale": col("2025年规模", "C", "D", 7),
        "emp": col("就业人数", "C", "E", 7), "zxzn": col("支小支农业务统计表", "C", "E", 7),
        "lj": col("近3年累计代偿率", "C", "J", 7), "dn": col("2025年代偿率", "C", "J", 7),
        "pj": col("2025年机构评价", "B", "C", 2), "net": col("2025年互联网落地", "B", "C", 3),
    }


def is_fixed_name(name, keywords=None):
    return any(k in name for k in (keywords or FIXED_KEYWORDS))


def import_direct(wv, found, filename):
    """无「计算表」时：按 PRD 列映射直接取数（基准名单=规模表 C 列 ∩ 申请表；固定额度按关键词识别）。"""
    m = _direct_map(wv, found)
    alerts = [{"level": "warn", "type": "无计算表", "msg": "工作簿中没有「计算表」：按 PRD 列映射取数；"
               "截止备案金额(AF)、工行(E/F/G)、固定额度金额无来源，需在基础数据中补录"}]
    base = base_list(wv, found)
    applicants = {norm_name(n) for n in m["app_c"]}
    names = [n for n in base if norm_name(n) in applicants]
    for n in m["app_c"]:  # 固定额度机构可能不在规模表中
        if is_fixed_name(n) and norm_name(n) not in {norm_name(x) for x in names}:
            names.append(n)
    skipped = [n for n in base if norm_name(n) not in applicants]
    if skipped:
        alerts.append({"level": "info", "type": "名单差异", "msg": f"规模表中 {len(skipped)} 家不在申请表、不参与测算：{'、'.join(skipped)}"})
    rows = []
    for i, n in enumerate(names):
        k = norm_name(n)
        g = lambda key: m[key].get(k)  # noqa: E731
        fixed = is_fixed_name(n)
        rec = {"row": None, "seq": i + 1, "name": n, "sources": {"mode": "PRD 列映射"}, "notes": [], "excel": {},
               "excel_formulas": {}, "scale": g("scale"), "mintou_prev": g("mintou") or 0.0, "emp": g("emp"),
               "base_const": 0.0, "zxzn_raw": g("zxzn"), "lj_raw": g("lj"), "dn_raw": g("dn"),
               "pj": g("pj"), "zc_total": g("app_i"), "v_const": None,
               "apply_regular_tech": (g("app_c") or 0) + (g("app_d") or 0), "mintou_special": g("app_f") or 0.0,
               "tech_apply": g("app_d"), "tech_quota": None, "filed_amount": None, "net_amount": g("net"),
               "icbc": {}, "fixed": fixed, "fixed_quota": None}
        for key, lab in (("scale", "2025年规模"), ("emp", "就业人数"), ("zxzn", "支小支农"), ("lj", "近3年累计代偿率"),
                         ("dn", "2025年代偿率"), ("pj", "2025年机构评价")):
            if not fixed and g(key) is None:
                alerts.append({"level": "warn", "type": "名单不匹配", "msg": f"「{n}」在分表「{lab}」中未匹配到，已忽略"})
        if fixed:
            alerts.append({"level": "warn", "type": "固定额度", "msg": f"「{n}」为固定额度机构，额度需在基础数据中填写"})
        rows.append(rec)
    return rows, alerts


def import_workbook(data: bytes, filename="附件2.xlsx"):
    wf = openpyxl.load_workbook(BytesIO(data))                  # 公式
    wv = openpyxl.load_workbook(BytesIO(data), data_only=True)  # 缓存值
    alerts = []
    smap = locate_sheets(wf)  # 缺核心表直接拒绝
    sheets_found = [CALC_SHEET] * (CALC_SHEET in wf.sheetnames) + [v for v in smap.values() if v]
    if not smap.get("2025年互联网落地"):
        alerts.append({"level": "warn", "type": "缺少分表", "msg": "未找到「2025年互联网落地」（参考 AG），不影响测算"})
    def sheet_const(std, coord):
        try:
            return eval_simple(wf, wv, smap[std], coord)
        except ValueError:
            return _num(wv[smap[std]][coord].value)
    if CALC_SHEET not in wf.sheetnames:
        rows, al = import_direct(wv, smap, filename)
        consts = {"emp_avg": sheet_const("就业人数", "F3"), "lj_avg": sheet_const("近3年累计代偿率", "J3"),
                  "dn_avg": sheet_const("2025年代偿率", "J3"),
                  "base_consts": {n["name"]: DEFAULT_BASE_CONST for n in rows if n["name"] == "示例担保公司A"},
                  "emp_avg_ref": f"{smap['就业人数']}!F3", "lj_avg_ref": f"{smap['近3年累计代偿率']}!J3",
                  "dn_avg_ref": f"{smap['2025年代偿率']}!J3", "base_const_ref": "默认（PRD 4.1：示例担保公司A 50000；其他机构如需请在参数配置中补充）"}
        return {"source_file": filename, "sheets": wf.sheetnames, "sheets_found": sheets_found, "sheet_map": smap,
                "mode": "direct", "weights_from_workbook": {"L": 0.2, "N": 0.2, "P": 0.1, "R": 0.1, "T": 0.4},
                "zxzn_threshold": 0.9, "constants": consts, "pool_rows": None, "header": {}, "rows": rows,
                "alerts": alerts + al, "base_list": base_list(wv, smap)}

    ws, vs = wf[CALC_SHEET], wv[CALC_SHEET]
    header = {c.column_letter: (c.value or "") for c in ws[1]}
    lk = Lookup(wv)

    # ---- 1. 全局口径：权重、全省均值、参与分配区间 ----
    first = 2
    u_f = ftext(ws[f"U{first}"].value) or ""
    wm = re.search(r"L\d+\*([\d.]+)\+N\d+\*([\d.]+)\+P\d+\*([\d.]+)\+R\d+\*([\d.]+)\+T\d+\*([\d.]+)", u_f)
    weights = dict(zip(["L", "N", "P", "R", "T"], [float(x) for x in wm.groups()])) if wm else \
        {"L": 0.2, "N": 0.2, "P": 0.1, "R": 0.1, "T": 0.4}
    if not wm:
        alerts.append({"level": "warn", "type": "公式解析", "msg": "未能从 U 列公式解析权重，使用默认 0.2/0.2/0.1/0.1/0.4"})
    n_f = ftext(ws[f"N{first}"].value) or ""
    nm = re.search(r">=([\d.]+)%", n_f)
    zxzn_threshold = float(nm.group(1)) / 100 if nm else 0.9

    const_refs = {}
    for col, key in (("L", "emp_avg"), ("P", "lj_avg"), ("R", "dn_avg")):
        f = ftext(ws[f"{col}{first}"].value) or ""
        part = f.split("/")[-1] if col == "L" else f[f.find("-(") + 2:]
        ref = parse_ref(part.split("-")[-1] if col != "L" else part)
        if col != "L":
            m = re.search(r"-\((?:[A-Z]+\d+)-((?:'[^']+'|[^'!,()=]+)!\$?[A-Z]+\$?\d+)\)", f)
            ref = parse_ref(m.group(1)) if m else None
        const_refs[key] = ref
    constants = {}
    for key, ref in const_refs.items():
        if ref:
            sh, coord = ref
            sh = _find_sheet(wf, sh) or sh
            constants[key] = eval_simple(wf, wv, sh, coord.replace("$", ""))
            constants[key + "_ref"] = f"{sh}!{coord}"
            constants[key + "_formula"] = ftext(wf[sh][coord].value)
        else:
            alerts.append({"level": "error", "type": "公式解析", "msg": f"未能解析 {key} 引用的均值单元格"})

    x_f = ftext(ws[f"X{first}"].value) or ""
    xm = re.search(r"SUM\(\$W\$(\d+):\$W\$(\d+)\)", x_f)
    pool_rows = (int(xm.group(1)), int(xm.group(2))) if xm else None

    # ---- 2. 逐机构 ----
    rows = []
    r = 2
    while r <= ws.max_row:
        name = ws[f"B{r}"].value
        if not name:
            r += 1
            continue
        name = str(name).strip()
        rec = {"row": r, "seq": len(rows) + 1, "name": name, "sources": {}, "notes": []}

        def from_formula(col, mode="plain"):
            f = ftext(ws[f"{col}{r}"].value)
            if f is None:
                v = _num(ws[f"{col}{r}"].value)
                rec["sources"][col] = "常数" if v is not None else "空"
                return v, True
            specs = parse_filters(f)
            if not specs:
                rec["sources"][col] = f
                return _num(vs[f"{col}{r}"].value), True
            total, ok_all = 0.0, True
            for sp in specs:
                sp["key_sheet"] = _find_sheet(wv, sp["key_sheet"]) or sp["key_sheet"]
                sp["value_sheet"] = _find_sheet(wv, sp["value_sheet"]) or sp["value_sheet"]
                found, v = lk.get(sp, name)
                if not found:
                    ok_all = False
                else:
                    total += v or 0.0
            rec["sources"][col] = " + ".join(f"{s['value_sheet']}!{s['value_col']}" for s in specs)
            if not ok_all:
                if "IFERROR" in f:
                    return (0.0 if mode == "iferror0" else None), True
                return None, False
            return total, True

        rec["scale"], ok_i = from_formula("I")
        rec["mintou_prev"], _ = from_formula("H", "iferror0")
        # K = 就业人数 / (I + 基期常数) × 10000
        kf = ftext(ws[f"K{r}"].value) or ""
        rec["emp"], ok_k = from_formula("K")
        bm = re.search(r"/\(I\d+\+([\d.]+)\)", kf)
        rec["base_const"] = float(bm.group(1)) if bm else 0.0
        if rec["base_const"]:
            rec["notes"].append(f"K 列含基期常数 {rec['base_const']}")
        rec["zxzn_raw"], _ = from_formula("M")
        rec["lj_raw"], _ = from_formula("O")
        rec["dn_raw"], _ = from_formula("Q")
        rec["pj"], _ = from_formula("S")
        # V：注册资本调整，取申请表「增资总额」
        vf = ftext(ws[f"V{r}"].value)
        if vf is None:
            rec["zc_total"] = None
            rec["v_const"] = _num(ws[f"V{r}"].value)
        else:
            specs = parse_filters(vf)
            rec["v_const"] = None
            if specs:
                found, v = lk.get(specs[0], name)
                rec["zc_total"] = v if found else None
                rec["sources"]["V"] = f"{specs[0]['value_sheet']}!{specs[0]['value_col']}（增资总额）"
                if not found:
                    rec["notes"].append("申请表中未找到该机构，增资按 0 处理")
        rec["apply_regular_tech"], _ = from_formula("C")
        rec["mintou_special"], _ = from_formula("D")
        rec["tech_quota"] = _num(vs[f"AE{r}"].value)
        rec["tech_quota1"] = _num(vs[f"AD{r}"].value)
        rec["filed_amount"] = _num(vs[f"AF{r}"].value)
        rec["net_amount"] = _num(vs[f"AG{r}"].value)
        rec["icbc"] = {c: _num(vs[f"{c}{r}"].value) for c in ("E", "F", "G")}
        xf = ftext(ws[f"X{r}"].value)
        in_pool = bool(xf) and pool_rows is not None and pool_rows[0] <= r <= pool_rows[1]
        rec["fixed"] = not in_pool
        rec["fixed_quota"] = _num(ws[f"X{r}"].value) if not xf else None
        if rec["fixed"] and rec["fixed_quota"] is None:
            rec["fixed_quota"] = _num(vs[f"X{r}"].value)
        if isinstance(vs[f"AJ{r}"].value, str) and "固定" in vs[f"AJ{r}"].value:
            rec["fixed"] = True
        rec["fixed_ac"] = _num(vs[f"AC{r}"].value) if rec["fixed"] else None
        rec["excel"] = {c: (vs[f"{c}{r}"].value if not isinstance(vs[f"{c}{r}"].value, str) else vs[f"{c}{r}"].value)
                        for c in CACHED_COLS}
        rec["excel_formulas"] = {c: ftext(ws[f"{c}{r}"].value) for c in ("I", "K", "V", "W", "X", "Z", "AA", "AB", "AC")}
        if not ok_i and not rec["fixed"]:
            alerts.append({"level": "error", "type": "名单不匹配", "msg": f"「{name}」在规模分表中未找到"})
        if not ok_k and not rec["fixed"]:
            alerts.append({"level": "error", "type": "名单不匹配", "msg": f"「{name}」在就业人数分表中未找到"})
        rows.append(rec)
        r += 1

    # ---- 3. 名单齐全性校验（分表 vs 计算表） ----
    calc_names = {norm_name(x["name"]) for x in rows}
    seen_specs = {}
    for c in ("I", "H", "K", "M", "O", "Q", "S", "V", "D"):
        f = ftext(ws[f"{c}{first}"].value)
        for sp in parse_filters(f):
            seen_specs[sp["key_sheet"]] = sp
    for sh, sp in seen_specs.items():
        if sh not in wv.sheetnames:
            continue
        names = lk.keys(sp)
        extra = [n for n in names if norm_name(n) not in calc_names]
        lack = [x["name"] for x in rows if not x["fixed"] and norm_name(x["name"]) not in {norm_name(n) for n in names}]
        if extra:
            alerts.append({"level": "info", "type": "名单差异", "msg": f"分表「{sh}」中有 {len(extra)} 家不在计算表：{'、'.join(extra)}"})
        if lack:
            alerts.append({"level": "warn", "type": "名单差异", "msg": f"计算表中 {len(lack)} 家在分表「{sh}」中缺失：{'、'.join(lack)}"})

    # ---- 4. 以「2025年规模」C 列为基准名单核对 ----
    bl = base_list(wv, smap)
    bln = {norm_name(n) for n in bl}
    not_in_base = [x["name"] for x in rows if norm_name(x["name"]) not in bln]
    if not_in_base:
        alerts.append({"level": "info", "type": "基准名单", "msg": f"计算表中 {len(not_in_base)} 家不在「2025年规模」C 列基准名单："
                       f"{'、'.join(not_in_base)}（均为固定额度机构则属正常）" if all(
                           x["fixed"] for x in rows if x["name"] in not_in_base) else
                       f"计算表中 {len(not_in_base)} 家不在「2025年规模」C 列基准名单：{'、'.join(not_in_base)}"})
    constants = dict(constants)
    constants["base_consts"] = {x["name"]: x["base_const"] for x in rows if x.get("base_const")}
    constants["base_const_ref"] = "计算表 K 列公式（仅这些机构的 K 列分母含基期常数）"

    payload = {
        "sheet_map": smap, "mode": "workbook", "base_list": bl,
        "source_file": filename,
        "sheets": wf.sheetnames,
        "sheets_found": sheets_found,
        "weights_from_workbook": weights,
        "zxzn_threshold": zxzn_threshold,
        "constants": constants,
        "pool_rows": pool_rows,
        "header": header,
        "rows": rows,
        "alerts": alerts,
    }
    return payload
