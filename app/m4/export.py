"""M4 导出：测算表 Excel、《请示》Word（套用模板）。"""
import copy
import re
from io import BytesIO

import docx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from ..config import QINGSHI_TEMPLATE
from . import service

THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="F3E3E4")


def _ws_table(ws, headers, rows, widths=None):
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BORDER
    for r in rows:
        ws.append(r)
        for c in ws[ws.max_row]:
            c.border = BORDER
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "C2"


def _cur_log_sheet(ws, sid):
    log = service.cur_log(sid)
    _ws_table(ws, ["时间", "机构", "调整项", "调整前(万元)", "调整后(万元)", "调整量(万元)", "原因", "操作人"],
              [[x["time"], x["institution"], x["field"], x["before"], x["after"], x["delta"], x["reason"], x["operator"]]
               for x in log], [20, 34, 24, 15, 15, 14, 50, 12])
    for row in ws.iter_rows(min_row=2):
        for c in row[3:6]:
            c.number_format = "#,##0.00"


def _formula_sheet(ws, s):
    """口径说明：公式文件版本、本次结果打戳、参数版本、系数缺失告警、列级白话公式。"""
    from .spec import load_spec
    spec = load_spec()
    res = s.get("result") or {}
    meta = res.get("meta") or {}
    p = res.get("param") or {}
    alerts = [a for a in res.get("alerts") or [] if a.get("type") == "数据缺失"]
    ws["A1"] = spec.get("title") or "授信测算口径说明"
    ws["A1"].font = Font(bold=True, color="B0242E")
    rows = [
        ["项目", "内容"],
        ["公式口径版本（结果打戳）", meta.get("formula_spec_version") or "（本次结果尚未打戳，请重新测算）"],
        ["当前公式文件版本", spec.get("version")],
        ["参数版本", f"v{p.get('version')}（id={meta.get('param_id', p.get('id'))}）" if p else ""],
        ["数据集", meta.get("dataset_id") if meta.get("dataset_id") is not None else s.get("dataset_id")],
        ["calc.py 哈希", meta.get("calc_module_hash") or ""],
        ["方案", f"v{s.get('version')} {s.get('name') or ''}"],
        ["缺系数规则", spec.get("missing_coefficient_rule") or ""],
        ["权重规则", spec.get("weight_rule") or ""],
        ["说明", spec.get("scope") or ""],
        ["本次系数缺失告警", "无" if not alerts else "；".join(a.get("msg") or "" for a in alerts)],
    ]
    if p.get("weights"):
        rows.append(["本次权重", " / ".join(f"w{k}={v}" for k, v in p["weights"].items())])
    _ws_table(ws, rows[0], rows[1:], [28, 88])
    ws.append([])
    ws.append(["列", "指标", "白话公式", "输入", "缺失时", "来源"])
    policies = spec.get("missing_policies") or {}
    for col in spec.get("columns") or []:
        pol = col.get("missing_policy")
        ws.append([col.get("id"), col.get("name") or "", col.get("expr_human"), "、".join(col.get("inputs") or []),
                   f"{pol}：{policies.get(pol, '')}", col.get("source")])
    for cell in ws[ws.max_row - len(spec.get("columns") or [])]:
        cell.font = Font(bold=True)
        cell.fill = HEAD
    for row in ws.iter_rows(min_row=2, max_col=6):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 72
    ws.column_dimensions["D"].width = 36
    ws.column_dimensions["E"].width = 42
    ws.column_dimensions["F"].width = 36
    ws.row_dimensions[1].height = 22


def formula_spec_xlsx(sid):
    s = service.get_scheme(sid)
    wb = Workbook()
    _formula_sheet(wb.active, s)
    wb.active.title = "口径说明"
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue(), f"授信口径说明_v{s['version']}.xlsx"


def cur_log_xlsx(sid):
    s = service.get_scheme(sid, with_result=False)
    wb = Workbook()
    _cur_log_sheet(wb.active, sid)
    wb.active.title = "调整记录"
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue(), f"授信额度调整记录_v{s['version']}.xlsx"


def scheme_xlsx(sid, include_spec=False):
    s = service.get_scheme(sid)
    res = s["result"]
    wb = Workbook()
    ws = wb.active
    ws.title = "测算结果"
    heads = ["序号", "机构", "固定额度", "上一年规模I", "规模系数J", "就业贡献K", "就业系数L", "支小支农M", "系数N",
             "累计代偿率O", "系数P", "当年代偿率Q", "系数R", "评价分S", "系数T", "总系数U", "注册资本调整V", "上一年民投H",
             "第一次调整W", "第二次调整X", "第三次调整(±)", "第三次调整原因", "第三次调整后", "民投完成率调整", "常规授信额度",
             "民间投资专项", "最终授信额度"]
    rows = []
    for r in res["rows"]:
        rows.append([r["seq"], r["name"], "是" if r["fixed"] else "", r["I"], r["J"], r["K"], r["L"], r["M"], r["N"], r["O"],
                     r["P"], r["Q"], r["R"], r["S"], r["T"], r["U"], r["V"], r["H"], r["W"], r["X"], r["third"],
                     r["third_reason"], r["after_third"], r["mintou_adj"], r["regular_final"], r["mintou_special"], r["final"]])
    t = res["totals"]
    rows.append(["", "合计", "", None, None, None, None, None, None, None, None, None, None, None, None, t["sumU"], None, None,
                 t["sumW"], t["sumX"], t["sumThird"], "", None, t["sumMintouAdj"], t["sumRegular"], t["sumMintouSpecial"], t["sumFinal"]])
    _ws_table(ws, heads, rows, [5, 34] + [12] * 25)
    for row in ws.iter_rows(min_row=2):
        for c in row[3:]:
            if isinstance(c.value, float):
                c.number_format = "0.0000" if c.column in range(5, 17) else "#,##0.00"
    ws.insert_rows(1)
    ws["A1"] = f"2026年各合作担保机构授信额度测算表 · 方案 v{s['version']}「{s['name']}」 · 状态：{s['status']} · 单位：万元 · 内部资料 注意保密"
    ws["A1"].font = Font(bold=True, color="B0242E")

    wsp = wb.create_sheet("初步授信额度", 0)
    _ws_table(wsp, ["序号", "机构", "固定额度", "规模系数J", "总系数U", "第一次调整W", "第二次调整X", "第三次Z", "第四次AA",
                    "第五次AB", "截止备案金额AF", "已备案占比AH", "初步授信额度AC", "机构总额度AI", "Excel 台账AC", "AC差异",
                    "目前最新常规业务授信额度", "目前最新民间投资授信额度"],
              [[r["seq"], r["name"], "固定额度" if r["fixed"] else "", r["J"], r["U"], r["W"], r["X"], r.get("Z"), r.get("AA"),
                r.get("AB"), r.get("AF"), r.get("AH"), r.get("AC"), r.get("AI"), r.get("excel_AC"),
                (r.get("AC") or 0) - r["excel_AC"] if isinstance(r.get("excel_AC"), (int, float)) else None,
                r.get("cur_regular"), r.get("cur_mintou")] for r in res["rows"]], [5, 34, 9] + [13] * 15)
    for row in wsp.iter_rows(min_row=2):
        for c in row[3:]:
            if isinstance(c.value, (int, float)):
                c.number_format = "0.0000" if c.column in (4, 5) else ("0.00%" if c.column == 12 else "#,##0.00")
    wsp.insert_rows(1)
    wsp["A1"] = "初步授信额度（AC，公式值）与目前最新授信额度（台账现值 C/D + 人工增减）分栏，互不覆盖 · 单位：万元 · 最终授信以请示批复为准"
    wsp["A1"].font = Font(bold=True, color="B0242E")
    _cur_log_sheet(wb.create_sheet("目前最新额度调整记录"), sid)

    ws2 = wb.create_sheet("人工调整留痕")
    hist = [h for h in service.adjustment_history(sid) if h["kind"] in ("third", "mintou")]
    _ws_table(ws2, ["时间", "操作人", "角色", "机构", "类型", "调整值(万元)", "原因"],
              [[h["created_at"], h["operator"], h["role"], h["institution"], "第三次调整" if h["kind"] == "third" else "民投完成率调整覆盖",
                h["value"], h["reason"]] for h in hist], [20, 12, 10, 34, 16, 14, 60])
    ws3 = wb.create_sheet("参数")
    p = res["param"]
    prow = [["参数版本", f"v{p['version']}"]] + [[f"权重 w{k}", v] for k, v in p["weights"].items()] + \
           [["全省平均户均就业人数", p["constants"].get("emp_avg")], ["累计代偿率全省均值", p["constants"].get("lj_avg")],
            ["当年代偿率全省均值", p["constants"].get("dn_avg")], ["支小支农阈值", p["constants"].get("zxzn_threshold")],
            ["系数下限", p["constants"].get("floor")]]
    _ws_table(ws3, ["参数", "值"], prow, [26, 24])
    ws4 = wb.create_sheet("状态记录")
    _ws_table(ws4, ["时间", "操作人", "原状态", "新状态", "备注"],
              [[x["at"], x["by"], x["from"], x["to"], x.get("remark", "")] for x in s["status_log"]], [20, 12, 10, 10, 40])
    if include_spec:
        _formula_sheet(wb.create_sheet("口径说明"), s)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue(), f"授信测算_v{s['version']}.xlsx"


# ---------------- 请示 Word ----------------
def _fmt(v):
    return f"{v:,.2f}"


def qingshi_defaults(sid, base_sid=None):
    """生成《请示》各段落默认文本：模板原文 + 依据方案自动生成的“调整汇总”段与附件表。

    调整口径：若给定 base_sid（对比基准方案），按两方案「常规授信额度」差额统计调增/调减；
    否则按本方案的「第三次调整 + 民投完成率调整」统计。"""
    d = docx.Document(str(QINGSHI_TEMPLATE))
    paras = [p.text for p in d.paragraphs]
    s = service.get_scheme(sid)
    rows = s["result"]["rows"]
    if base_sid:
        b = {r["name"]: r for r in service.get_scheme(base_sid)["result"]["rows"]}
        deltas = [(r["name"], r["regular_final"] - (b.get(r["name"], {}).get("regular_final") or 0)) for r in rows]
    else:
        deltas = [(r["name"], (r["third"] or 0) + (r["mintou_adj"] or 0)) for r in rows]
    dec = sorted([x for x in deltas if x[1] < -1e-6], key=lambda x: x[1])
    inc = sorted([x for x in deltas if x[1] > 1e-6], key=lambda x: -x[1])
    sd, si = -sum(v for _, v in dec), sum(v for _, v in inc)
    summary = (f"综上，本次调整共涉及{len(dec) + len(inc)}家机构，其中需调减额度的机构{len(dec)}家，合计{sd / 10000:.2f}亿元；"
               f"需调增额度的机构{len(inc)}家，合计{si / 10000:.2f}亿元。")
    idx = next((i for i, t in enumerate(paras) if t.startswith("综上，本次调整")), None)
    if idx is not None:
        paras[idx] = summary
    return {"paragraphs": paras, "summary_index": idx,
            "decrease": [{"name": n, "value": -v} for n, v in dec],
            "increase": [{"name": n, "value": v} for n, v in inc],
            "scheme": {"id": s["id"], "version": s["version"], "name": s["name"], "status": s["status"]}}


def _set_para_text(p, text):
    if p.text == text:
        return
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def qingshi_docx(sid, paragraphs=None, base_sid=None):
    dflt = qingshi_defaults(sid, base_sid)
    paras = paragraphs or dflt["paragraphs"]
    d = docx.Document(str(QINGSHI_TEMPLATE))
    for p, t in zip(d.paragraphs, paras):
        _set_para_text(p, t)
    # 附件1表格（模板第 2 张表：序号|机构名称|调减额度|机构名称|调增额度）
    tbl = d.tables[1] if len(d.tables) > 1 else None
    if tbl is not None:
        dec, inc = dflt["decrease"], dflt["increase"]
        n = max(len(dec), len(inc), 1)
        tmpl_row = tbl.rows[1]._tr
        total_tr = tbl.rows[-1]._tr
        for r in list(tbl.rows)[1:-1]:
            tbl._tbl.remove(r._tr)
        for i in range(n):
            new = copy.deepcopy(tmpl_row)
            total_tr.addprevious(new)
        data_rows = list(tbl.rows)[1:-1]
        for i, row in enumerate(data_rows):
            vals = [str(i + 1),
                    dec[i]["name"] if i < len(dec) else "", _fmt(dec[i]["value"]) if i < len(dec) else "",
                    inc[i]["name"] if i < len(inc) else "", _fmt(inc[i]["value"]) if i < len(inc) else ""]
            for c, v in zip(row.cells, vals):
                _set_para_text(c.paragraphs[0], v)
        tot = tbl.rows[-1].cells
        _set_para_text(tot[2].paragraphs[0], f"{sum(x['value'] for x in dec):,.0f}")
        _set_para_text(tot[4].paragraphs[0], f"{sum(x['value'] for x in inc):,.0f}")
    bio = BytesIO()
    d.save(bio)
    s = dflt["scheme"]
    return bio.getvalue(), f"关于调整部分合作担保机构授信额度的请示_方案v{s['version']}.docx"
