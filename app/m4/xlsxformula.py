"""读取附件2工作簿的「公式」并据此解析输入口径。

不是通用 Excel 引擎：只识别附件2计算表中实际出现的公式形态（FILTER 查找、IFERROR、ROUND、
四则运算），用于确定每个系数的数据来源（哪张分表、哪一列）、基期常数、权重、全省均值单元格等。
"""
import re
from decimal import Decimal, ROUND_HALF_UP

from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.formula import ArrayFormula


def xround(x, n=0):
    """Excel ROUND：四舍五入（远离 0），而非 Python 的银行家舍入。"""
    if x is None:
        return None
    d = Decimal(repr(float(x))).quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP)
    return int(d) if n == 0 else float(d)


def ftext(cell_value):
    """返回单元格公式文本（含数组公式），非公式返回 None。"""
    v = cell_value
    if isinstance(v, ArrayFormula):
        v = v.text
    if isinstance(v, str) and v.startswith("="):
        return v
    return None


_SHEET = r"(?:'(?P<{0}q>[^']+)'|(?P<{0}u>[^'!,()=]+))"
FILTER_RE = re.compile(
    r"FILTER\(" + _SHEET.format("vs") + r"!\$?(?P<vc>[A-Z]+)\$?(?P<v1>\d+):\$?[A-Z]+\$?(?P<v2>\d+),"
    + _SHEET.format("ks") + r"!\$?(?P<kc>[A-Z]+)\$?(?P<k1>\d+):\$?[A-Z]+\$?(?P<k2>\d+)=B\d+\)"
)


def parse_filters(formula):
    out = []
    for m in FILTER_RE.finditer(formula or ""):
        g = m.groupdict()
        out.append({
            "value_sheet": g["vsq"] or g["vsu"], "value_col": g["vc"],
            "key_sheet": g["ksq"] or g["ksu"], "key_col": g["kc"],
            "row_from": int(g["v1"]), "row_to": int(g["v2"]),
        })
    return out


def norm_name(s):
    return re.sub(r"\s+", "", str(s or "")).replace("(", "（").replace(")", "）")


class Lookup:
    """模拟 FILTER(值列, 键列=机构名)：返回 (found, value)。空单元格按 Excel 规则取 0。"""

    def __init__(self, wb_values):
        self.wb = wb_values

    def get(self, spec, name):
        ws = self.wb[spec["key_sheet"]]
        wv = self.wb[spec["value_sheet"]]
        kc = column_index_from_string(spec["key_col"])
        vc = column_index_from_string(spec["value_col"])
        target = norm_name(name)
        for r in range(spec["row_from"], spec["row_to"] + 1):
            if norm_name(ws.cell(r, kc).value) == target:
                v = wv.cell(r, vc).value
                if v in (None, ""):
                    return True, 0.0
                try:
                    return True, float(v)
                except (TypeError, ValueError):
                    return True, None
        return False, None

    def keys(self, spec):
        ws = self.wb[spec["key_sheet"]]
        kc = column_index_from_string(spec["key_col"])
        out = []
        for r in range(spec["row_from"], spec["row_to"] + 1):
            v = ws.cell(r, kc).value
            if v not in (None, "") and str(v).strip() not in ("合计", "总计"):
                out.append(str(v).strip())
        return out


REF_RE = re.compile(r"(?:(?:'(?P<sq>[^']+)'|(?P<su>[^'!,()=+\-*/]+))!)?\$?(?P<c>[A-Z]{1,3})\$?(?P<r>\d+)")


def eval_simple(wb_f, wb_v, sheet, coord, depth=0):
    """递归求值只含四则运算与单元格引用的公式（如 就业人数!F3 = E6/8124032.18*10000）。"""
    if depth > 20:
        raise ValueError("公式嵌套过深")
    cell = wb_f[sheet][coord]
    f = ftext(cell.value)
    if f is None:
        v = wb_v[sheet][coord].value
        return float(v or 0)
    expr = f[1:]
    if re.search(r"[A-Z]+\(", expr):  # 含函数，退回缓存值
        return float(wb_v[sheet][coord].value or 0)

    def rep(m):
        sh = m.group("sq") or m.group("su") or sheet
        return repr(eval_simple(wb_f, wb_v, sh, f"{m.group('c')}{m.group('r')}", depth + 1))

    py = REF_RE.sub(rep, expr).replace("%", "/100")
    if not re.fullmatch(r"[0-9eE.+\-*/() ]+", py):
        raise ValueError(f"无法求值的公式 {sheet}!{coord}: {f}")
    return float(eval(py, {"__builtins__": {}}, {}))  # noqa: S307 - 已白名单校验


def parse_ref(formula_part):
    m = REF_RE.search(formula_part)
    if not m:
        return None
    return (m.group("sq") or m.group("su")), f"{m.group('c')}{m.group('r')}"
