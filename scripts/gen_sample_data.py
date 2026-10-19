"""生成 samples/ 下全部虚构示例数据（固定随机种子，可复现）。

    python scripts/gen_sample_data.py      # 需要 LibreOffice（soffice）用于重算公式缓存值

所有机构名称与数值均为随机虚构，不对应任何真实机构或业务。
"""
import random
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

ROOT = Path(__file__).resolve().parent.parent
M4_OUT = ROOT / "samples" / "m4" / "附件2：2026年各合作担保机构授信额度计算表（示例）.xlsx"
M5_DIR = ROOT / "samples" / "m5"
R = random.Random(20261019)
POOL = [f"示例担保公司{c}" for c in "ABCDEFGH"]
FIXED = ["示例担保公司I", "示例担保公司J"]
ALL = POOL + FIXED
DIRECT = "示例省级融资再担保有限责任公司"
BASE_CONST = 50000.0  # 仅「示例担保公司A」K 列分母含基期常数（演示该口径）


def rnd(a, b, nd=2):
    return round(R.uniform(a, b), nd)


def fa(ws, coord, text):
    ws[coord] = ArrayFormula(coord, text)


def ffilter(vs, vc, ks, kc, r1, r2, b):
    q = lambda s: f"'{s}'" if s[0].isdigit() else s  # noqa: E731
    return f"_xlfn._xlws.FILTER({q(vs)}!${vc}${r1}:${vc}${r2},{q(ks)}!${kc}${r1}:${kc}${r2}={b})"


def scale_sheet(ws, title, period, names_vals, value_label="规模"):
    ws["B2"] = title
    ws["B3"] = period
    ws["K3"] = "单位：万元、笔、户、%"
    ws["B4"], ws["C4"], ws["D4"], ws["G4"], ws["J4"] = "序号", "机构名称", value_label, "笔数", "户数"
    for c, v in zip("DEFGHIJKL", ["规模", "占比", "同比", "笔数", "占比", "同比", "户数", "占比", "同比"]):
        ws[f"{c}5"] = v
    tot = sum(v for _, v in names_vals)
    cnt = [R.randint(50, 3000) for _ in names_vals]
    ws["B6"], ws["D6"], ws["E6"], ws["F6"], ws["G6"], ws["J6"] = "合计", round(tot, 6), "-", rnd(-10, 30), sum(cnt), sum(int(c * .8) for c in cnt)
    r = 7
    for i, ((n, v), c) in enumerate(zip(names_vals, cnt)):
        ws.append([None, i + 1, n, v, v / tot * 100 if tot else 0, rnd(-20, 60), c, c / sum(cnt) * 100, rnd(-20, 40), int(c * .8),
                   int(c * .8) / max(1, sum(int(x * .8) for x in cnt)) * 100, rnd(-20, 40)])
        r += 1
    ws.append([None, "备注：", "1、业务规模：统计期内已确认备案的业务累计主债权金额合计（示例数据，虚构）。"])
    ws.append([None, None, "2、占比：统计金额÷总数×100%。"])
    return r - 1


def build_m4():
    wb = openpyxl.Workbook()
    calc = wb.active
    calc.title = "计算表"
    names = ["合作业务授信额度申请表", "2025年民投", "2025年规模", "就业人数", "支小支农业务统计表", "近3年累计代偿率",
             "2025年代偿率", "2025年机构评价", "2025年互联网落地"]
    sh = {n: wb.create_sheet(n) for n in names}
    scale = {n: rnd(30000, 400000, 6) for n in POOL + FIXED[:1]}
    # 申请表
    ap = sh["合作业务授信额度申请表"]
    ap["B1"] = "合作业务授信额度申请表（示例）"
    for c, v in zip("ABCDEFGHI", ["序号", "合作机构", "2026年度常规业务申请额度（不含科技担保专项、民间投资专项及国担“总对总”业务额度）",
                                  "2026年度科技担保专项申请额度", "民间投资专项担保计划申请额度", None, None, "2025年注册资本增加情况", None]):
        ap[f"{c}3"] = v
    for c, v in zip("EFGHI", ["小计", "2026年度", "2027年度", "是否增资", "增资总额"]):
        ap[f"{c}4"] = v
    for i, n in enumerate(ALL):
        f26 = R.choice([0, 20000, 30000, 50000, 80000]) if n != FIXED[1] else 0
        zc = rnd(500, 15000) if R.random() < .5 else None
        ap.append([i + 1, n, R.randrange(50, 700) * 1000, R.choice([0, 10000, 20000, 50000]), f26 * 2, f26, f26, "是" if zc else "否", zc])
    last = 4 + len(ALL)
    ap.append(["合计", None, f"=SUM(C5:C{last})", f"=SUM(D5:D{last})", f"=SUM(E5:E{last})", f"=SUM(F5:F{last})", f"=SUM(G5:G{last})"])
    # 民投
    mt = sh["2025年民投"]
    mt.append([None, "机构", "符合民投", "邮储", "符合民投（扣除邮储）", None, "直担机构名称", "求和项:主债权金额（万元）"])
    for n in POOL:
        v = rnd(0, scale[n] * .3, 6)
        p = rnd(0, v * .2, 6)
        mt.append([None, n, v, p, v - p, None, n, v - p])
    mt.append([None, "总计"])
    # 规模
    scale_sheet(sh["2025年规模"], "机构业务规模分析", "备案时间： 2025-01-01 - 2025-12-31（示例数据）", [(n, scale[n]) for n in scale])
    # 就业人数
    em = sh["就业人数"]
    emp = {n: int(scale[n] / 10000 * R.uniform(700, 1300)) for n in scale}
    tot_scale = round(sum(scale.values()) * 1.25, 2)
    em["B2"] = "新增业务就业规模机构分析（示例）"
    em["B3"], em["E3"], em["F3"] = "统计年度：2025年（示例数据）", "全省平均", f"=E6/{tot_scale}*10000"
    em["B4"], em["C4"], em["D4"], em["D5"], em["E5"], em["F5"] = "序号", "机构名称", "合计", "户数", "人数", "户均"
    em["B6"], em["D6"], em["E6"] = "合计", sum(int(v / 8) for v in emp.values()), int(sum(emp.values()) * 1.2)
    for i, n in enumerate(scale):
        em.append([None, i + 1, n, int(emp[n] / 8), emp[n], 8])
    em.append([None, "备注：", "1、示例数据，虚构。"])
    # 支小支农
    zx = sh["支小支农业务统计表"]
    zx["B2"], zx["B4"], zx["C4"], zx["D4"], zx["D5"], zx["E5"] = "支小支农业务统计表（示例）", "序号", "机构名称", "支小支农", "业务规模", "规模占比"
    zx["B6"] = "合计"
    for i, n in enumerate(scale):
        zx.append([None, i + 1, n, rnd(scale[n] * .6, scale[n], 6), rnd(70, 100, 4)])
    # 代偿率两张
    for sname in ("近3年累计代偿率", "2025年代偿率"):
        s = sh[sname]
        h, c = rnd(2_000_000, 6_000_000), None
        c = round(h * R.uniform(.015, .035), 2)
        s["B2"], s["G3"], s["H3"], s["I3"], s["J3"] = "机构代偿业务分析（示例）", "再保代偿率", h, c, "=I3/H3"
        s["B4"], s["C4"], s["J4"] = "序号", "机构名称", "担保代偿率\n（监管口径）"
        s["B6"] = "合计"
        for i, n in enumerate(scale):
            s.append([None, i + 1, n, None, None, None, None, None, None, rnd(0.3, 4.5, 6)])
    pj = sh["2025年机构评价"]
    pj.append(["序号", "机构名称", "分数"])
    for i, n in enumerate(scale):
        pj.append([f"=ROW()-1", n, rnd(60, 95)])
    net = sh["2025年互联网落地"]
    net.append([None, None, None])
    net.append([None, "直担机构名称", "求和项:主债权金额（万元）"])
    for n in POOL[:5]:
        net.append([None, n, rnd(5000, 120000, 6)])
    net.append([None, "总计"])

    # 计算表
    hdr = ["序号", "机构", "2026年常规（含创业、科技）额度", "2026年民间专项额度", "2026年工行预计规模(地方版，不含总对总)", "2025年工行实际发生金额",
           "2026年预计工行发生额（以2025年增加30%测算）", "2025年符合民投金额（扣工行、邮储、展期类、扣除银行分险不超过20%）", "2025年规模", "系数",
           "就业贡献指标", "系数", "支小支农比例", "系数", "累计担保代偿率", "系数", "2025年担保代偿率", "系数", "2025年机构评价数", "系数", "总系数",
           "根据注册资本调整", "2026年第一次调整后", "2026年第二次调整后", "是否满足工行", "第三次调整", "第四次调整", "第五次调整", "第六次调整",
           "其中科技创新计划额度(第一次)", "其中科技创新计划额度（第二次）", "截止2026年3月16日备案金额", "2025年互联网银行落地规模", "已备案占比",
           "机构总额度（第六次+民投）", "变小机构"]
    calc.append(hdr)
    AP, N0, N1 = "合作业务授信额度申请表", 5, 5 + len(ALL) - 1
    p_last = 1 + len(POOL)
    for i, n in enumerate(ALL):
        r = i + 2
        b = f"B{r}"
        calc[f"A{r}"] = f"=MAX($A$1:A{r - 1})+1"
        calc[b] = n
        fa(calc, f"C{r}", "=" + ffilter(AP, "C", AP, "B", N0, N1, b) + "+" + ffilter(AP, "D", AP, "B", N0, N1, b))
        fa(calc, f"D{r}", "=" + ffilter(AP, "F", AP, "B", N0, N1, b))
        fa(calc, f"H{r}", "=IFERROR(" + ffilter("2025年民投", "H", "2025年民投", "G", 2, 2 + len(POOL), b) + ",0)")
        fa(calc, f"AD{r}", f"=ROUND(IFERROR(400000*{ffilter(AP, 'D', AP, 'B', N0, N1, b)}/SUM({AP}!$D${N0}:$D${N1}),\"\")/1000,0)*1000")
        fa(calc, f"AG{r}", "=IFERROR(" + ffilter("2025年互联网落地", "C", "2025年互联网落地", "B", 3, 9, b) + ',"")')
        calc[f"AE{r}"] = R.randrange(0, 60) * 1000
        calc[f"AI{r}"] = f"=AC{r}+D{r}"
        if i < 3:
            calc[f"E{r}"] = R.randrange(5, 30) * 10000
            calc[f"F{r}"] = rnd(5000, 80000)
        calc[f"G{r}"] = f"=F{r}*1.3"
        if n != FIXED[1]:
            fa(calc, f"I{r}", "=" + ffilter("2025年规模", "D", "2025年规模", "C", 7, 7 + len(scale), b))
            calc[f"J{r}"] = f"=I{r}/10000"
            kden = f"(I{r}+{BASE_CONST})" if n == POOL[0] else f"I{r}"
            fa(calc, f"K{r}", "=" + ffilter("就业人数", "E", "就业人数", "C", 7, 7 + len(scale), b) + f"/{kden}*10000")
            calc[f"L{r}"] = f"=K{r}/就业人数!$F$3"
            fa(calc, f"M{r}", "=ROUND(" + ffilter("支小支农业务统计表", "E", "支小支农业务统计表", "C", 7, 7 + len(scale), b) + ",2)/100")
            calc[f"N{r}"] = f"=_xlfn.SWITCH(TRUE,M{r}>=90%,1,M{r}/90%)"
            fa(calc, f"O{r}", "=" + ffilter("近3年累计代偿率", "J", "近3年累计代偿率", "C", 7, 7 + len(scale), b) + "/100")
            calc[f"P{r}"] = f"=MAX(1-(O{r}-近3年累计代偿率!$J$3)/近3年累计代偿率!$J$3,0.5)"
            fa(calc, f"Q{r}", "=" + ffilter("2025年代偿率", "J", "2025年代偿率", "C", 7, 7 + len(scale), b) + "/100")
            calc[f"R{r}"] = f"=MAX(1-(Q{r}-'2025年代偿率'!$J$3)/'2025年代偿率'!$J$3,0.5)"
            fa(calc, f"S{r}", "=" + ffilter("2025年机构评价", "C", "2025年机构评价", "B", 2, 1 + len(scale), b))
            calc[f"T{r}"] = f"=S{r}/100"
            calc[f"U{r}"] = f"=J{r}*(L{r}*0.2+N{r}*0.2+P{r}*0.1+R{r}*0.1+T{r}*0.4)"
            fa(calc, f"V{r}", f"=ROUND(IFERROR({ffilter(AP, 'I', AP, 'B', N0, N1, b)},\"\")*2/1000,0)*1000")
            calc[f"W{r}"] = f"=I{r}+V{r}-H{r}"
        if n in POOL:
            calc[f"X{r}"] = f"=ROUND((SUM($W$2:$W${p_last}))*U{r}/SUM($U$2:$U${p_last}),0)"
            calc[f"Y{r}"] = f'=IF(G{r}>0,X{r}-G{r},"")'
            calc[f"Z{r}"] = f"=X{r}"
            calc[f"AA{r}"] = f"=Z{r}"
            calc[f"AB{r}"] = f"=ROUND(IF(ROUND(AA{r}/10000,0)*10000=0,10000,ROUND(AA{r}/1000,0)*1000)/10000,0)*10000"
            calc[f"AC{r}"] = f"=AB{r}+IF(AH{r}>=40%,20000,0)"
            calc[f"AF{r}"] = rnd(scale[n] * .1, scale[n] * .6, 6)
            calc[f"AH{r}"] = f"=AF{r}/AB{r}"
            calc[f"AJ{r}"] = f"=AC{r}-I{r}"
        else:
            q = 10000 if n == FIXED[0] else 2000
            for c in ("X", "Z", "AA", "AB", "AC"):
                calc[f"{c}{r}"] = q
            calc[f"AF{r}"] = rnd(q * .2, q * .8)
            calc[f"AJ{r}"] = "固定额度"
            calc[f"W{r}"] = f"=I{r}+V{r}-H{r}"
    tmp = Path(tempfile.mkdtemp()) / "m4.xlsx"
    wb.save(tmp)
    return tmp


def ledger_book(path, period, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append([None] * 13)
    scale_sheet(ws, "机构业务规模分析", period, rows)
    wb.save(path)


def build_m5(m4_vals):
    """m4_vals: {机构: (常规额度, 民间额度)}（取自重算后的 M4 计算表，并做小幅“年中调整”）。"""
    period = "备案时间： 2026-01-01 - 2026-09-19（数据更新日期：2026-09-29）"
    regular = {n: rnd(m4_vals[n][0] * .3, m4_vals[n][0] * .95, 6) for n in ALL}
    tech = {n: rnd(0, 60000, 6) for n in ALL[:6]}
    private = {n: rnd(0, max(m4_vals[n][1], 1) * .9, 6) for n in ALL if m4_vals[n][1]}
    t2t = {n: rnd(0, 150000, 6) for n in ALL[:4]}
    direct = {"regular": rnd(800000, 1500000, 6), "tech": rnd(100000, 300000, 6), "private": rnd(50000, 200000, 6), "t2t": rnd(100000, 400000, 6)}
    for fn, port, d in (("2-1.常规.xlsx", "regular", regular), ("2-2.科创.xlsx", "tech", tech), ("2-3.民投.xlsx", "private", private),
                        ("2-4.总对总.xlsx", "t2t", t2t)):
        ledger_book(M5_DIR / fn, period, [(DIRECT, direct[port])] + list(d.items()))
    # 汇总表
    wb = openpyxl.Workbook()
    s1 = wb.active
    s1.title = "业务整体情况"
    s2 = wb.create_sheet("分机构额度使用情况")
    s2.append(["序号", "机构", "常规业务授信额度（万元）(示例：年中调整后)", "常规业务", None, "科技创新计划额度（万元）", "科技专项", None,
               "民间专项计划额度（万元）", "民间专项", None, "总对总", "机构合计备案金额", None, None])
    s2.append([None, None, None, "截止2026年9月19日常规业务备案金额", "额度使用比例", None, "截止2026年9月19日科技创新计划备案金额", "额度使用比例", None,
               "截止2026年9月19日民间专项备案金额", "额度使用比例", "截止2026年9月19日总对总备案金额", None, "常规备案占再保总备案比例", "民投备案占再保总备案比例"])
    tr = 3 + len(ALL)
    for i, n in enumerate(ALL):
        r = i + 3
        lim = round(m4_vals[n][0] * R.choice([1, 1, .95, 1.05]) / 1000) * 1000
        pl = m4_vals[n][1]
        s2.append([f"=ROW()-2", n, lim, regular.get(n, 0), f"=D{r}/C{r}", "无上限额度，先到先得" if i == 0 else None, tech.get(n, 0), None,
                   pl, private.get(n, 0), f"=J{r}/I{r}" if pl else "未申报额度", t2t.get(n, 0), f"=SUM(D{r},G{r},J{r},L{r})",
                   f"=D{r}/$D${tr}", f"=J{r}/$J${tr}"])
    s2.append([None, "合计", f"=SUM(C3:C{tr - 1})", f"=SUM(D3:D{tr - 1})", f"=D{tr}/C{tr}", None, f"=SUM(G3:G{tr - 1})", None,
               f"=SUM(I3:I{tr - 1})", f"=SUM(J3:J{tr - 1})", f"=J{tr}/I{tr}", f"=SUM(L3:L{tr - 1})", f"=SUM(M3:M{tr - 1})",
               f"=D{tr}/$D${tr}", f"=J{tr}/$J${tr}"])
    tot = {k: direct[k] + sum(d.values()) for k, d in (("regular", regular), ("tech", tech), ("private", private), ("t2t", t2t))}
    reg_alloc = sum(round(m4_vals[n][0] * 1) for n in ALL)
    g_reg = round((tot["regular"] * 1.15 + reg_alloc) / 100000) * 100000 + 500000
    pri_const = round(sum(m4_vals[n][1] for n in ALL) * 0.97 / 10000) * 10000  # 手填常数（演示与分机构合计不一致的告警）
    s1.append(["项目", "常规额度", None, None, None, "科技专项额度", None, None, None, "民间专项额度", None, None, None, "总对总额度", None, None])
    s1.append([None] + ["截止2026年9月19日备案金额", "额度分配", "使用率", "剩余额度"] * 3 + ["截止2026年9月19日备案金额", "使用率", "剩余额度"])
    dalloc = round(direct["regular"] * 1.1 / 10000) * 10000
    dpri = round(direct["private"] * 1.2 / 10000) * 10000
    s1.append(["直保", direct["regular"], dalloc, "=B3/C3", "=C3-B3", direct["tech"], "无上限额度，先到先得", "=F3/$F$6", None,
               direct["private"], dpri, "=J3/K3", "=K3-J3", direct["t2t"], "=N3/N6", None])
    s1.append(["再保", "=B5-B3", f"='分机构额度使用情况'!C{tr}", "=B4/C4", "=C4-B4", "=F5-F3", "无上限额度，先到先得", "=F4/$F$6", None,
               "=J5-J3", pri_const, "=J4/K4", "=K4-J4", "=N5-N3", "=N4/N6", None])
    s1.append(["合计", tot["regular"], "=SUM(C3:C4)", "=B5/B6", "=B6-B5", tot["tech"], None, "=F5/$F$6", "=F6-F5", tot["private"], "=SUM(K3:K4)",
               "=J5/K5", "=K5-J5", tot["t2t"], "=N5/N6", None])
    s1.append(["2026年国担额度", g_reg, None, None, None, 500000, None, None, None, dpri + pri_const, None, None, None, round(tot["t2t"] * 1.4 / 10000) * 10000,
               None, "=N6-N5"])
    s1.append(["剩余未分配额度", None, "=B6-C5", None, None, None, None, None, None, 0])
    s1.append(["需调整", None, "10000（示例，已上会确认）"])
    tmp = Path(tempfile.mkdtemp()) / "m5.xlsx"
    wb.save(tmp)
    return tmp


def recalc(src, dst):
    out = Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                   check=True, capture_output=True)
    shutil.move(str(out / src.name), dst)


def main():
    M4_OUT.parent.mkdir(parents=True, exist_ok=True)
    recalc(build_m4(), M4_OUT)
    ws = openpyxl.load_workbook(M4_OUT, data_only=True)["计算表"]
    vals = {}
    for r in range(2, ws.max_row + 1):
        n = ws[f"B{r}"].value
        if n:
            vals[n] = (ws[f"AC{r}"].value or 0, ws[f"D{r}"].value or 0)
    recalc(build_m5(vals), M5_DIR / "1.业务额度使用情况 - 截止2026年9月19日v1.xlsx")
    print("ok", vals)


if __name__ == "__main__":
    main()
